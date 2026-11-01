"""R461. The parameter sheet of button 1, the twin of ParameterSheet of Common: config/shared_parameters.xlsx that addin.parameterSheet
of config names, sheet Parameters, one row per shared parameter with the columns Parameter, GUID, Data type, Group, Category, Filled
by, Sample value and Keep. A row of Keep No is left out, Keep is Yes or No and Filled by a button from 1 to 9 or none, and
RRE_ZZ_SP_Prototype-FFL is never given a button (R465). The messages are those of the C# side word for word.
"""
import uuid
from pathlib import Path

from openpyxl import load_workbook

SHEET = "Parameters"
COLUMNS = ("Parameter", "GUID", "Data type", "Group", "Category", "Filled by", "Sample value", "Keep")
NONE = "none"
FFL = "RRE_ZZ_SP_Prototype-FFL"


class SheetError(ValueError):
    pass


def path_of(config_path, settings):
    """addin.parameterSheet, relative to the repo root, the folder above the one config is in."""
    named = (settings.get("addin") or {}).get("parameterSheet") if isinstance(settings, dict) else None
    if not isinstance(named, str) or not named.strip():
        raise SheetError("config names no addin.parameterSheet, the parameter sheet of button 1 (R461)")
    path = Path(named)
    return path if path.is_absolute() else Path(config_path).resolve().parent.parent / path


def cell(value):
    text = None if value is None else str(value).strip()
    return text or None


def rows_of(path):
    """The rows of sheet Parameters as openpyxl gives them read only with values only, a workbook closed after the read."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in book.sheetnames:
            raise SheetError("the parameter sheet %s has no sheet %s (R462)" % (Path(path).name, SHEET))
        return [list(row) for row in book[SHEET].iter_rows(values_only=True)]
    finally:
        book.close()


def keep(value, name, sheet):
    if value is not None and value.lower() == "yes":
        return True
    if value is not None and value.lower() == "no":
        return False
    raise SheetError("the parameter sheet %s gives %s the Keep %s, Yes or No (R461)" % (sheet, name, value or "blank"))


def filled_by(value, name, sheet):
    if value is not None and value.lower() == NONE:
        return None
    if value is not None and "1" <= value[0] <= "9" and (len(value) == 1 or value[1] == " "):
        return int(value[0])
    raise SheetError("the parameter sheet %s gives %s the Filled by %s, a button from 1 to 9 or %s (R465)" % (sheet, name, value or "blank", NONE))


def from_rows(sheet, rows):
    """The kept parameters in sheet order, each a dict of name, guid, dataType, group, categories and filledBy, and the names left out."""
    rows = [row for row in rows or []]
    if not rows:
        raise SheetError("the parameter sheet %s holds no row (Q155)" % sheet)
    header = [(cell(value) or "").lower() for value in rows[0]]
    at = [header.index(column.lower()) if column.lower() in header else -1 for column in COLUMNS]
    missing = [column for column, index in zip(COLUMNS, at) if index < 0]
    if missing:
        raise SheetError("the parameter sheet %s has no column %s in its first row (R461)" % (sheet, ", ".join(missing)))
    kept, left_out = [], []
    for row in rows[1:]:
        cells = [cell(row[index]) if index < len(row) else None for index in at]
        if all(value is None for value in cells):
            continue
        if cells[0] is None:
            raise SheetError("the parameter sheet %s holds a parameter with no name (Q155)" % sheet)
        if not keep(cells[7], cells[0], sheet):
            left_out.append(cells[0])
            continue
        kept.append({"name": cells[0], "guid": cells[1], "dataType": cells[2].upper() if cells[2] else None, "group": cells[3],
                     "categories": [part.strip() for part in (cells[4] or "").split(",") if part.strip()], "filledBy": filled_by(cells[5], cells[0], sheet)})
    names = [item["name"] for item in kept] + left_out
    twice = sorted({name for name in names if names.count(name) > 1}, key=names.index)
    if twice:
        raise SheetError("the parameter sheet %s names %s twice, one parameter per name (Q155)" % (sheet, ", ".join(twice)))
    for item in kept:
        if item["guid"] is None:
            continue
        try:
            parsed = uuid.UUID(item["guid"])
        except ValueError:
            parsed = None
        if parsed is None or len(item["guid"]) != 36:
            raise SheetError("the parameter sheet %s gives %s the GUID %s, which is no GUID (Q155)" % (sheet, item["name"], item["guid"]))
        item["guid"] = str(parsed)
    ffl = next((item for item in kept if item["name"] == FFL), None)
    if ffl is not None and ffl["filledBy"] is not None:
        raise SheetError("the parameter sheet %s gives %s to button %d, it is never written, Filled by %s (R465)" % (sheet, FFL, ffl["filledBy"], NONE))
    return kept, left_out


def from_file(path):
    return from_rows(Path(path).name, rows_of(path))
