import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from common import write_json

HEADER_MARKER = "Plot No."
# Columns that keep 4 decimals, matched on the header text without trailing stars
WIDE_COLUMNS = {"el @ entry", "height"}
EXTRA_LIMIT = 200


def find_sheet(workbook, wanted):
    for name in workbook.sheetnames:
        if name.strip().lower() == wanted.strip().lower():
            return workbook[name]
    raise ValueError("Sheet '" + wanted + "' not found. Sheets present: " + ", ".join(workbook.sheetnames))


def unique_headers(first_row):
    headers = []
    seen = {}
    for index, cell in enumerate(first_row, start=1):
        text = "" if cell is None else str(cell).strip()
        if not text:
            text = "col_" + get_column_letter(index)
        seen[text] = seen.get(text, 0) + 1
        headers.append(text if seen[text] == 1 else text + "_" + str(seen[text]))
    return headers


def is_blank(cell):
    return cell is None or str(cell).strip() == ""


def find_header_row(rows):
    marker = HEADER_MARKER.lower()
    for index, row in enumerate(rows):
        if any(not is_blank(cell) and marker in str(cell).lower() for cell in row):
            return index
    return None


def is_total_row(row, plot_column):
    if any(isinstance(cell, str) and cell.strip().lower().startswith("total") for cell in row):
        return True
    # The total row of the real workbook has no label, only a number and an empty plot number cell
    return is_blank(row[plot_column]) if plot_column < len(row) else True


def wide_headers(headers):
    return {header for header in headers if header.strip().rstrip("*").strip().lower() in WIDE_COLUMNS}


def read_sheet(xlsx_path, sheet_name):
    workbook = load_workbook(str(xlsx_path), read_only=True, data_only=True)
    try:
        sheet = find_sheet(workbook, sheet_name)
        rows = [row for row in sheet.iter_rows(values_only=True)]
        sheet_title = sheet.title
    finally:
        workbook.close()
    header_index = find_header_row(rows)
    if header_index is None:
        raise ValueError("No row containing '" + HEADER_MARKER + "' in sheet '" + sheet_title + "'")
    header_row = rows[header_index]
    headers = unique_headers(header_row)
    plot_column = next(index for index, cell in enumerate(header_row) if not is_blank(cell) and HEADER_MARKER.lower() in str(cell).lower())
    records = []
    for offset, row in enumerate(rows[header_index + 1:]):
        if all(is_blank(cell) for cell in row) or is_total_row(row, plot_column):
            break
        record = {"sheetRow": header_index + 2 + offset}
        for header, cell in zip(headers, row):
            record[header] = cell.isoformat() if hasattr(cell, "isoformat") else cell
        records.append(record)
    return sheet_title, header_index + 1, headers, records


def cells_of(row, number):
    cells = {}
    for index, cell in enumerate(row, start=1):
        if not is_blank(cell):
            cells[get_column_letter(index)] = cell.isoformat() if hasattr(cell, "isoformat") else cell
    return {"row": number, "cells": cells} if cells else None


def read_extra(xlsx_path, sheet_name, header_row, row_count):
    """Rows outside the plot table and the formulas of the first plot row, kept because notes and legends hold the rules."""
    workbook = load_workbook(str(xlsx_path), read_only=True, data_only=False)
    try:
        rows = [row for row in find_sheet(workbook, sheet_name).iter_rows(values_only=True)]
    finally:
        workbook.close()
    end = header_row + row_count
    above = [entry for entry in (cells_of(row, number) for number, row in enumerate(rows[:header_row - 1], start=1)) if entry]
    below = [entry for entry in (cells_of(row, number) for number, row in enumerate(rows[end:], start=end + 1)) if entry]
    formulas = {}
    if row_count and header_row < len(rows):
        headers = unique_headers(rows[header_row - 1])
        for header, cell in zip(headers, rows[header_row]):
            if isinstance(cell, str) and cell.startswith("="):
                formulas[header] = cell
    return {"above": above[:EXTRA_LIMIT], "below": below[:EXTRA_LIMIT], "formulas": formulas}


def write_outputs(xlsx_path, out_dir, sheet_name):
    sheet_title, header_row, headers, records = read_sheet(xlsx_path, sheet_name)
    out_dir = Path(out_dir)
    wide = wide_headers(headers)
    write_json(out_dir / "coords.json", {"source": Path(xlsx_path).name, "sheet": sheet_title, "headerRow": header_row, "rows": records}, wide)
    write_json(out_dir / "coords_headers.json", {"sheet": sheet_title, "headerRow": header_row, "headers": headers})
    extra = read_extra(xlsx_path, sheet_name, header_row, len(records))
    write_json(out_dir / "coords_extra.json", dict({"source": Path(xlsx_path).name, "sheet": sheet_title}, **extra))
    return headers, records


def main(argv=None):
    parser = argparse.ArgumentParser(description="Read the coordinates sheet of a workbook.")
    parser.add_argument("xlsx")
    parser.add_argument("out_dir")
    parser.add_argument("--sheet", default="Coordinates")
    args = parser.parse_args(argv)
    if not Path(args.xlsx).is_file():
        print("Workbook not found: " + args.xlsx, file=sys.stderr)
        return 2
    try:
        headers, records = write_outputs(args.xlsx, args.out_dir, args.sheet)
    except ValueError as problem:
        print(str(problem), file=sys.stderr)
        return 3
    print("headers: " + ", ".join(headers))
    print("rows " + str(len(records)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
