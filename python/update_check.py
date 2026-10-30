"""R402. The second route of Update from report: an edited copy of the level Excel and the report Excel of one NBH read with openpyxl,
the changes against the plan, their rule checks, the tiles and chips and the override entries worked out again with json and csv only,
apart from the C# of Common/Update, and held cell by cell against the result the C# route wrote.

python python/update_check.py --plan PLAN --excel FILE [--excel FILE] --overrides LIST --rules RULES --run-id ID [--tick all] --cs RESULT --out DIR

Counts, plot numbers and row numbers only to --out, outside the repo, never a level.
"""
import argparse
import csv
import hashlib
import json
import math
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
LEVELS, PLACED, SETTINGS = "Levels", "Placed", "Settings"
LEVEL_COLUMNS = ("Plot No", "Unit", "Cluster", "Road at front", "El @ Entry", "FFL", "SSL", "Height", "Steps", "Source")
PLACED_COLUMNS = ("Plot", "Kind", "Type", "Mirror", "Element id")
LINK = "link"
NO_CLUSTER = "none"
NO_BUILD = "none, the plan holds no run stamp"
BUILD_ID, PLAN_SHA = "Plan build id", "Plan SHA-256"
# The override field of each editable column, as R360 names them, unit for the villa type
FIELDS = {"FFL": "heightM", "El @ Entry": "entryM", "Steps": "steps", "Villa type": "unit", "Mirror": "mirrored"}
TILES = ("Changes found", "Follow the rules", "Break a manual rule", "Not editable ignored")
CHIPS = ("All", "FFL", "El @ Entry", "Steps", "Villa type", "Mirror", "Breaks a rule")
FOLLOWS = "follows the rules"
NO_RULE = "no rule checks this field"
# The plain name of each check by its id, the words of UpdateText.RuleNames of C# (12-ui-spec.md, at most 6 words)
PLAIN = {"fflOverEntry": "FFL over entry by the base", "stepsRule": "Steps by the steps rule", "stepsAtMost": "Front steps within their range",
         "stepHeights": "Height in the step legend", "riser": "Riser height within its limits", "oneStepUnits": "One step for one step units",
         "oneFflPerPiece": "One FFL per piece", "villaType": "Villa type of the profile", "handRule": "Mirror by the hand rule"}
# The source plan_from_cad gives a mirror the hand rules set
BY_RULE = "computed"
NOT_SAVABLE = "the plot has no override entry of the job %s, so Apply can not save it (Q169)"
TWICE = "%s and %s are both %s, Update from report takes one level Excel and one report Excel of one NBH (R402)"
# The decimals of the level Excel columns of R364
ENTRY_DEC, FFL_DEC = 4, 3


class Refused(Exception):
    pass


def fixed(value, decimals):
    """R385 again by its words: the text of 15 significant digits as a decimal, half away from zero, back to a float, no minus zero."""
    text = Decimal(format(value, ".15g"))
    return float(text.quantize(Decimal(1).scaleb(-decimals), rounding=ROUND_HALF_UP)) + 0.0


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def cell(row, column):
    return row[column] if row is not None and 0 <= column < len(row) else None


def key_text(value):
    if value is None or value == "":
        return None
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value)


def read_book(path):
    """The sheets Levels, Placed and Settings the file holds, every row from row 1 as openpyxl gives it, read only."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {name: [list(row) for row in book[name].iter_rows(values_only=True)] for name in (LEVELS, PLACED, SETTINGS) if name in book.sheetnames}
    finally:
        book.close()
    return {"file": Path(path).name, "sheets": sheets}


def build_id(plan):
    run = plan.get("run")
    found = run.get("runId") if isinstance(run, dict) else None
    return found if isinstance(found, str) else NO_BUILD


def refusal(book, plan, sha):
    settings = {}
    for row in book["sheets"].get(SETTINGS, []):
        if isinstance(cell(row, 0), str) and cell(row, 0) not in settings:
            settings[cell(row, 0)] = cell(row, 1)
    if SETTINGS not in book["sheets"] or not isinstance(settings.get(BUILD_ID), str) or not isinstance(settings.get(PLAN_SHA), str):
        return "%s holds no sheet Settings with the Plan build id and the Plan SHA-256, so its build is not known (R402)" % book["file"]
    if settings[BUILD_ID] != build_id(plan):
        return "%s is of another build, its Plan build id is %s, the plan of the model %s (R402)" % (book["file"], settings[BUILD_ID], build_id(plan))
    if settings[PLAN_SHA].lower() != sha.lower():
        return "%s is of another build, its Plan SHA-256 is %s, the plan of the model %s (R402)" % (book["file"], settings[PLAN_SHA], sha)
    return None


def header_row(rows, header):
    for index, row in enumerate(rows):
        if all(cell(row, column) == name for column, name in enumerate(header)):
            return index
    return None


def number_cell(value, model, decimals):
    """(differs, in model, in Excel, invalid) of a number cell against the plan at the decimals of its column, 0 a whole number."""
    shown = None if not is_number(model) else (int(fixed(model, 0)) if decimals == 0 else fixed(model, decimals))
    if value is None:
        return shown is not None, shown, None, None if shown is None else "the cell is empty"
    if not is_number(value):
        return True, shown, value, "the cell holds no number"
    if decimals == 0:
        if float(value) != math.floor(float(value)):
            return True, shown, value, "the cell holds no whole number"
        read = int(value)
    else:
        read = fixed(float(value), decimals)
    return shown is None or read != shown, shown, read, None


def changes_of(books, plan, sha, rules, profile, overrides=None):
    """The changes, the ignored cells and the lacks, or the refusal, of the edited books against the plan, a change of a plot with no entry
    of the job in the override list marked as one Apply can not save."""
    found = {"refusal": None, "changes": [], "ignored": [], "lacks": []}
    for sheet, kind in ((LEVELS, "level Excels"), (PLACED, "report Excels")):
        named = [book["file"] for book in books if sheet in book["sheets"]]
        if len(named) > 1:
            found["refusal"] = TWICE % (named[0], named[1], kind)
            return found
    for book in books:
        found["refusal"] = refusal(book, plan, sha)
        if found["refusal"] is None and LEVELS not in book["sheets"] and PLACED not in book["sheets"]:
            found["refusal"] = "%s holds no sheet Levels or Placed, so it is no level Excel or report Excel (R402)" % book["file"]
        if found["refusal"]:
            return found
    level_profile = (rules.get("levelProfiles") or {}).get(profile)
    if not isinstance(level_profile, dict):
        found["refusal"] = "the rules profile holds no level profile %s (R402)" % profile
        return found
    written = {item["plot"]: item for item in plan.get("plots") or []}
    not_written = {item["plot"] for item in plan.get("missing") or [] if item.get("written") is False}
    edits, order = {}, []

    def why_not(plot):
        return "the plan does not write the plot" if plot in not_written else "no plot of the plan"

    def ignore(book, sheet, index, plot, column, model, excel, why):
        found["ignored"].append({"file": book["file"], "sheet": sheet, "row": index + 1, "plot": plot, "column": column, "inModel": model, "inExcel": excel, "why": why})

    def change(book, sheet, index, plot, field, model, excel, invalid):
        made = {"file": book["file"], "sheet": sheet, "row": index + 1, "plot": plot, "field": field, "inModel": model, "inExcel": excel, "invalid": invalid,
                "notSavable": None, "checks": [], "breaks": [], "manualRules": []}
        found["changes"].append(made)
        return made

    for book in books:
        rows = book["sheets"].get(LEVELS)
        if rows is not None:
            top = header_row(rows, LEVEL_COLUMNS)
            if top is None:
                found.update(refusal="%s holds no header %s in its sheet Levels (R402)" % (book["file"], ", ".join(LEVEL_COLUMNS)), changes=[], ignored=[])
                return found
            col = {name: index for index, name in enumerate(LEVEL_COLUMNS)}
            seen = set()
            for index in range(top + 1, len(rows)):
                row = rows[index]
                plot = key_text(cell(row, 0))
                if plot is None:
                    break
                if plot in seen or plot not in written:
                    why = "a second row of the plot" if plot in edits else why_not(plot)
                    seen.add(plot)
                    for name in LEVEL_COLUMNS:
                        if name not in ("Plot No", "Cluster", "Source") and cell(row, col[name]) is not None:
                            ignore(book, LEVELS, index, plot, name, None, cell(row, col[name]), why)
                    continue
                seen.add(plot)
                planned = written[plot]
                levels = planned.get("levels") or {}
                edit = {"cluster": cell(row, col["Cluster"]) if isinstance(cell(row, col["Cluster"]), str) else None, "entry": None, "ffl": None, "steps": None, "changes": []}
                edits[plot] = edit
                order.append(plot)
                unit = (planned.get("unit") or {}).get("value")
                if cell(row, col["Unit"]) != unit:
                    ignore(book, LEVELS, index, plot, "Unit", unit, cell(row, col["Unit"]), "not editable")
                for name in LEVEL_COLUMNS:
                    value = cell(row, col[name])
                    if name in ("Road at front", "SSL", "Height"):
                        key = {"Road at front": "roadAtFrontM", "SSL": "ssl", "Height": "height"}[name]
                        differs, model, _, _ = number_cell(value, (levels.get(key) or {}).get("value"), FFL_DEC if name == "SSL" else ENTRY_DEC)
                        if differs:
                            ignore(book, LEVELS, index, plot, name, model, value, "not editable")
                    elif name in ("El @ Entry", "FFL", "Steps"):
                        key, decimals = {"El @ Entry": ("entry", ENTRY_DEC), "FFL": ("ffl", FFL_DEC), "Steps": ("steps", 0)}[name]
                        differs, model, excel, invalid = number_cell(value, (levels.get(key) or {}).get("value"), decimals)
                        if differs:
                            edit["changes"].append(change(book, LEVELS, index, plot, name, model, excel, invalid))
                            edit[key] = excel if invalid is None else None
        rows = book["sheets"].get(PLACED)
        if rows is not None:
            top = header_row(rows, PLACED_COLUMNS)
            if top is None:
                found.update(refusal="%s holds no header %s in its sheet Placed (R402)" % (book["file"], ", ".join(PLACED_COLUMNS)), changes=[], ignored=[])
                return found
            seen = set()
            for index in range(top + 1, len(rows)):
                row = rows[index]
                plot = key_text(cell(row, 0))
                if plot is None:
                    break
                kind = cell(row, 1) if isinstance(cell(row, 1), str) else None
                first = (plot, kind) not in seen
                seen.add((plot, kind))
                planned = written.get(plot) if first else None
                floor = ((planned or {}).get("floors") or {}).get(kind) if kind != LINK else None
                if planned is None or (kind != LINK and not isinstance(floor, dict)):
                    why = "a second row of the plot and kind" if not first else why_not(plot) if plot not in written else "no kind of the plot in the plan"
                    ignore(book, PLACED, index, plot, "Kind", None, cell(row, 1), why)
                    continue
                if cell(row, 4) is not None:
                    ignore(book, PLACED, index, plot, "Element id", None, cell(row, 4), "not editable")
                if floor is not None:
                    model = (floor.get("type") or {}).get("value")
                    model = model if isinstance(model, str) else None
                    if cell(row, 2) != model:
                        ignore(book, PLACED, index, plot, "Type", model, cell(row, 2), "the type of a floor is not editable")
                    if cell(row, 3) is not None:
                        ignore(book, PLACED, index, plot, "Mirror", None, cell(row, 3), "a floor holds no mirror flag")
                    continue
                unit = (planned.get("unit") or {}).get("value")
                typed = cell(row, 2)
                if typed != unit:
                    change(book, PLACED, index, plot, "Villa type", unit, typed, None if isinstance(typed, str) and typed.strip() else "the cell holds no villa type")
                flag = (planned.get("mirrored") or {}).get("value")
                flag = flag if isinstance(flag, bool) else None
                read = cell(row, 3)
                if not (read is flag or (isinstance(read, bool) and read == flag)):
                    change(book, PLACED, index, plot, "Mirror", flag, read, None if isinstance(read, bool) else "the cell is empty" if read is None else "the cell holds no TRUE or FALSE")
    # The rebuild keeps the cell of the model for a change Apply can not save, so the checks take it as no edit
    job = level_profile.get("job")
    if overrides is not None and isinstance(job, str):
        held = {entry.get("plot") for entry in overrides.get("overrides") or [] if entry.get("source") == job}
        for item in found["changes"]:
            if item["invalid"] is None and item["plot"] not in held:
                item["notSavable"] = NOT_SAVABLE % job
                key = {"El @ Entry": "entry", "FFL": "ffl", "Steps": "steps"}.get(item["field"])
                if key and item["plot"] in edits and any(change is item for change in edits[item["plot"]]["changes"]):
                    edits[item["plot"]][key] = None
    level_checks(found, edits, order, written, rules, level_profile)
    type_checks(found, rules)
    mirror_checks(found, written)
    manual_checks(found, rules)
    for item in found["changes"]:
        item["ticked"] = can_apply(item) and not item["manualRules"]
        named = [PLAIN.get(check, check) for check in item["breaks"]] + item["manualRules"]
        item["ruleCheck"] = item["invalid"] or item["notSavable"] or (", ".join(named) if item["breaks"] else FOLLOWS if item["checks"] else NO_RULE)
    return found


def can_apply(item):
    return item["invalid"] is None and item.get("notSavable") is None


def lack(found, key, check):
    line = "the rules profile holds no %s, so the check %s is not run (R402)" % (key, PLAIN.get(check, check))
    if line not in found["lacks"]:
        found["lacks"].append(line)


def base_of(rules, unit, base_rule):
    """The base step of a unit in metres: the typology table of keyRules.baseHeight, else steps.base_height in mm."""
    table = base_rule.get("table") if isinstance(base_rule, dict) and base_rule.get("rule") == "typology" else None
    if isinstance(table, dict) and table and unit in table:
        return float(table[unit])
    bases = rules["steps"]["base_height"]
    return bases[unit if unit in bases else "default"] / 1000.0


def steps_of(rule, height, unit, rules, level_profile, base_rule):
    if rule == "stepsFromHeight":
        over = fixed((fixed(height, 6) - base_of(rules, unit, base_rule)) / (rules["steps"]["riser_added"] / 1000.0), 6)
        return 1 + max(0, math.ceil(over - 0.5))
    if rule == "stepBands":
        for band in level_profile.get("stepBands") or []:
            if fixed(height, 6) < band[0]:
                return band[1]
        return None
    raise Refused("keyRules.steps codes %s, a build gives stepsFromHeight or stepBands (R402)" % rule)


def in_legend(legend, height, count):
    heights = [pair[0] for pair in legend if isinstance(pair, list) and len(pair) == 2 and is_number(pair[0]) and is_number(pair[1]) and int(fixed(pair[1], 0)) == count]
    return bool(heights) and min(heights) <= height <= max(heights)


def level_checks(found, edits, order, written, rules, level_profile):
    """A check the plan row held and the edited row breaks is broken by the change, named on each valid level change of the plot."""
    steps_block = rules.get("steps") if isinstance(rules.get("steps"), dict) else None
    key_rules = level_profile.get("keyRules") or {}
    rule = (key_rules.get("steps") or {}).get("rule")
    base_rule = key_rules.get("baseHeight")
    after = {}

    def number(key, check):
        value = (steps_block or {}).get(key)
        if is_number(value):
            return float(value)
        lack(found, "steps." + key, check)
        return None

    def listed(key, check):
        value = (steps_block or {}).get(key)
        if isinstance(value, list):
            return value
        lack(found, "steps." + key, check)
        return None

    for plot in order:
        edit, planned = edits[plot], written[plot]
        levels = planned.get("levels") or {}
        entry0, ffl0, steps0 = ((levels.get(key) or {}).get("value") for key in ("entry", "ffl", "steps"))
        if not is_number(ffl0):
            continue
        after[plot] = edit["ffl"] if edit["ffl"] is not None else ffl0
        valid = [item for item in edit["changes"] if can_apply(item)]
        if not valid or not is_number(entry0) or not is_number(steps0):
            continue
        unit = (planned.get("unit") or {}).get("value")
        entry1 = edit["entry"] if edit["entry"] is not None else entry0
        ffl1 = edit["ffl"] if edit["ffl"] is not None else ffl0
        height0, height1 = fixed(ffl0 - entry0, ENTRY_DEC), fixed(ffl1 - entry1, ENTRY_DEC)
        base = base_of(rules, unit, base_rule) if isinstance((steps_block or {}).get("base_height"), dict) else None

        def by_rule(height):
            if rule is None:
                lack(found, "keyRules.steps.rule", "stepsRule")
                return None
            return steps_of(rule, height, unit, rules, level_profile, base_rule)

        # plan_from_book.py sets the count of a steps override after the steps rule, so a plot that holds one keeps it in a rebuild
        held = "steps" in ((planned.get("overrides") or {}).get("applied") or [])
        moved = edit["entry"] is not None or edit["ffl"] is not None
        steps1 = edit["steps"] if edit["steps"] is not None else (by_rule(height1) if moved and not held else steps0)
        if steps1 is None:
            continue
        broken, ran = [], []

        def hold(name, before, now):
            if before is not None and now is not None:
                ran.append(name)
            if before is True and now is False:
                broken.append(name)

        if base is not None:
            slack = 0.5 * 10 ** -FFL_DEC
            hold("fflOverEntry", height0 >= base - slack, height1 >= base - slack)
        else:
            lack(found, "steps.base_height", "fflOverEntry")
        if edit["steps"] is not None:
            rule0, rule1 = by_rule(height0), by_rule(height1)
            hold("stepsRule", None if rule0 is None else rule0 == steps0, None if rule1 is None else rule1 == steps1)
        low, high = number("front_min", "stepsAtMost"), number("front_max", "stepsAtMost")
        if low is not None and high is not None:
            hold("stepsAtMost", low <= steps0 <= high, low <= steps1 <= high)
        legend = listed("legend", "stepHeights")
        if legend is not None:
            hold("stepHeights", in_legend(legend, height0, steps0), in_legend(legend, height1, steps1))
        least, most = number("riser_min", "riser"), number("riser_max", "riser")
        if least is not None and most is not None and base is not None:
            def riser(height, count):
                return count < 2 or least <= fixed((height - base) * 1000 / (count - 1), 6) <= most
            hold("riser", riser(height0, steps0), riser(height1, steps1))
        single = listed("one_step_only", "oneStepUnits")
        if single is not None and unit in single:
            hold("oneStepUnits", steps0 == 1, steps1 == 1)
        for item in valid:
            item["checks"].extend(ran)
            item["breaks"].extend(broken)
    pieces = {}
    for plot in order:
        ffl_rule = (written[plot].get("levelRule") or {}).get("ffl")
        if plot in after and edits[plot]["cluster"] not in (None, NO_CLUSTER) and is_number(ffl_rule):
            pieces.setdefault((edits[plot]["cluster"], float(ffl_rule)), []).append(plot)
    for members in pieces.values():
        moved = [item for plot in members for item in edits[plot]["changes"] if item["field"] == "FFL" and can_apply(item)]
        if len(members) < 2 or not moved:
            continue
        broken = len({written[plot]["levels"]["ffl"]["value"] for plot in members}) == 1 and len({after[plot] for plot in members}) != 1
        for item in moved:
            item["checks"].append("oneFflPerPiece")
            if broken:
                item["breaks"].append("oneFflPerPiece")


def type_checks(found, rules):
    types = [item for item in found["changes"] if item["field"] == "Villa type" and can_apply(item)]
    if not types:
        return
    typologies = rules.get("typologies")
    if not isinstance(typologies, dict):
        found["lacks"].append("the rules profile holds no typologies, so the check %s is not run (R402)" % PLAIN["villaType"])
        return
    for item in types:
        item["checks"].append("villaType")
        if item["inExcel"] not in typologies:
            item["breaks"].append("villaType")


def mirror_checks(found, written):
    """A mirror of source computed is the hand the rules of the profile gave, so an edit of it breaks them, an override mirror is not checked."""
    for item in found["changes"]:
        if item["field"] != "Mirror" or not can_apply(item):
            continue
        mirrored = written[item["plot"]].get("mirrored") or {}
        if mirrored.get("source") == BY_RULE and isinstance(mirrored.get("value"), bool):
            item["checks"].append("handRule")
            if item["inExcel"] is not mirrored["value"]:
                item["breaks"].append("handRule")


def manual_checks(found, rules):
    block = rules.get("manualRules")
    if isinstance(block, list):
        entries = [(item.get("id"), item) for item in block if isinstance(item, dict)]
    elif isinstance(block, dict):
        entries = [(item.get("id") if isinstance(item.get("id"), str) else name, item) for name, item in block.items() if isinstance(item, dict)]
    else:
        found["lacks"].append("the rules profile holds no block manualRules, so no change breaks a manual rule (R402)")
        return
    by_check, bare = {}, []
    for ident, item in entries:
        if not isinstance(ident, str):
            continue
        checks = item.get("updateChecks")
        if not isinstance(checks, list) or not checks:
            bare.append(ident)
            continue
        for check in checks:
            if isinstance(check, str):
                by_check.setdefault(check, []).append(ident)
    if bare:
        found["lacks"].append("%d manual rules of the block name no updateChecks, so no change is held against them: %s (R402)" % (len(bare), ", ".join(bare)))
    for item in found["changes"]:
        for check in item["breaks"]:
            for ident in by_check.get(check, []):
                if ident not in item["manualRules"]:
                    item["manualRules"].append(ident)


def tiles_of(found):
    changes = found["changes"]
    breaks = [item for item in changes if item["invalid"] is not None or item["breaks"]]
    follow = sum(1 for item in changes if can_apply(item) and not item["breaks"] and item["checks"])
    tiles = dict(zip(TILES, (len(changes), follow, sum(1 for item in changes if item["manualRules"]), len(found["ignored"]))))
    chips = {"All": len(changes), "Breaks a rule": len(breaks)}
    chips.update({field: sum(1 for item in changes if item["field"] == field) for field in FIELDS})
    return tiles, {chip: chips[chip] for chip in CHIPS}


def apply(ticked, plan, overrides, rules, profile, run_id):
    """Each ticked change as an entry of R330 in the form of R360, replacing an entry of its plot, field and source in place, else after the
    last entry of its plot, the key and model taken from an entry of that source of the plot."""
    job = ((rules.get("levelProfiles") or {}).get(profile) or {}).get("job")
    if not isinstance(job, str):
        raise Refused("the level profile %s names no job, so an override has no source (R402)" % profile)
    listing = json.loads(json.dumps(overrides))
    entries = listing["overrides"]
    plots = {item["plot"]: item for item in plan.get("plots") or []}
    done = {"entries": [], "replaced": [], "plots": [], "notApplied": []}
    for item in ticked:
        where = "plot %s %s: " % (item["plot"], item["field"])
        field = FIELDS.get(item["field"])
        if not can_apply(item) or field is None:
            done["notApplied"].append(where + (item["invalid"] or item.get("notSavable") or "no field of the override list") + " (R402)")
            continue
        if item["field"] in ("El @ Entry", "FFL"):
            rule = (plots.get(item["plot"], {}).get("levelRule") or {}).get("entry" if item["field"] == "El @ Entry" else "ffl")
            if not is_number(rule):
                done["notApplied"].append(where + "the plan holds no levelRule of the plot, so no gap to the rule (R402)")
                continue
            value = {"gapToRuleM": fixed(item["inExcel"] - rule, ENTRY_DEC)}
        else:
            value = item["inExcel"]
        mine = [entry for entry in entries if entry.get("plot") == item["plot"]]
        same = [entry for entry in mine if entry.get("source") == job and entry.get("field") == field]
        origin = same[0] if same else next((entry for entry in mine if entry.get("source") == job), None)
        if origin is None:
            done["notApplied"].append(where + "no entry of source %s of the plot names its key and model (R402)" % job)
            continue
        made = {"plot": item["plot"], "field": field, "value": value, "rule": None, "source": job, "model": origin.get("model"), "key": origin.get("key"),
                "decision": "Update from report", "reason": "Update from report " + run_id, "byHand": True}
        if same:
            place = next(index for index, entry in enumerate(entries) if entry is same[0])
            for old in same:
                if not any(old is entry for entry in done["entries"]):
                    done["replaced"].append(old)
                done["entries"] = [entry for entry in done["entries"] if entry is not old]
            entries[:] = [entry for entry in entries if not any(entry is old for old in same)]
        else:
            place = next(index for index, entry in enumerate(entries) if entry is mine[-1]) + 1
        entries.insert(place, made)
        done["entries"].append(made)
        if item["plot"] not in done["plots"]:
            done["plots"].append(item["plot"])
    done["list"] = listing
    return done


def list_text(listing):
    """The list as majority.overrides_text writes it, by its own words: one entry per line with json.dumps."""
    def block(items):
        return ",\n".join("    " + json.dumps(item, ensure_ascii=False) for item in items) + ("\n" if items else "")
    text = "{\n  \"note\": " + json.dumps(listing.get("note"), ensure_ascii=False) + ",\n  \"overrides\": [\n" + block(listing["overrides"]) + "  ]"
    for key in ("cannot", "skip"):
        if key in listing:
            text += ",\n  \"" + key + "\": [\n" + block(listing[key]) + "  ]"
    return text + "\n}\n"


def result(books, plan, sha, rules, profile, overrides, run_id, tick="default"):
    found = changes_of(books, plan, sha, rules, profile, overrides)
    tiles, chips = tiles_of(found)
    found.update(tiles=tiles, chips=chips)
    if found["refusal"] is None and overrides is not None:
        done = apply([item for item in found["changes"] if tick == "all" or item["ticked"]], plan, overrides, rules, profile, run_id)
        found.update(entries=done["entries"], replaced=done["replaced"], plots=done["plots"], notApplied=done["notApplied"], listText=list_text(done["list"]))
    return found


def same_value(one, other):
    if isinstance(one, bool) or isinstance(other, bool):
        return one is other or (isinstance(one, bool) and isinstance(other, bool) and one == other)
    if is_number(one) and is_number(other):
        return one == other
    return one == other


def compare(python, sharp):
    """Matches and mismatches cell by cell: each field of each change and ignored cell by place, the tiles, chips, lacks, plots and not
    applied lines, each key of each entry, and the list text line by line."""
    tally = {"matches": 0, "mismatches": 0, "conflicts": []}

    def held(where, one, other):
        if same_value(one, other):
            tally["matches"] += 1
        else:
            tally["mismatches"] += 1
            tally["conflicts"].append(where)

    held("refusal", python.get("refusal"), sharp.get("refusal"))
    for name, keys in (("changes", ("file", "sheet", "row", "plot", "field", "inModel", "inExcel", "invalid", "notSavable", "checks", "breaks", "manualRules", "ruleCheck", "ticked")),
                       ("ignored", ("file", "sheet", "row", "plot", "column", "inModel", "inExcel", "why")),
                       ("entries", ("plot", "field", "value", "rule", "source", "model", "key", "decision", "reason", "byHand")),
                       ("replaced", ("plot", "field", "value", "source", "key"))):
        mine, theirs = python.get(name) or [], sharp.get(name) or []
        held(name + " count", len(mine), len(theirs))
        for index, (one, other) in enumerate(zip(mine, theirs)):
            for key in keys:
                held("%s %d %s plot %s row %s" % (name, index, key, one.get("plot"), one.get("row")), one.get(key), other.get(key))
    for name in ("tiles", "chips"):
        for key in sorted(set(python.get(name) or {}) | set(sharp.get(name) or {})):
            held(name + " " + key, (python.get(name) or {}).get(key), (sharp.get(name) or {}).get(key))
    for name in ("lacks", "plots", "notApplied"):
        held(name, python.get(name) or [], sharp.get(name) or [])
    if "listText" in python or "listText" in sharp:
        mine, theirs = (python.get("listText") or "").split("\n"), (sharp.get("listText") or "").split("\n")
        held("listText lines", len(mine), len(theirs))
        for index, (one, other) in enumerate(zip(mine, theirs)):
            held("listText line %d" % (index + 1), one, other)
    return tally


def outside_repo(path):
    resolved = Path(path).resolve()
    if resolved == REPO or REPO in resolved.parents:
        raise Refused("--out %s is inside the repo, the counts go outside it (R402)" % path)
    return resolved


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--plan", required=True)
    parser.add_argument("--excel", action="append", required=True, help="an edited copy of the level Excel or the report Excel, read only")
    parser.add_argument("--overrides", required=True, help="the override list the changes apply to, read only")
    parser.add_argument("--rules", required=True)
    parser.add_argument("--profile", default="parsons2024Rev6")
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--tick", choices=("default", "all"), default="default")
    parser.add_argument("--cs", help="the result JSON the C# route wrote of the same inputs")
    parser.add_argument("--out", required=True, help="a folder outside the repo for the counts")
    args = parser.parse_args(argv)
    out = outside_repo(args.out)
    plan_bytes = Path(args.plan).read_bytes()
    plan = json.loads(plan_bytes)
    rules = json.loads(Path(args.rules).read_text(encoding="utf-8"))
    overrides = json.loads(Path(args.overrides).read_text(encoding="utf-8"))
    found = result([read_book(path) for path in args.excel], plan, hashlib.sha256(plan_bytes).hexdigest(), rules, args.profile, overrides, args.run_id, args.tick)
    report = {"refusal": found["refusal"], "tiles": found["tiles"], "chips": found["chips"], "lacks": found["lacks"], "ignored": len(found["ignored"]),
              "entries": len(found.get("entries") or []), "plots": len(found.get("plots") or []), "notApplied": len(found.get("notApplied") or [])}
    if args.cs:
        report["compare"] = compare(found, json.loads(Path(args.cs).read_text(encoding="utf-8")))
    out.mkdir(parents=True, exist_ok=True)
    (out / "update_check.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    with (out / "update_changes_py.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(("plot", "field", "sheet", "row", "ruleCheck", "ticked"))
        for item in found["changes"]:
            writer.writerow((item["plot"], item["field"], item["sheet"], item["row"], item["ruleCheck"], item["ticked"]))
    print(json.dumps(report.get("compare", {"matches": None}) | {"tiles": found["tiles"]}))
    return 1 if found["refusal"] or report.get("compare", {}).get("mismatches") else 0


if __name__ == "__main__":
    sys.exit(main())
