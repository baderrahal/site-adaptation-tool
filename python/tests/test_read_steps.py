"""R348 to R350. The Python side of the read steps of the parity harness, python/read_steps.py, on made up CadReader outputs, LandXML
files and workbooks built in code (R10). dotnet/SiteAdaptation.Tests/ReadStepTests.cs holds the C# side on the same made up data."""
import json

import openpyxl
import pytest

import cad_analysis
import parity
import plan_from_book
import project_profile
import read_steps
from landxml_fixture import GROUND, GROUND_FACES, ROAD, ROAD_FACES, landxml_text

FIELDS = {"rings": {"kind": "points", "decimals": 3}, "types": {"kind": "noGeometry"}}


def shape(label, *xy, kind="plot_shape", type_name="DP1", area=100.0):
    return {"kind": kind, "plot": label, "vertices": [[xy[i], xy[i + 1]] for i in range(0, len(xy), 2)], "type": type_name, "area": area}


# The made up baseplan of the C# tests: label 1 drawn in two rings, 02 and 3a, a polyline and a shape of two vertices
PLOTS = [shape("1", 1000, 5000, 1002, 5000, 1002, 5003, 1000, 5003), shape("02", 0, 0, 10, 0, 10, 10, type_name="OS", area=50.0),
         shape("1", 1500, 4500, 1501, 4500, 1501, 4501, area=0.5), shape("3a", 0, 30, 10, 30, 10, 40, type_name=None, area=None),
         shape("5", 0, 60, 10, 60, kind="polyline"), shape("4", 0, 50, 10, 50, area=0.0)]


def test_r348_the_labels_step_holds_every_plot_shape_of_a_label_as_read():
    found = read_steps.label_step(PLOTS)
    assert list(found) == ["1", "02", "3a", "4"]
    assert found["1"] == {"rings": [[[1000, 5000], [1002, 5000], [1002, 5003], [1000, 5003]], [[1500, 4500], [1501, 4500], [1501, 4501]]],
                          "types": ["DP1", "DP1"], "areas": [100.0, 0.5], "whole": True}
    assert found["02"]["whole"] is True and found["3a"]["whole"] is False
    assert found["3a"]["types"] == [None] and found["4"]["rings"] == [[[0, 50], [10, 50]]]


BLOCKS = [{"name": "Villa_C10", "layer": "MP_VL_C10", "handle": "A1", "insertPoint": [1.0, 2.0], "scale": [-1, 1, 1], "rotationDeg": 30.5, "mirrored": True,
           "outline": [[0, 0], [4, 0], [4, 3]]},
          {"name": "JM", "layer": "MP_BLDG_Footprint", "handle": "A1", "insertPoint": [5.0, 6.0], "scale": [1, 1, 1], "rotationDeg": 0.0, "mirrored": False},
          {"name": "DP2", "layer": "MP_DP2", "handle": "B7", "insertPoint": [7.25, 8.5], "scale": [1, 1, 1], "rotationDeg": 359.999, "mirrored": False,
           "outline": [[1, 1], [2, 1], [2, 2], [1, 2]]}]


def test_r348_the_blocks_step_holds_every_block_by_its_handle_a_handle_met_again_by_its_count():
    found = read_steps.block_step(BLOCKS)
    assert list(found) == ["A1", "A1#2", "B7"]
    assert found["A1"] == {"name": "Villa_C10", "layer": "MP_VL_C10", "mirrored": True, "scale": [-1, 1, 1], "insertPoint": [1.0, 2.0], "rotationDeg": 30.5,
                           "outline": [[0, 0], [4, 0], [4, 3]]}
    assert found["A1#2"]["outline"] is None


PIECES = [{"kind": "polyline", "layer": "RD-CURB-LINE", "handle": "C1", "pts": [[0, 0, 0], [10, 0, 0], [10, 5, 0]], "closed": True},
          {"kind": "line", "layer": "Ext Road$1$rd-curb-line", "handle": "C2", "pts": [[20, 0, 0], [30, 0, 0]]},
          {"kind": "polyline", "layer": "ROW", "handle": "R1", "pts": [[0, 10, 0], [50, 10, 0]], "closed": False},
          {"kind": "polyline", "layer": "RD-CURB-LINE", "handle": "C3", "pts": [[5, 5, 0]], "closed": False},
          {"kind": "arc", "layer": "RD-CURB-LINE", "handle": "C4", "center": [0, 0], "radius": 1},
          {"kind": "polyline", "layer": "RD-CURB-LINE-OLD", "handle": "C5", "pts": [[0, 0, 0], [1, 1, 0]], "closed": False}]


def test_r348_the_edge_steps_take_the_lines_read_road_edges_takes_by_their_place_in_the_read(tmp_path):
    curbs = read_steps.edge_step(PIECES, ["RD-CURB-LINE"])
    assert list(curbs) == ["0", "1"]
    assert curbs["0"] == {"handle": "C1", "layer": "RD-CURB-LINE", "kind": "polyline", "closed": True, "points": [[0, 0], [10, 0], [10, 5], [0, 0]]}
    assert curbs["1"] == {"handle": "C2", "layer": "Ext Road$1$rd-curb-line", "kind": "line", "closed": None, "points": [[20, 0], [30, 0]]}
    # The lines of the build are the same ones in the same order
    (tmp_path / "pieces.json").write_text(json.dumps({"pieces": PIECES}), encoding="utf-8")
    lines = cad_analysis.read_road_edges(tmp_path, ["RD-CURB-LINE"])
    assert [[list(point) for point in line.coords] for line in lines] == [item["points"] for item in curbs.values()]
    rows = read_steps.edge_step(PIECES, ["ROW"])
    assert list(rows) == ["0"] and rows["0"]["handle"] == "R1"


def write_roads(folder):
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "A__Road.xml").write_text(landxml_text([("Road", ROAD, ROAD_FACES)]), encoding="utf-8")
    (folder / "B__Ground.xml").write_text(landxml_text([("Ground", GROUND, GROUND_FACES)]), encoding="utf-8")
    return [folder / "A__Road.xml", folder / "B__Ground.xml"]


def test_r349_the_surfaces_step_holds_the_unit_counts_bounds_and_levels_of_each_surface(tmp_path):
    files = write_roads(tmp_path / "landxml")
    plan_from_book.load_roads(files, tmp_path / "scratch")
    reports = [json.loads((tmp_path / "scratch" / path.stem / "landxml_summary.json").read_text(encoding="utf-8")) for path in files]
    found = read_steps.surface_step(reports)
    assert list(found) == ["Road", "Ground"]
    assert found["Road"] == {"file": "A__Road.xml", "linearUnit": "meter", "epsgCode": "32637", "pointCount": 6, "faceCount": 5, "visibleFaceCount": 3,
                             "invisibleFaceCount": 1, "badFaceCount": 1, "badPointCount": 0, "declaredZRange": {"zMin": "1", "zMax": "40"},
                             "bboxMin": [1000.0, 5000.0], "bboxMax": [1020.0, 5010.0], "zMin": 10.0, "zMax": 40.0}


def test_r349_the_road_levels_step_gives_each_face_vertex_the_level_of_the_first_surface_that_covers_it(tmp_path):
    surfaces = plan_from_book.load_roads(write_roads(tmp_path / "landxml"), tmp_path / "scratch")
    faces = read_steps.face_rings(PLOTS, {"1": "101", "02": "102", "4": "104"})
    # 02 is the whole label 2, which plots.csv does not name, and 4 has two vertices
    assert list(faces) == ["101"]
    found = read_steps.road_level_step(["101", "103"], faces, surfaces, 0.005)
    assert list(found) == ["101/%d" % index for index in range(7)]
    assert found["101/0"] == {"point": [1000, 5000], "level": pytest.approx(10.0), "surface": "Road"}
    assert found["101/2"]["level"] == pytest.approx(15.0) and found["101/2"]["surface"] == "Road"
    # The second ring lies on the ground surface alone
    assert found["101/4"] == {"point": [1500, 4500], "level": pytest.approx(5.0), "surface": "Ground"}


def test_r350_the_area_key_step_gives_the_zone_and_neighbourhood_of_each_plot_of_the_dlus_and_the_sales_sheet():
    dlus = [{"plot": 101, "zone": "Zone 2", "nbh": "012"}, {"plot": "102", "zone": 3, "nbh": "01a"}, {"plot": 103, "zone": "-", "nbh": None},
            {"plot": 101.0, "zone": "Zone 2", "nbh": "013"}]
    sales = [{"plot": 102, "zone": 3.0, "nbh": "01A"}, {"plot": 104, "zone": 4, "nbh": "020"}]
    found = read_steps.area_key_step(dlus, sales)
    assert list(found) == ["101", "102", "104"]
    assert found["101"] == {"inDlus": True, "zone": "Zone 2", "nbhId": "013", "twoRows": True, "inSales": False, "salesZone": None, "salesNbhId": None, "salesTwoRows": False}
    assert found["102"] == {"inDlus": True, "zone": "Zone 3", "nbhId": "01A", "twoRows": False, "inSales": True, "salesZone": "Zone 3", "salesNbhId": "01A", "salesTwoRows": False}
    assert found["104"]["inDlus"] is False and found["104"]["salesNbhId"] == "020"


def test_r350_the_setback_and_sales_steps_hold_the_rows_as_the_python_readers_give_them():
    setbacks = read_steps.dlus_setback_step([{"plot": 101, "typology": "DP2", "front": 4.5, "rear": "-"}, {"plot": 101.0, "typology": "DP2", "front": 3}])
    assert setbacks == {"101": {"plotCell": 101, "typology": "DP2", "front": 4.5, "rear": "-"}, "101#2": {"plotCell": 101.0, "typology": "DP2", "front": 3}}
    rows = {"101": {"Plot No": 101, "Orientation": "EN", None: 7}}
    found = read_steps.sales_step(rows, [{"plot": 101, "typology": "DP2", "landWidth": 12.5}, {"plot": "x", "typology": "VL1"}])
    assert found == {"101": {"row": {"Plot No": 101, "Orientation": "EN", None: 7}, "plain": {"plot": 101, "typology": "DP2", "landWidth": 12.5}}}
    # A header cell with no value is the key null, as json writes a key None
    assert read_steps.entry(found["101"], {"row": {"kind": "noGeometry"}})["raw"]["row"]["null"] == 7


def made_settings(tmp_path):
    """A copy of the config with cadReads in tmp_path, the made up reads there with their layers.json in Meters, and two road surfaces."""
    values = json.loads(json.dumps(project_profile.load_settings().values))
    values["cadReads"] = dict(values["cadReads"], root=str(tmp_path / "reads"))
    values["siteBuild"]["roadSurfaces"] = ["A__Road", "B__Ground"]
    drawings = values["drawings"]
    for key, name, body in (("baseplan", "plots.json", {"source": "b.dwg", "plots": PLOTS}), ("villaBlocks", "blocks.json", {"source": "f.dwg", "blocks": BLOCKS}),
                            ("villaBlocks", "pieces.json", {"source": "f.dwg", "pieces": [dict(PIECES[0], layer="MP_Setback Line")]}),
                            ("roadLayout", "pieces.json", {"source": "r.dwg", "pieces": PIECES})):
        folder = tmp_path / "reads" / drawings[key]
        folder.mkdir(parents=True, exist_ok=True)
        (folder / "layers.json").write_text(json.dumps({"units": "Meters", "layers": []}), encoding="utf-8")
        (folder / name).write_text(json.dumps(body), encoding="utf-8")
    path = tmp_path / "project.json"
    path.write_text(json.dumps(values), encoding="utf-8")
    return path, values


def made_workbooks(tmp_path, values):
    """The DLUS and the sales workbook with the sheets of dataBook in config, a header row on the row each names."""
    book = values["dataBook"]
    dlus, sales = openpyxl.Workbook(), openpyxl.Workbook()
    area = book["data5"]["areaKey"]["dlus"]
    sheet = dlus.active
    sheet.title = area["sheet"]
    sheet.append([area["columns"]["plot"], area["columns"]["zone"], area["columns"]["nbh"]])
    sheet.append([101, "Zone 2", "012"])
    sheet.append(["102", 3, "01a"])
    sales_sheet = sales.active
    sales_sheet.title = book["locationPlan"]["sales"]["sheet"]
    columns = book["rev9"]["sales"]["columns"]
    sales_sheet.append([columns["plot"], columns["typology"], columns["typologyOld"], columns["nbh"], columns["zone"], columns["landWidth"], columns["landLength"], "Orientation"])
    sales_sheet.append([101, "DP2", "D02", "012", 2, 12.5, 25, "EN"])
    setback = sales.create_sheet(book["rev9"]["dlus"]["sheet"])
    setback.append(["title"])
    setback.append([])
    names = book["rev9"]["dlus"]["columns"]
    setback.append([names[key] for key in ("plot", "typology", "nbh", "front", "right", "left", "rear")])
    setback.append([101, "DP2", "012", 4.5, 0, "1.5", None])
    paths = tmp_path / "dlus.xlsx", tmp_path / "sales.xlsx"
    dlus.save(paths[0])
    sales.save(paths[1])
    return paths


def test_r348_r349_r350_read_steps_writes_every_step_file_with_its_inputs(tmp_path, capsys):
    settings, values = made_settings(tmp_path)
    write_roads(tmp_path / "landxml")
    dlus, sales = made_workbooks(tmp_path, values)
    (tmp_path / "plots.csv").write_text("baseplanLabel,locationPlanNumber\n1,101\n", encoding="utf-8")
    (tmp_path / "labels.csv").write_text("number,code,kind,nbh\n101,DP1,clientCode,012\n", encoding="utf-8")
    steps = tmp_path / "steps"
    code = read_steps.main(["--settings", str(settings), "--steps", str(steps), "--landxml", str(tmp_path / "landxml"), "--scratch", str(tmp_path / "scratch"),
                            "--plots-csv", str(tmp_path / "plots.csv"), "--labels", str(tmp_path / "labels.csv"), "--dlus", str(dlus), "--sales", str(sales)])
    assert code == 0, capsys.readouterr().err
    counts = {}
    for step in read_steps.STEPS:
        payload = parity.load_step(steps / parity.step_name(step, "python"), "python")
        assert payload["scope"] == "project"
        counts[step] = len(payload["plots"])
    assert counts == {"labels": 4, "blocks": 3, "curbs": 2, "rows": 1, "setbacks": 1, "surfaces": 2, "roadLevels": 7, "areaKey": 2, "dlusSetbacks": 1, "sales": 1}
    sale = parity.load_step(steps / parity.step_name("sales", "python"), "python")["plots"]["101"]["raw"]
    assert sale["row"]["Orientation"] == "EN" and sale["plain"]["landWidth"] == 12.5
    setback = parity.load_step(steps / parity.step_name("dlusSetbacks", "python"), "python")["plots"]["101"]["raw"]
    assert setback == {"plotCell": 101, "typology": "DP2", "nbh": "012", "front": 4.5, "right": 0, "left": "1.5", "rear": None}
    roles = {item["role"] for item in parity.load_step(steps / parity.step_name("roadLevels", "python"), "python")["run"]["inputs"]}
    assert {"config", "baseplanPlots", "roadSurface:A__Road", "roadSurface:B__Ground", "plotsCsv", "labels"} <= roles


def test_r348_read_steps_refuses_a_cad_output_not_in_meters_and_r349_a_landxml_not_in_meter(tmp_path, capsys):
    settings, values = made_settings(tmp_path)
    layers = tmp_path / "reads" / values["drawings"]["baseplan"] / "layers.json"
    layers.write_text(json.dumps({"units": "Millimeters"}), encoding="utf-8")
    assert read_steps.main(["--settings", str(settings), "--steps", str(tmp_path / "steps"), "--only", "labels"]) == 2
    assert "Millimeters" in capsys.readouterr().err
    assert not (tmp_path / "steps" / parity.step_name("labels", "python")).exists()
    files = write_roads(tmp_path / "landxml")
    files[1].write_text(files[1].read_text(encoding="utf-8").replace('linearUnit="meter"', 'linearUnit="foot"'), encoding="utf-8")
    assert read_steps.main(["--settings", str(settings), "--steps", str(tmp_path / "steps"), "--only", "surfaces", "--landxml", str(tmp_path / "landxml"),
                            "--scratch", str(tmp_path / "scratch")]) == 2
    assert "B__Ground.xml is in foot" in capsys.readouterr().err
