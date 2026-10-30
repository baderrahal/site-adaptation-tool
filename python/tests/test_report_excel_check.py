"""R365. The second route of the report Excel, python/report_excel_check.py, on small workbooks written here with openpyxl in the layout the
worker writes, from a made up plan with its status, walls and missing list, override list, step checkRows and plots.csv, the same as
ReportExcelTests of C# (R10). The counts it works out are those the C# test asserts, and each planted fault is caught by its own check."""
import hashlib
import json

import openpyxl
import pytest
from openpyxl.styles import Border, Font, PatternFill, Side

import report_excel_check as check

KEY = "rules/r.json levelProfiles.parsons2024Rev6.keyRules."


def valued(value, source, status, came):
    return {"value": value, "source": source, "status": status, "statusFrom": came}


DETAILS = {
    "unit from locationPlan": {"entries": [{"status": "DOCUMENT", "document": "the location plan"}]},
    "mirrored from computed hand villa": {"entries": [{"entry": KEY + "hand villa.rule", "status": "MEASURED", "share": 0.96, "strength": "strong"}], "strength": "strong"},
    "levels.entry from estimate": {"entries": [{"entry": KEY + "entry.rule", "status": "MEASURED", "share": 0.088, "strength": "weak"}], "strength": "weak"},
    "levels.entry from override by entryM": {"entries": [{"status": "RECOMMENDED", "question": "Q053", "where": "section 6"}], "overrides": ["entryM"]},
    # A strong entry before the UNKNOWN one, so the basis is that of the worst entry by worstFirst and never the first entry
    "floors type from computed": {"entries": [{"entry": KEY + "generic.rule", "status": "MEASURED", "share": 0.99, "strength": "strong"},
                                              {"entry": "config/project.json floors.types", "status": "UNKNOWN", "why": "no entry"}]},
    # Review 1 of PR 96: a height the heightM override touched still reads keyRules.entry, and counts under no keyRules entry
    "levels.height from estimate by heightM": {"entries": [{"status": "RECOMMENDED", "question": "Q053", "where": "section 6"},
                                                           {"entry": KEY + "entry.rule", "status": "MEASURED", "share": 0.088, "strength": "weak"}], "overrides": ["heightM"]},
    "walls.base from keyRules": {"entries": [{"entry": KEY + "wallBase.rule", "status": "MEASURED", "share": 0.98, "strength": "strong"}], "strength": "strong"},
    "walls.top from keyRules by heightM": {"entries": [{"status": "RECOMMENDED", "question": "Q053", "where": "section 6"}], "overrides": ["heightM"]},
}
TYPES = {"generic": "Generic 150mm", "yard": "yard", "ramp": "Ramps 2%"}


def plot(number, unit, mirrored, moved, floors, overrides=None):
    item = {"plot": number, "unit": valued(unit, "locationPlan", "DOCUMENT", "unit from locationPlan"),
            "mirrored": valued(mirrored, "computed", "MEASURED", "mirrored from computed hand villa"),
            "levels": {"entry": valued(10.5, "override", "RECOMMENDED", "levels.entry from override by entryM") if moved
                       else valued(10.25, "estimate", "MEASURED", "levels.entry from estimate")},
            "floors": {kind: {"type": valued(TYPES[kind], "computed", "UNKNOWN", "floors type from computed")} for kind in floors}}
    if overrides is not None:
        item["overrides"] = overrides
        item["levels"]["height"] = valued(0.2, "estimate", "RECOMMENDED", "levels.height from estimate by heightM")
    return item


PLAN = {"format": "site-adaptation-plan", "run": {"runId": "b7"}, "statusFrom": DETAILS,
        "summary": {"skipped": [], "rules": {"levelProfile": "parsons2024Rev6", "overrideSource": ["parsons"], "overrideModel": []}},
        "plots": [plot("101", "VD4", False, True, ["generic", "ramp"], {"applied": ["entryM", "parameter CASE", "rampFrontM", "parameter RRE_ZZ_SP_RVT-Name"], "notApplied": {"rampFrontFallM": "no direction"}}),
                  plot("102", "DP1", True, False, ["generic", "yard"]), plot("201", "VD1", False, False, ["generic"]), plot("999", "VD1", False, False, ["generic"])],
        "walls": [{"plot": "101", "edge": 1, "plots": valued(["101", "102"], "locationPlan", "DOCUMENT", "unit from locationPlan"),
                   "base": valued(10.1, "keyRules", "MEASURED", "walls.base from keyRules"), "top": valued(10.4, "keyRules", "RECOMMENDED", "walls.top from keyRules by heightM"),
                   "basePlot": "102", "topPlot": "101"}],
        "missing": [{"plot": "102", "written": True, "missing": [{"key": "floors.ramp", "why": "no ramp type"}]},
                    {"plot": "103", "written": False, "missing": [{"key": "unit", "why": "no unit code"}],
                     "overrides": {"applied": [], "notApplied": {"mirrored": "plot not written"}}}],
        "conflicts": []}


def entry(number, field, value, reason, source="parsons"):
    return {"plot": number, "field": field, "value": value, "source": source, "model": "model_%s.rvt" % source, "reason": reason}


OVERRIDES = {"overrides": [entry("101", "entryM", {"gapToRuleM": 0.01}, "breaks roadLevel"), entry("101", "heightM", {"gapToRuleM": 0.5}, "native", "native"),
                           entry("101", "parameter CASE", "D", "breaks typologyMirrored"), entry("101", "rampFrontFallM", 0.02, "breaks roadAtEachCorner"),
                           entry("101", "rampFrontM", {"gapToRuleM": 0.003}, "breaks roadLevel"),
                           entry("101", "parameter RRE_ZZ_SP_RVT-Name", "101 | VD4 | 1", "breaks plotCodeStepsAndFacade"),
                           entry("103", "mirrored", True, "breaks blockMirrorRoadSide")]}


def row(check_type, field, came, use, why=None):
    return {"check": check_type, "field": field, "from": came, "use": use, "rule": None, "toolValue": None, "past": None, "why": why}


CHECK_ROWS = {"plots": {number: {"written": {"rows": rows}} for number, rows in (
    ("101", [row("Level", "entryM", "override", "Past"), row("Level", "rampFrontFallM", "override", "Rule", "no direction"), row("Parameter", "parameter CASE", "override", "Past")]),
    ("102", [row("Floor", "floors.ramp", "missing", "Rule")]),
    ("103", [row("Hand", "mirrored", "override", "Rule", "plot not written"), row("Unit code", "unit", "missing", "Rule")]), ("201", []))}}
RULES = {"levelProfiles": {"parsons2024Rev6": {"keyRules": {"note": "made up", "hand villa": {"rule": "blockMirrorRoadSide", "share": 0.96, "strong": True},
                                                            "entry": {"rule": "roadLevel", "share": 0.088, "strong": False},
                                                            "wallBase": {"rule": "yardLevelOfTheHigherPlot", "share": 0.98, "strong": True},
                                                            "wallTop": {"rule": "higherFflPlusRev6", "share": 0.35, "strong": False},
                                                            "ssl": {"rule": "fflPlusGroundSsl", "notMeasured": "no key"}}}}}
# The step summary of the build, 103 with no unit, 150 of plots.csv in 012 outside the scope as a key cover holds part of an NBH
SUMMARY = {"plots": {number: {"written": {"unit": unit, "written": unit is not None}} for number, unit in
                     (("101", "VD4"), ("102", "DP1"), ("103", None), ("201", "VD1"), ("999", "VD1"))}}
# R381: the steps of the order of Bader's changes of 28 Oct 2026, the link parameters under 7 Place villas, as config/project.json holds them
STEP_OF_FIELD = {"rampType": 5, "rampFrontM": 5, "genericFloor": 5, "rampOutline": 5, "yardOutline": 5, "rampFrontFallM": 5, "entryM": 6, "heightM": 6, "steps": 6,
                 "parameter RRE_ZZ_SP_Prototype-FFL": 7, "parameter RRE_ZZ_SP_RVT-Name": 7}
CONFIG = {"profile": "rules/r.json", "addin": {"report": {"stepOfField": STEP_OF_FIELD, "defaultStep": 7, "wallDecisions": ["wallBase", "wallTop", "wallLine", "frontOpen"]}},
          "valueStatus": {"assumed": ["RECOMMENDED", "CHOSEN by Code", "UNKNOWN", "MEASURED weak"], "worstFirst": ["UNKNOWN", "CHOSEN by Code", "RECOMMENDED", "MEASURED weak"]},
          }
PLOTS = "locationPlanNumber,typologyRev9,zone,nbhId\n101,VD4,Zone 2,012\n102,DP1,Zone 2,012\n103,,Zone 2,012\n150,,Zone 2,012\n201,VD1,Zone 3,013\n"
FILL = PatternFill(fill_type="solid", fgColor="FFDDEBF7")
THIN = Side(style="thin")


def build(tmp_path, rules=RULES, summary=SUMMARY, revit=None):
    paths = {}
    for name, content in (("plan", PLAN), ("overrides", OVERRIDES), ("checkRows", CHECK_ROWS), ("summary", summary), ("rules", rules), ("config", CONFIG)):
        paths[name] = tmp_path / (name + ".json")
        paths[name].write_text(json.dumps(content), encoding="utf-8")
    paths["plots"] = tmp_path / "plots.csv"
    paths["plots"].write_text(PLOTS, encoding="utf-8")
    if revit is not None:
        paths["revit"] = tmp_path / "revit_counts.json"
        paths["revit"].write_text(json.dumps(revit), encoding="utf-8")
    mock = check.mock_labels()
    inputs = check.load_inputs(paths["plan"], paths["overrides"], paths["checkRows"], paths["summary"], paths["plots"], paths["rules"], paths["config"], paths.get("revit"))
    return paths, mock, check.expected(inputs, "parsons2024Rev6", mock)


@pytest.fixture
def made(tmp_path):
    return build(tmp_path)


def plain(value):
    return json.dumps(value[1], separators=(",", ":")) if isinstance(value, tuple) else value


def write_book(path, want, mock, settings):
    """A workbook in the layout of the worker, filled from want."""
    book = openpyxl.Workbook()
    summary = book.active
    summary.title = "Summary"
    summary.append([want["topLine"]])
    summary["A1"].font = Font(bold=True)
    summary.merge_cells("A1:F1")
    summary.append([])
    summary.append(mock["stepColumns"])
    for line in want["steps"]:
        summary.append(line)
    for start, name, header, lines in ((13, mock["countsBox"], mock["countColumns"], want["counts"]), (19, mock["rulesBox"], list(check.RULE_COLUMNS), want["box"])):
        summary.cell(row=start, column=1, value=name).font = Font(bold=True)
        for offset, values in enumerate([header] + lines):
            for column, value in enumerate(values, 1):
                target = summary.cell(row=start + 1 + offset, column=column, value=value)
                target.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
                if start == 19 and column == 3 and offset > 0 and check.is_number(value):
                    target.number_format = check.SHARE_FORMAT
    headers = {"Placed": mock["placed"], "Skipped": mock["skipped"], "Overrides": list(check.OVERRIDE_COLUMNS), "Assumed": mock["assumed"],
               "Issues": mock["issueColumns"], "Rules": list(check.RULE_SHEET_COLUMNS)}
    for name in mock["sheets"][1:-1]:
        sheet = book.create_sheet(name)
        sheet.append(headers[name])
        for line in want["sheets"][name]:
            sheet.append([plain(value) for value in line])
    for column, name in enumerate(mock["placed"], 1):
        if name in mock["editable"]:
            book["Placed"].cell(row=1, column=column).fill = FILL
    sheet = book.create_sheet("Settings")
    sheet.append(["Setting", "Value"])
    for name in check.SETTINGS:
        sheet.append([name, settings.get(name, want["settings"].get(name))])
    book.save(path)


SETTINGS = {"Run id": "rx1", "Run time": "2026-10-29T08:00:00Z"}


def test_r365_the_plan_side_works_out_the_counts_the_c_sharp_test_asserts(made):
    paths, mock, want = made
    book = want["books"]["012"]
    assert sorted(want["books"]) == ["012", "013"]
    assert want["gaps"]["plots with no nbhId"] == 1
    assert book["topLine"] == "[file name].rvt, NBH 012, 3 plots, parsons2024Rev6, run time [n] min"
    # R381 moved this from the buttons 1 to 7 of the mock of 25 Oct to the steps 1 to 8 of the new order: rampFrontM and rampFrontFallM under 5,
    # entryM under 6, CASE and RVT-Name under 7 Place villas, 2 Read data over the faces of the scope, 150 left out as the scope is no whole
    # project, and the weak wallTop the Issues of 8 Walls, each line with the words of no manual rule block
    none = "no manual rule block"
    assert book["steps"] == [["1 Set up site"] + ["not run yet"] * 4 + [none], ["2 Read data", 2, 1, 0, 0, none], ["3 Review data"] + ["not run yet"] * 4 + [none],
                             ["4 Check", 3, 0, 6, 6, none], ["5 Floors and ramps", 4, 4, 1, 1, none], ["6 Levels", 2, 1, 1, 0, none],
                             ["7 Place villas", 2, 1, 2, 1, none], ["8 Walls", 1, 0, 1, 1, none]]
    assert book["sheets"]["Issues"][0] == ["4 Check", "101", "Level", "entryM", "override", "Past", None, none]
    assert book["sheets"]["Issues"][-1] == ["8 Walls", None, "Weak rule", "wallTop", KEY + "wallTop.rule", None, "higherFflPlusRev6 share 0.35, weak", none]
    assert len(book["sheets"]["Issues"]) == 7
    assert book["counts"] == [["Links", "not placed yet", 2], ["Floors", "not placed yet", 4], ["Levels", "not placed yet", "no run of button 6 yet"]]
    assert book["box"] == [["hand villa", "blockMirrorRoadSide", 0.96, "strong", "2 plots"], ["entry", "roadLevel", 0.088, "weak", "1 plots"],
                           ["wallBase", "yardLevelOfTheHigherPlot", 0.98, "strong", "1 plots"], ["wallTop", "higherFflPlusRev6", 0.35, "weak", "0 plots"], ["ssl", "fflPlusGroundSsl", "not measured", "not measured", "0 plots"]]
    assert book["sheets"]["Assumed"][2] == ["101", "floors.generic.type", "Generic 150mm", "UNKNOWN", "config/project.json floors.types"]
    assert book["sheets"]["Assumed"][1] == ["101", "levels.height", 0.2, "RECOMMENDED", "Q053"]
    assert book["sheets"]["Assumed"][4] == ["102", "levels.entry", 10.25, "MEASURED weak", KEY + "entry.rule share 0.088"]
    assert len(book["sheets"]["Assumed"]) == 7
    assert book["sheets"]["Rules"][0] == ["Q053", "RECOMMENDED", "Q053", 1, 3]
    assert book["sheets"]["Rules"][2] == [KEY + "entry.rule", "MEASURED weak", KEY + "entry.rule share 0.088", 2, 2]
    assert book["sheets"]["Rules"][-1] == ["the location plan", "DOCUMENT", "the location plan", 2, 3]
    assert book["settings"]["Plan SHA-256"] == hashlib.sha256(paths["plan"].read_bytes()).hexdigest()
    assert book["settings"]["Editable columns"] == "Type, Mirror"


def test_r365_a_workbook_in_the_layout_of_the_worker_holds_with_no_mismatch(made, tmp_path):
    paths, mock, want = made
    folder = tmp_path / "out"
    folder.mkdir()
    for book in want["books"].values():
        write_book(folder / book["file"], book, mock, SETTINGS)
    result, bad = check.check_folder(folder, want, mock, "rx1")
    assert not bad
    assert sum(count for key, count in result.items() if key.startswith("mismatch")) == 0
    # The Mirror cell of each of the 8 Placed rows of the two workbooks, the flag of the 3 links and empty on the 5 floors
    assert result["match Placed Mirror"] == 8


@pytest.mark.parametrize("fault, caught", [
    (lambda book: book["Summary"].cell(row=6, column=5, value=5), "Summary row 6 Issues"),
    (lambda book: book["Summary"].cell(row=21, column=5, value="3 plots"), "Rules applied row 21 Plots"),
    (lambda book: book["Summary"].cell(row=7, column=6, value="M01"), "Summary row 7 Manual rule"),
    (lambda book: book["Issues"].cell(row=8, column=8, value="M11"), "Issues row 8 Manual rule"),
    (lambda book: setattr(book["Summary"].cell(row=14, column=1), "border", Border()), "Summary box borders"),
    (lambda book: setattr(book["Placed"].cell(row=1, column=5), "fill", FILL), "fills on"),
    (lambda book: book["Placed"].cell(row=3, column=4, value=False), "Placed row 3 Mirror"),
    (lambda book: book["Placed"].insert_rows(3), "Placed max_row"),
    (lambda book: book["Assumed"].cell(row=2, column=5, value="Q053 share 1"), "Assumed row 2 Basis"),
    (lambda book: book["Settings"].cell(row=9, column=2, value="0" * 64), "Settings Plan SHA-256"),
    (lambda book: book.move_sheet("Assumed", offset=-1), "sheets"),
])
def test_r365_each_planted_fault_is_caught_by_its_check(made, tmp_path, fault, caught):
    paths, mock, want = made
    book = want["books"]["012"]
    path = tmp_path / book["file"]
    write_book(path, book, mock, SETTINGS)
    opened = openpyxl.load_workbook(path)
    fault(opened)
    opened.save(path)
    result, bad = check.check_xlsx(path, book, mock, "rx1")
    assert any(caught in item for item in bad), bad


def test_r365_the_labels_come_from_the_mock_files_and_their_additions_of_7_7(made):
    paths, mock, want = made
    assert mock["sheets"] == ["Summary", "Placed", "Skipped", "Overrides", "Assumed", "Issues", "Rules", "Settings"]
    assert mock["placed"] == ["Plot", "Kind", "Type", "Mirror", "Element id"]
    assert mock["editable"] == ["Type", "Mirror"]
    # R381: the steps 1 to 8 of the order of docs/design/13-workflow-changes.md, 9 Report the step of the report
    assert mock["steps"] == ["1 Set up site", "2 Read data", "3 Review data", "4 Check", "5 Floors and ramps", "6 Levels", "7 Place villas", "8 Walls"]
    assert mock["reportStep"] == "9 Report"
    assert mock["stepColumns"][-1] == "Manual rule"


def test_r365_a_placed_sheet_short_of_one_floor_row_gives_the_mismatch_as_an_issue_row(made):
    paths, mock, want = made
    written = [plot for plot in PLAN["plots"] if plot["plot"] in ("101", "102")]
    placed = [row for row in want["books"]["012"]["sheets"]["Placed"]][:-1]
    rows, issues = check.counts_box(placed, written, mock, None, "M15")
    assert rows[1] == ["Floors", "not placed yet", 3]
    assert issues == [["9 Report", None, "Counts", "Floors", None, None, "the plan holds 4, Revit not placed yet, Excel 3 (R365)", "M15"]]


def test_r381_read_data_counts_the_faces_of_plots_csv_in_the_scope_with_a_unit_as_done(tmp_path):
    paths, mock, want = build(tmp_path, summary=dict(SUMMARY, scope="project"))
    assert want["books"]["012"]["steps"][1] == ["2 Read data", 2, 2, 0, 0, "no manual rule block"]


MANUAL = {"manual": {"name": "made up"},
          "rules": [{"id": "M01", "text": "made up", "group": "Plots and clusters", "manualRef": "2.1", "checkedIn": ["4 Check", "7 Place villas"]},
                    {"id": "M04", "text": "made up", "group": "Levels and steps", "manualRef": "2.2.A", "checkedIn": ["6 Levels", "8 Walls"]},
                    {"id": "M02", "text": "made up", "group": "Plots and clusters", "manualRef": "2.1.E", "checkedIn": ["4 Check"]},
                    {"id": "M15", "text": "made up", "group": "Structure", "manualRef": "3.2.A", "checkedIn": ["9 Report"]}]}


def test_r381_every_summary_and_issues_line_names_the_manual_rule_ids_whose_steps_take_its_step(tmp_path):
    paths, mock, want = build(tmp_path, rules=dict(RULES, manualRules=MANUAL))
    book = want["books"]["012"]
    assert [line[-1] for line in book["steps"]] == ["no manual rule", "no manual rule", "no manual rule", "M01, M02", "no manual rule", "M04", "M01", "M04"]
    assert book["sheets"]["Issues"][0][-1] == "M01, M02" and book["sheets"]["Issues"][-1][-1] == "M04"
    with pytest.raises(ValueError, match="no list rules each with a text id"):
        build(tmp_path, rules=dict(RULES, manualRules={"rules": [{"id": "M01", "text": "made up"}]}))


def test_r381_the_revit_counts_fill_the_revit_column_and_a_mismatch_of_any_two_is_an_issue(tmp_path):
    floors = lambda generic, yard, ramp: {"generic": generic, "yard": yard, "ramp": ramp}
    revit = {"format": "siteAdaptationRevitCounts", "formatVersion": 1, "levelRun": "r6",
             "plots": {"101": {"links": 1, "floors": floors(1, 0, 1), "level": True}, "102": {"links": 1, "floors": floors(1, 0, 0), "level": False},
                       "201": {"links": 1, "floors": floors(1, 0, 0), "level": True}}}
    paths, mock, want = build(tmp_path, revit=revit)
    book = want["books"]["012"]
    assert book["counts"] == [["Links", 2, 2], ["Floors", 3, 4], ["Levels", 1, "no run of button 6 yet"]]
    assert book["sheets"]["Issues"][-2:] == [["9 Report", None, "Counts", "Floors", None, None, "the plan holds 4, Revit 3, Excel 4 (R381)", "no manual rule block"],
                                             ["9 Report", None, "Counts", "Levels", None, None, "the plan holds 2, Revit 1, Excel no run of button 6 yet (R381)", "no manual rule block"]]
    assert want["books"]["013"]["counts"] == [["Links", 1, 1], ["Floors", 1, 1], ["Levels", 1, "no run of button 6 yet"]]
    with pytest.raises(ValueError, match="no file of the counts"):
        build(tmp_path, revit={"format": "another"})


def test_r365_settings_with_no_block_addin_report_are_refused(made, tmp_path):
    paths, mock, want = made
    bare = dict(CONFIG)
    bare.pop("addin")
    paths["config"].write_text(json.dumps(bare), encoding="utf-8")
    inputs = check.load_inputs(paths["plan"], paths["overrides"], paths["checkRows"], paths["summary"], paths["plots"], paths["rules"], paths["config"])
    with pytest.raises(ValueError, match="no block addin.report"):
        check.expected(inputs, "parsons2024Rev6", mock)
