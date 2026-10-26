"""R351. The location plan steps of the parity harness, python/read_steps.py, on a made up page, baseplan, schedules and sales workbook
built in code (R10), and the settings of the location plan in config/project.json. dotnet/SiteAdaptation.Tests/LocationPlanTests.cs holds
the C# side on the same made up pieces."""
import json

import numpy as np
import openpyxl
import pytest

import c3d_run
import data_book
import location_plan
import parity
import project_profile
import read_steps
from test_read_steps import made_settings

TRUTH = np.array([[1.7, 0.3], [0.3, -1.7], [1000.0, 2000.0]])
STEPS = ["locationLabels", "locationFit", "controlCheck", "controlFaces"]


def page_of(site):
    return tuple((np.asarray(site) - TRUTH[2]) @ np.linalg.inv(TRUTH[:2]))


def chars(text, x, top, size=1.48, step=0.7):
    return [{"text": ch, "x0": x + k * step, "x1": x + k * step + 0.65, "top": top, "bottom": top + 1.0, "size": size, "upright": True} for k, ch in enumerate(text)]


def label_chars(code, number, site):
    """A code line over its number line whose label point, the middle of the code top and the number bottom, lies on the page at site."""
    x, y = page_of(site)
    left = lambda text: x - ((len(text) - 1) * 0.7 + 0.65) / 2
    return chars(code, left(code), y - 1.28) + chars(number, left(number), y + 0.28)


def plot(label, x0, y0, x1, y1):
    return {"kind": "plot_shape", "plot": label, "vertices": [[x0, y0], [x1, y0], [x1, y1], [x0, y1]], "type": "DP1", "area": 100.0}


# Six faces of 10 m in two rows of three, labelled 101 to 106, numbered 111 to 116 by the IMC schedule of shift 10
PLOTS = [plot(str(101 + k), 1000.0 + 10 * (k % 3), 1900.0 + 10 * (k // 3), 1010.0 + 10 * (k % 3), 1910.0 + 10 * (k // 3)) for k in range(6)]
CODES = ["DP1 (QP)", "DP1-SA", "TH3-0 (BB)", "VL3-0", "DP1", "XX9"]


def made_page():
    found = []
    for k, code in enumerate(CODES):
        found += label_chars(code, str(111 + k), (1005.0 + 10 * (k % 3), 1905.0 + 10 * (k // 3)))
    marks = [page_of(site) for site in ((500.0, 1880.0), (1030.0, 1880.5), (1560.0, 1881.0))]
    words, curves = [], []
    for mark in marks:
        east, north = TRUTH[2] + np.asarray(mark) @ TRUTH[:2]
        words += [{"text": t, "x0": mark[0] + 20.0 + k * 5.0, "top": mark[1] - 8.0} for k, t in enumerate(("E", "=", "%.3f" % east, "N", "=", "%.3f" % north))]
        curves.append({"x0": mark[0] - 1.98, "x1": mark[0] + 1.98, "top": mark[1] - 1.98, "bottom": mark[1] + 1.98, "pts": [mark] * 6})
    return {"chars": found, "words": words, "curves": curves, "lines": [], "small": [], "meta": {}}


def schedule(path, sheet, numbers):
    workbook = openpyxl.Workbook()
    rows = workbook.active
    rows.title = sheet
    rows.append(["Schedule of plots"])
    rows.append(["Plot No.", "Unit", "FFL", "SSL", "El @ Entry", "Steps Needed"])
    for number in numbers:
        rows.append([number, "DP1", 10.5, 10.3, 10.1, "2"])
    rows.append(["Total", None, None, None, None, None])
    workbook.save(path)
    return path


def sales_book(path, values):
    plan = values["dataBook"]["locationPlan"]
    workbook = openpyxl.Workbook()
    sales = workbook.active
    sales.title = plan["sales"]["sheet"]
    columns = plan["sales"]["columns"]
    sales.append([columns["plot"], columns["typology"], columns["nbh"], columns["zone"]])
    for k in range(5):
        sales.append([111 + k, "DP1", "012", "Zone 2"])
    table = workbook.create_sheet(plan["codeTable"]["sheet"])
    table.append(["title"])
    names = plan["codeTable"]["columns"]
    table.append([names["form"], names["old"], names["new"]])
    for row in (("Duplex", "DP1", "DP1"), ("Villa", "DP1", "VD1")):
        table.append(list(row))
    workbook.save(path)
    return path


def test_r351_the_location_plan_settings_and_the_neighbourhood_schedules_live_in_config_and_no_longer_in_the_probe_settings():
    settings = project_profile.load_settings()
    probe = json.loads((c3d_run.CONFIG).read_text(encoding="utf-8"))
    plan = settings.get("dataBook.locationPlan")
    for key in ("pdf", "patterns", "codeKinds", "pairDyPt", "fitNeighbourhood", "fitDropFactor", "codeTable", "controlMark", "controlSearchPt", "sales"):
        assert key in plan, key
    assert "locationPlan" not in probe["dataBook"] and "neighbourhoods" not in probe["connect"]
    connect = settings.get("connect")
    assert plan["fitNeighbourhood"] in connect["neighbourhoods"] and connect["scheduleSheet"]
    merged = c3d_run.load_config()
    assert merged["connect"]["neighbourhoods"] == connect["neighbourhoods"] and merged["dataBook"]["locationPlan"]["pdf"] == plan["pdf"]
    # The probe settings that did not move stay, and a key both files hold is refused with its path
    assert merged["connect"]["snapM"] == probe["connect"]["snapM"]
    with pytest.raises(ValueError, match="connect.scheduleSheet"):
        c3d_run.lay_over({"scheduleSheet": "a"}, {"scheduleSheet": "b"}, "connect")


def test_r351_schedule_rows_read_the_sheet_config_names_and_a_file_a_run_names(tmp_path):
    values = project_profile.load_settings()
    connect = {"neighbourhoods": {"IMC": {"schedule": "imc.xlsx", "shift": 10}, "IME": None}, "scheduleSheet": "MP Sheet"}
    path = schedule(tmp_path / "imc.xlsx", "MP Sheet", [2290, "2291 "])
    found = data_book.schedule_rows(values, connect, {"IMC": path})
    assert sorted(found) == [("IMC", "2290"), ("IMC", "2291")]
    connect["scheduleSheet"] = "Another"
    with pytest.raises(ValueError, match="Another"):
        data_book.schedule_rows(values, connect, {"IMC": path})


def test_r351_the_written_values_of_a_step_take_the_decimals_of_their_field():
    found = read_steps.entry({"scale": 1.76545, "site": [[2.0465, 1.0005]], "face": "12"}, {"scale": {"kind": "points", "decimals": 4}, "site": {"kind": "points", "decimals": 3},
                                                                                             "face": {"kind": "noGeometry"}})
    assert found["written"] == {"scale": 1.7655, "site": [[2.047, 1.001]], "face": "12"}
    assert found["raw"]["scale"] == 1.76545


def test_r351_read_steps_numbers_a_made_up_page_by_the_schedule_fit_and_checks_it_by_the_control_point_fit(tmp_path, monkeypatch, capsys):
    path, values = made_settings(tmp_path)
    folder = tmp_path / "reads" / values["drawings"]["baseplan"]
    (folder / "plots.json").write_text(json.dumps({"source": "b.dwg", "plots": PLOTS}), encoding="utf-8")
    values["connect"]["neighbourhoods"] = {"IMC": {"schedule": "imc.xlsx", "shift": 10}, "IMD": {"schedule": "imd.xlsx", "shift": 11}}
    path.write_text(json.dumps(values), encoding="utf-8")
    sheet = values["connect"]["scheduleSheet"]
    imc = schedule(tmp_path / "imc.xlsx", sheet, [111, 112, 113, 114, 115, 116])
    imd = schedule(tmp_path / "imd.xlsx", sheet, [900])
    sales = sales_book(tmp_path / "sales.xlsx", values)
    pdf = tmp_path / "plan.pdf"
    pdf.write_bytes(b"made up, read_page is replaced")
    monkeypatch.setattr(location_plan, "read_page", lambda found, plan, check=None: made_page())
    steps = tmp_path / "steps"
    code = read_steps.main(["--settings", str(path), "--steps", str(steps), "--only", *STEPS, "--plot-numbers", str(pdf), "--sales", str(sales),
                            "--schedule", "IMC=%s" % imc, "--schedule", "IMD=%s" % imd])
    assert code == 0, capsys.readouterr().err
    labels = parity.load_step(steps / parity.step_name("locationLabels", "python"), "python")
    assert sorted(labels["plots"]) == ["111", "112", "113", "114", "115", "116"]
    first = labels["plots"]["111"]["raw"]
    assert first["face"] == "101" and first["offset"] == 10 and first["group"] == "IMC" and first["newCode"] == "DP1" and first["rev9"] == "DP1" and first["zone"] == "Zone 2"
    assert labels["plots"]["116"]["raw"]["rev9"] is None and labels["plots"]["112"]["raw"]["newCode"] == "VD1" and labels["plots"]["115"]["raw"]["newCode"] is None
    assert labels["plots"]["113"]["raw"]["site"] == [pytest.approx(1025.0), pytest.approx(1905.0)]
    roles = {item["role"] for item in labels["run"]["inputs"]}
    assert {"config", "baseplanPlots", "plotNumbers", "schedule:IMC", "schedule:IMD", "sales"} <= roles
    fit = parity.load_step(steps / parity.step_name("locationFit", "python"), "python")["plots"]
    assert list(fit) == ["IMC"] and fit["IMC"]["raw"]["kept"] == ["111", "112", "113", "114", "115", "116"] and fit["IMC"]["raw"]["role"] == "fit"
    assert fit["IMC"]["raw"]["metresPerPoint"] == pytest.approx(np.sqrt(abs(np.linalg.det(TRUTH[:2]))))
    check = parity.load_step(steps / parity.step_name("controlCheck", "python"), "python")["plots"]["control"]["raw"]
    assert check["marks"] == 3 and check["refineFaces"] == 6 and check["facesNumberedOtherwise"] == 0 and check["labelsOnAFaceSchedule"] == 6
    assert parity.load_step(steps / parity.step_name("controlFaces", "python"), "python")["plots"] == {}


def test_r351_the_location_plan_steps_need_the_plan_and_the_sales_workbook(tmp_path, capsys):
    path, _ = made_settings(tmp_path)
    assert read_steps.main(["--settings", str(path), "--steps", str(tmp_path / "steps"), "--only", "locationLabels"]) == 2
    assert "--plot-numbers and --sales" in capsys.readouterr().err
    assert not (tmp_path / "steps" / parity.step_name("locationLabels", "python")).exists()
