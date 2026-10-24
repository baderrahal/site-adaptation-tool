import json
from pathlib import Path

import pytest
from openpyxl import Workbook

import full_run
import levels
import levels_proof
import link_table
import project_profile

SETTINGS = project_profile.load_settings()
PROFILE = project_profile.load_profile(project_profile.profile_path(SETTINGS))
SCHEDULES = SETTINGS.get("schedules")
SHIFTS = link_table.load_shifts(levels_proof.DEFAULT_SHIFTS)
HEADERS = ["Plot & Unit", "Plot No.", "Unit", "FFL", "SSL", "El @ Entry*", "Height**", "Site Steps Needed", "Steps Needed"]


def schedule(path, rows, header_at=3):
    book = Workbook()
    sheet = book.active
    sheet.title = full_run.DEFAULT_SHEET
    for _ in range(header_at - 1):
        sheet.append(["a note above the table"])
    sheet.append(HEADERS)
    for plot, unit, ffl, entry, steps in rows:
        sheet.append([str(plot) + "-" + unit, plot, unit, ffl, ffl - 0.075, entry, ffl - entry, "x", steps])
    sheet.append([None, None, None, None])
    book.save(path)
    return path


ROWS = [
    (1001, "DT1", 10.2, 10.1, "1 steps"),
    (1002, "DP2", 20.3, 20.0, "2 step"),
    (1003, "DP2", 20.3, 20.1, "1 steps"),
    (1004, "VL4", 30.05, 29.9, "1 steps"),
    (1005, "DP2", 41.0, 40.2, "3 steps"),
    (1006, "DP2", 41.0, 40.8, "1 steps"),
]


def lone(plot, ffl, entry=10.0, unit="VL4", steps=1):
    return {"sheetRow": 1, "plot": str(plot), "unit": unit, "ffl": ffl, "ssl": ffl - 0.075, "entry": entry, "steps": steps}


def test_r67_a_schedule_is_read_from_its_header_row_with_the_steps_as_numbers(tmp_path):
    rows = levels_proof.read_schedule(schedule(tmp_path / "a.xlsx", ROWS, header_at=5))
    assert [(row["plot"], row["unit"], row["ffl"], row["steps"]) for row in rows] == [("1001", "DT1", 10.2, 1), ("1002", "DP2", 20.3, 2), ("1003", "DP2", 20.3, 1), ("1004", "VL4", 30.05, 1), ("1005", "DP2", 41.0, 3), ("1006", "DP2", 41.0, 1)]
    assert rows[0]["sheetRow"] == 6


def test_r67_consecutive_rows_with_the_same_ffl_are_one_cluster_and_a_difference_in_the_third_decimal_is_two():
    rows = [{"ffl": 5.0}, {"ffl": 5.0}, {"ffl": 6.0}, {"ffl": 5.0}, {"ffl": 5.00001}]
    assert [len(group) for group in levels_proof.clusters_by_ffl(rows)] == [2, 1, 2]
    assert [len(group) for group in levels_proof.clusters_by_ffl([{"ffl": 5.0}, {"ffl": 5.004}, {"ffl": 5.004}])] == [1, 2]


def test_r67_the_report_counts_matches_and_lists_each_mismatch_with_both_ffl_values(tmp_path):
    path = schedule(tmp_path / "a.xlsx", ROWS)
    report = levels_proof.build_report(PROFILE, [path])
    summary = report["summary"]
    assert (summary["plots"], summary["clusters"], summary["fflMatches"], summary["stepsMatches"], summary["bothMatch"], summary["sslMatches"]) == (6, 4, 5, 6, 5, 5)
    assert summary["countRuleAloneMatches"] == 6 and summary["unitsNotInProfile"] == [] and summary["cadClusters"] is None
    assert [(m["plot"], m["unit"], m["sheetFfl"], m["computedFfl"], m["difference"]) for m in report["fflMismatches"]] == [("1004", "VL4", 30.05, 30.0, -0.05)]
    assert report["stepsMismatches"] == []
    assert levels_proof.summary_line(report) == "FFL matches 5 of 6, steps matches 6 of 6, both 5, SSL 5, count rule alone 6"
    assert levels_proof.cad_line(report) is None


def test_r67_a_computed_ffl_within_half_a_millimetre_of_the_sheet_matches_and_a_hair_over_does_not():
    rule = levels.count_rule(PROFILE)
    # A lone VL4 plot at 10.0 has the computed FFL 10.1. 10.1005 is half a millimetre away, plus float noise
    rows = [lone(1, 10.1004), lone(2, 10.1005), lone(3, 10.1006), lone(4, 10.0996), lone(5, 10.0994)]
    found, _ = levels_proof.compare_schedule(PROFILE, rows, rule)
    assert [item["fflMatch"] for item in found] == [True, True, False, True, False]
    assert [item["sslMatch"] for item in found] == [True, True, False, True, False]


def test_r67_a_unit_code_the_profile_lacks_is_listed_and_not_mapped(tmp_path):
    rows = [(1, "VL4", 10.0, 9.9, "1 steps"), (2, "ZZ9", 11.0, 10.9, "1 steps"), (3, "QQ1", 12.0, 11.9, "1 steps")]
    report = levels_proof.build_report(PROFILE, [schedule(tmp_path / "a.xlsx", rows)])
    assert report["summary"]["unitsNotInProfile"] == ["QQ1", "ZZ9"]
    assert report["summary"]["plots"] == 3


def test_r67_a_steps_difference_is_reported_and_not_forced(tmp_path):
    rows = [(1, "VL4", 10.0, 9.9, "2 step"), (2, "VL4", 11.0, 10.9, "1 steps")]
    report = levels_proof.build_report(PROFILE, [schedule(tmp_path / "a.xlsx", rows)])
    assert [(m["plot"], m["sheetSteps"], m["computedSteps"]) for m in report["stepsMismatches"]] == [("1", 2, 1)]
    assert report["summary"]["stepsMatches"] == 1 and report["summary"]["countRuleAloneMatches"] == 1


def test_r67_the_input_folder_is_the_environment_variable_or_the_project_setting(monkeypatch):
    monkeypatch.delenv(levels_proof.ROOT_VARIABLE, raising=False)
    assert levels_proof.input_root(SETTINGS) == Path(SETTINGS.get("inputRoot")) and str(levels_proof.input_root(SETTINGS)).endswith("02-Site Adaptation")
    monkeypatch.setenv(levels_proof.ROOT_VARIABLE, "somewhere else")
    assert levels_proof.input_root(SETTINGS) == Path("somewhere else")
    assert [path.name for path in levels_proof.schedule_paths("root", SETTINGS)] == SCHEDULES
    assert all(path.parent == Path("root") / full_run.XLSX_FOLDER for path in levels_proof.schedule_paths("root", SETTINGS))


def books_for(folder):
    imc = schedule(folder / SCHEDULES[0], [(1011, "DP2", 10.3, 10.0, "2 step"), (1012, "DP2", 10.3, 10.1, "1 steps"), (1013, "VL4", 11.1, 11.0, "1 steps")])
    imd = schedule(folder / SCHEDULES[1], [(2021, "DP4", 20.2, 20.0, "1 steps"), (2022, "DP4", 20.4, 20.2, "1 steps")])
    return [imc, imd]


CAD_CLUSTERS = [
    {"code": "M02", "plots": ["1001", "1002"]},
    {"code": "partly", "plots": ["1003", "1004"]},
    {"code": "M04", "plots": ["2010", "2011"]},
    {"code": "elsewhere", "plots": ["7001", "7002"]},
]


def test_r67_the_rule_fed_with_cad_clusters_can_fail_on_which_plots_belong_together(tmp_path):
    report = levels_proof.build_report(PROFILE, books_for(tmp_path), CAD_CLUSTERS, SHIFTS)
    # The sheet FFL grouping agrees with itself in every case, the CAD clusters do not: the sheet splits the M04 pair in two FFLs
    assert (report["summary"]["fflMatches"], report["summary"]["stepsMatches"], report["summary"]["plots"]) == (5, 5, 5)
    found = report["summary"]["cadClusters"]
    assert (found["plots"], found["inAnalysisClusters"], found["clustersPartlyInTheSheet"]) == (5, 5, 1)
    assert (found["fflMatches"], found["stepsMatches"]) == (4, 4)
    assert found["clustersWhollyInTheSheet"] == {"2": {"clusters": 2, "withMoreThanOneSheetFfl": 1}}
    assert levels_proof.cad_line(report) == "CAD clusters: FFL matches 4 of 5, steps matches 4 of 5, clusters with more than one sheet FFL: 2 units 1 of 2"


def test_r67_a_sheet_plot_enters_the_cad_run_under_its_baseplan_label_and_a_bad_mapping_is_refused(tmp_path):
    imc = schedule(tmp_path / SCHEDULES[0], [(1011, "DP2", 10.3, 10.0, "2 step")])
    imd = schedule(tmp_path / SCHEDULES[1], [(1012, "DP2", 10.3, 10.0, "2 step")])
    with pytest.raises(project_profile.ProfileError, match="two sheet plots end on the baseplan label 1001"):
        levels_proof.build_report(PROFILE, [imc, imd], [], SHIFTS)
    other = schedule(tmp_path / "other.xlsx", [(1011, "DP2", 10.3, 10.0, "2 step")])
    with pytest.raises(project_profile.ProfileError, match="no plot number shift for the schedule other.xlsx"):
        levels_proof.build_report(PROFILE, [other], [], SHIFTS)
    assert levels_proof.shift_of(imc, SHIFTS) == 10 and levels_proof.shift_of(imd, SHIFTS) == 11


def write_analysis(path, clusters):
    path.write_text(json.dumps({"clusters": clusters}), encoding="utf-8")
    return path


def test_r67_the_proof_command_reads_the_two_schedules_in_place_and_writes_its_report_outside_git(tmp_path, capsys):
    folder = tmp_path / "input" / full_run.XLSX_FOLDER
    folder.mkdir(parents=True)
    schedule(folder / SCHEDULES[0], ROWS[:3])
    schedule(folder / SCHEDULES[1], ROWS[3:])
    assert levels_proof.main([str(tmp_path / "out"), "--root", str(tmp_path / "input")]) == 0
    lines = capsys.readouterr().out.strip().splitlines()
    assert lines[0] == "FFL matches 5 of 6, steps matches 6 of 6, both 5, SSL 5, count rule alone 6"
    assert lines[1].startswith("base height source: past job schedules")
    assert lines[2] == "FFL mismatches 1, steps mismatches 0, largest FFL difference 0.05 m"
    assert lines[3] == "FFL plot 1004 VL4 sheet 30.05 computed 30.0 (-0.05)"
    report = json.loads((tmp_path / "out" / "levels_proof.json").read_text(encoding="utf-8"))
    assert report["summary"]["files"] == SCHEDULES and report["summary"]["cadClusters"] is None


def test_r67_the_proof_command_adds_the_run_on_the_cad_clusters_when_it_gets_analysis_json(tmp_path, capsys):
    folder = tmp_path / "input" / full_run.XLSX_FOLDER
    folder.mkdir(parents=True)
    books_for(folder)
    analysis = write_analysis(tmp_path / "analysis.json", CAD_CLUSTERS)
    assert levels_proof.main([str(tmp_path / "out"), "--root", str(tmp_path / "input"), "--analysis", str(analysis)]) == 0
    lines = capsys.readouterr().out.strip().splitlines()
    assert lines[3] == "CAD clusters: FFL matches 4 of 5, steps matches 4 of 5, clusters with more than one sheet FFL: 2 units 1 of 2"
    report = json.loads((tmp_path / "out" / "levels_proof.json").read_text(encoding="utf-8"))
    assert report["summary"]["cadClusters"]["fflMatches"] == 4


@pytest.mark.parametrize("what", ["no schedule", "no sheet", "no column", "no number", "no analysis", "bad analysis", "no clusters"])
def test_r67_the_proof_command_says_what_is_wrong_in_one_line(tmp_path, capsys, what):
    folder = tmp_path / "input" / full_run.XLSX_FOLDER
    folder.mkdir(parents=True)
    extra = []
    if what != "no schedule":
        for name in SCHEDULES:
            path = schedule(folder / name, ROWS[:2])
            if what == "no sheet":
                book = Workbook()
                book.active.title = "Other"
                book.save(path)
            elif what == "no column":
                book = Workbook()
                book.active.title = full_run.DEFAULT_SHEET
                book.active.append(["Plot No.", "Unit"])
                book.active.append([1, "DT1"])
                book.save(path)
            elif what == "no number":
                book = Workbook()
                book.active.title = full_run.DEFAULT_SHEET
                book.active.append(HEADERS)
                book.active.append(["1-DT1", 1, "DT1", "tbd", 1.0, 1.0, 0, "x", "1 steps"])
                book.save(path)
    if what == "no analysis":
        extra = ["--analysis", str(tmp_path / "missing.json")]
    elif what == "bad analysis":
        (tmp_path / "bad.json").write_text("{not json", encoding="utf-8")
        extra = ["--analysis", str(tmp_path / "bad.json")]
    elif what == "no clusters":
        extra = ["--analysis", str(write_analysis(tmp_path / "empty.json", None))]
    assert levels_proof.main([str(tmp_path / "out"), "--root", str(tmp_path / "input")] + extra) == 2
    error = capsys.readouterr().err
    assert error.startswith("levels_proof: ") and "Traceback" not in error and error.count("\n") == 1
    assert not (tmp_path / "out").exists()


def test_r67_proof_on_the_real_schedules_reports_matches_out_of_747_and_changes_no_rule(capsys):
    paths = levels_proof.schedule_paths(levels_proof.input_root(SETTINGS), SETTINGS)
    if not all(path.is_file() for path in paths):
        pytest.skip("the schedules are not in " + str(levels_proof.input_root(SETTINGS)))
    report = levels_proof.build_report(PROFILE, paths)
    summary = report["summary"]
    with capsys.disabled():
        print("\nR67 proof: " + levels_proof.summary_line(report))
    assert summary["plots"] == 747
    assert summary["fflMatches"] + len(report["fflMismatches"]) == 747
    assert summary["stepsMatches"] + len(report["stepsMismatches"]) == 747
    assert all(m["computedFfl"] is not None for m in report["fflMismatches"])
    # The step count rule on its own, from the sheet's own FFL, is the rule of the sheet
    assert summary["countRuleAloneMatches"] == 747
    # The Excel unit codes are the profile codes as they are
    assert summary["unitsNotInProfile"] == []
