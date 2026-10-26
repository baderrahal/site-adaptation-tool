"""R347, R350, R388. The cells, texts and settings of the reads of button 2 where the C# went apart from openpyxl and the Python readers
with no failure, found in the review of PR 53, on made up workbooks (R10). dotnet/SiteAdaptation.Tests/ReadEdgeTests.cs holds the C#
side of the same cases."""
import datetime
import json
import zipfile
from pathlib import Path

import openpyxl
import pytest

import area_key
import c3d_run
import cad_units
import corner_front
import parity
import project_profile
import read_steps
import rev9_workbooks
from test_read_steps import made_settings, made_workbooks

ROOT = Path(__file__).resolve().parents[2]
MAIN = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
KIND = "application/vnd.openxmlformats-officedocument.spreadsheetml."
RELATION = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def raw_book(path, sheet_xml):
    """A workbook of one sheet S whose sheet XML is written as given, so a cell of type d or a dimension out of date is what the file holds."""
    with zipfile.ZipFile(path, "w") as book:
        book.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                      '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
                      '<Override PartName="/xl/workbook.xml" ContentType="' + KIND + 'sheet.main+xml"/>'
                      '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="' + KIND + 'worksheet+xml"/></Types>')
        book.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                      '<Relationship Id="rId1" Type="' + RELATION + '/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        book.writestr("xl/workbook.xml", '<?xml version="1.0" encoding="UTF-8"?><workbook ' + MAIN + ' xmlns:r="' + RELATION + '">'
                      '<sheets><sheet name="S" sheetId="1" r:id="rId1"/></sheets></workbook>')
        book.writestr("xl/_rels/workbook.xml.rels", '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                      '<Relationship Id="rId1" Type="' + RELATION + '/worksheet" Target="worksheets/sheet1.xml"/></Relationships>')
        book.writestr("xl/worksheets/sheet1.xml", '<?xml version="1.0" encoding="UTF-8"?><worksheet ' + MAIN + '>' + sheet_xml + '</worksheet>')
    return path


def text(ref, value):
    return '<c r="%s" t="inlineStr"><is><t>%s</t></is></c>' % (ref, value)


def cell(ref, value, kind=None):
    return '<c r="%s"%s><v>%s</v></c>' % (ref, ' t="%s"' % kind if kind else "", value)


COLUMNS = {"plot": "Plot No", "nbh": "NBH"}


def test_r350_a_date_cell_of_type_d_reads_as_openpyxl_from_iso8601_gives_it(tmp_path):
    values = ["2024-03-05", "12:30:00", "2024-03-05T07:08:09.5Z", "2024-03-05T07:08", "2024-03-05T00:00:00+03:00"]
    rows = '<row r="1">%s%s</row>' % (text("A1", "Plot No"), text("B1", "NBH"))
    rows += "".join('<row r="%d">%s%s</row>' % (index, cell("A%d" % index, 100 + index), cell("B%d" % index, value, "d")) for index, value in enumerate(values, 2))
    path = raw_book(tmp_path / "dates.xlsx", '<dimension ref="A1:B6"/><sheetData>' + rows + "</sheetData>")
    found = [row["nbh"] for row in rev9_workbooks.sheet_rows(path, "S", 1, COLUMNS)]
    # The texts ReadEdgeTests.cs holds for the C# reader, str() and isoformat() of what openpyxl gives
    assert [str(value) for value in found] == ["2024-03-05", "12:30:00", "2024-03-05 07:08:09.500000", "2024-03-05 07:08:00", "2024-03-05 00:00:00"]
    assert [read_steps.jsonable(value) for value in found] == ["2024-03-05", "12:30:00", "2024-03-05T07:08:09.500000", "2024-03-05T07:08:00", "2024-03-05T00:00:00"]
    # A duration openpyxl gives as a timedelta is refused by the step, as the C# refuses the cell
    with pytest.raises(ValueError, match="a duration"):
        read_steps.jsonable({"nbh": datetime.timedelta(hours=12)})


def test_r350_a_date_in_a_key_column_gives_the_text_python_str_gives():
    day = datetime.datetime(2024, 3, 5)
    found, _ = area_key.by_plot([{"plot": 103, "nbh": day, "zone": 4}])
    assert found["103"] == ("Zone 4", "2024-03-05 00:00:00")
    assert read_steps.setback_key(day) == "2024-03-05 00:00:00"
    assert read_steps.setback_key(datetime.date(2024, 3, 5)) == "2024-03-05"


def test_r350_a_zone_or_plot_digit_is_an_ascii_digit_on_both_sides():
    assert area_key.zone_text("Zone ٤") == "Zone ٤"
    assert area_key.zone_text("Zone 04") == "Zone 4"
    found, _ = area_key.by_plot([{"plot": "١٠١", "zone": 4, "nbh": "012"}, {"plot": "101", "zone": 4, "nbh": "012"}])
    assert list(found) == ["101"]


def stale(tmp_path, dimension="A1:B1", extra=""):
    rows = '<row r="1">%s%s</row><row r="2">%s%s%s</row>' % (text("A1", "Plot No"), text("B1", "NBH"), cell("A2", 101), text("B2", "012"), extra)
    return raw_book(tmp_path / "stale.xlsx", '<dimension ref="%s"/><sheetData>%s</sheetData>' % (dimension, rows))


def test_r350_a_sheet_with_a_value_past_its_stated_dimension_is_refused(tmp_path):
    path = stale(tmp_path)
    # openpyxl alone reads to the stated dimension and gives the header row only
    assert len(list(openpyxl.load_workbook(path, read_only=True, data_only=True)["S"].iter_rows(values_only=True))) == 1
    with pytest.raises(ValueError, match="past its stated dimension A1:B1"):
        rev9_workbooks.sheet_rows(path, "S", 1, COLUMNS)
    with pytest.raises(ValueError, match="past its stated dimension A1:B1"):
        corner_front.sales_rows(path, "S", "Plot No")
    with pytest.raises(ValueError, match="C2"):
        rev9_workbooks.sheet_rows(stale(tmp_path, "A1:B2", text("C2", "x")), "S", 1, COLUMNS)
    # A row of no value past the dimension, a formatted empty cell, is no refusal
    formatted = raw_book(tmp_path / "formatted.xlsx", '<dimension ref="A1:B2"/><sheetData><row r="1">%s%s</row><row r="2">%s%s</row><row r="3"><c r="A3" s="0"/></row></sheetData>'
                         % (text("A1", "Plot No"), text("B1", "NBH"), cell("A2", 101), text("B2", "012")))
    assert rev9_workbooks.sheet_rows(formatted, "S", 1, COLUMNS) == [{"plot": 101, "nbh": "012"}]


def test_r350_read_steps_refuses_a_table_of_no_row(tmp_path, capsys):
    settings, values = made_settings(tmp_path)
    dlus, sales = made_workbooks(tmp_path, values)
    book = openpyxl.load_workbook(dlus)
    sheet = book[values["dataBook"]["data5"]["areaKey"]["dlus"]["sheet"]]
    sheet.delete_rows(2, 2)
    book.save(dlus)
    code = read_steps.main(["--settings", str(settings), "--steps", str(tmp_path / "steps"), "--only", "areaKey", "--dlus", str(dlus), "--sales", str(sales)])
    assert code == 2
    assert "holds no row" in capsys.readouterr().err
    assert not (tmp_path / "steps" / parity.step_name("areaKey", "python")).exists()


def test_r347_q083_the_config_names_the_peak_of_a_whole_read_with_its_reason():
    peaks = project_profile.load_settings().get("addin.workerPeakMiB")
    assert peaks["commands"] == {"read": 3900}
    assert "4,008" in peaks["commandsNote"] and "CHOSEN" in peaks["commandsNote"]


def test_r347_the_peak_of_a_whole_read_is_the_highest_measured_with_the_read_order_plus_the_largest_swing_measured_rounded_up():
    # The whole reads with the read order of Q084: 2,692 to 2,702 MiB on net48 and 2,550 to 2,554 on net8.0 in 4d, 2,464 to 2,493 and
    # 2,550 to 2,559 at the head of the review of PR 54. The largest swing of one build on one runtime over its done whole reads, 2,752 to
    # 3,899 MiB of 0a58496 on net48, where 2,800 left 98 MiB over the highest
    peaks = project_profile.load_settings().get("addin.workerPeakMiB")
    note = peaks["commandsNote"]
    measured = [2692, 2702, 2550, 2554, 2464, 2480, 2493, 2559]
    swing = 3899 - 2752
    assert peaks["commands"]["read"] == -(-(max(measured) + swing) // 100) * 100 == 3900
    assert "Q084" in note and "2,702" in note and "1,147" in note and "ReadOrder" in note and "5,948" in note
    assert "scaled by its file size" in project_profile.load_settings().get("addin.workerPeakNote")


def test_r350_every_workbook_setting_names_a_key_of_the_rev9_workbooks():
    book = c3d_run.load_config()["dataBook"]
    keys = set(book["rev9"]["workbooks"])
    for path in ("rev9.dlus", "rev9.sales", "data5.areaKey.dlus", "data5.areaKey.sales", "locationPlan.sales", "locationPlan.dlus"):
        entry = book
        for part in path.split("."):
            entry = entry[part]
        assert entry["workbook"] in keys, path


def test_r350_read_steps_takes_the_sales_rows_from_the_workbook_their_setting_names(tmp_path, capsys):
    settings, values = made_settings(tmp_path)
    dlus, sales = made_workbooks(tmp_path, values)
    values["dataBook"]["locationPlan"]["sales"]["workbook"] = "tbd"
    settings.write_text(json.dumps(values), encoding="utf-8")
    # The DLUS workbook holds no sales sheet, so the sales rows of the workbook tbd are refused, where the workbook of rev9.sales read before
    code = read_steps.main(["--settings", str(settings), "--steps", str(tmp_path / "steps"), "--only", "sales", "--dlus", str(dlus), "--sales", str(sales)])
    assert code == 2
    assert values["dataBook"]["locationPlan"]["sales"]["sheet"] in capsys.readouterr().err


def test_r350_parity_refuses_a_workbook_input_other_than_the_file_config_names(tmp_path):
    values = project_profile.load_settings().values
    names = values["dataBook"]["rev9"]["workbooks"]
    right = tmp_path / Path(names["accommodation"]).name
    wrong = tmp_path / Path(names["tbd"]).name
    for path in (right, wrong):
        path.write_text("a made up workbook", encoding="utf-8")
    payload = {"run": {"inputs": [{"role": "sales", "path": str(right), "sha256": "x"}]}}
    parity.workbook_inputs(payload, "python", values)
    payload["run"]["inputs"][0]["path"] = str(wrong)
    with pytest.raises(ValueError, match="the input sales of the python step file is %s" % wrong.name.replace("(", r"\(").replace(")", r"\)")):
        parity.workbook_inputs(payload, "python", values)


def test_r388_the_probe_settings_take_the_baseplan_and_the_fallback_reads_through_cad_reads():
    settings = project_profile.load_settings()
    merged = c3d_run.load_config()
    assert Path(merged["connect"]["baseplanPlots"]) == cad_units.named_input(settings, "baseplan")
    assert Path(merged["dataBook"]["join"]["fallbackOutputs"][0]) == cad_units.named_input(settings, "roads").parent
    # What still names the reads of Round 12 stays for a reason the note gives: the LandXML read, no CadReader read, and the folder
    # of the drawings cadReads does not hold
    raw = json.loads((ROOT / "python" / "c3d_probe.json").read_text(encoding="utf-8"))
    assert "baseplanPlots" not in raw["connect"]
    held = [key for key, value in (("connect.landxmlDir", raw["connect"]["landxmlDir"]),
                                   *(("dataBook.join.fallbackOutputs", item) for item in raw["dataBook"]["join"]["fallbackOutputs"])) if "round12" in value]
    assert held == ["connect.landxmlDir", "dataBook.join.fallbackOutputs"]
    assert "cadReads" in raw["connect"]["cadReadsNote"]
