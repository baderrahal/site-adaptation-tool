import csv
import json
import os
import time

import openpyxl
import pytest
from shapely.geometry import LineString, Point, Polygon

import c3d_run
import c3d_window
import data_book
import data_book_analysis
import data_book_border
import data_book_cad
import data_book_check
import data_book_excel
import data_book_readme

SQUARE = Polygon([(0, 0), (100, 0), (100, 100), (0, 100)])


def write_border(path, ring=((0, 0), (100, 0), (100, 100), (0, 100))):
    lines = ["x,y,vertex,area,perimeter,drawing,layer,handle"] + ["%s,%s,%d,1,1,d,l,h" % (x, y, k) for k, (x, y) in enumerate(ring)]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def book_config(tmp_path):
    root = tmp_path / "book"
    (root / "work" / "Base Plans").mkdir(parents=True)
    (tmp_path / "source").mkdir()
    return {"sourceRoot": str(tmp_path / "source"), "language": "en-US", "probeDll": str(tmp_path / "probe.dll"),
            "dataBook": {"root": str(root), "drawings": {"cutsheet": "Base Plans/cut.dwg"}, "acad": "acad.exe", "watcher": "watch.exe",
                         "minFreeGB": 2, "timeoutMinutes": 1, "knownSurfaces": {}}}


class Fake:
    def __init__(self, pid, polls=None):
        self.pid, self.returncode, self.polls, self.waited = pid, None, list(polls or []), False

    def poll(self):
        if self.polls:
            value = self.polls.pop(0)
            if value is not None:
                self.returncode = value
            return value
        return self.returncode

    def wait(self, timeout=None):
        self.waited = True
        self.returncode = 0


# R160 full window runner

def test_r160_refuses_source_outside_and_missing(tmp_path):
    config = book_config(tmp_path)
    (tmp_path / "source" / "a.dwg").write_bytes(b"x")
    assert "source folder" in c3d_window.refusal(config, tmp_path / "source" / "a.dwg")
    (tmp_path / "a.dwg").write_bytes(b"x")
    assert "not under" in c3d_window.refusal(config, tmp_path / "a.dwg")
    assert "missing" in c3d_window.refusal(config, tmp_path / "book" / "work" / "Base Plans" / "cut.dwg")


def test_r160_script_and_arguments():
    assert c3d_window.script_text("p.dll").splitlines() == ["NETLOAD", '"p.dll"', "C3DDATA", "C3DQUIT"]
    args = c3d_window.window_args("acad.exe", "a.dwg", "s.scr", "en-US")
    assert args[args.index("/product") + 1] == "C3D" and args[args.index("/b") + 1] == "s.scr" and args[1] == "a.dwg"
    assert c3d_window.watch_args("w.exe", 7, "p.dll", "o.json", 5)[1:4] == ["--pid", "7", "--probe"]


def test_r160_skips_when_memory_is_low_and_starts_nothing(tmp_path):
    config = book_config(tmp_path)
    (tmp_path / "book" / "work" / "Base Plans" / "cut.dwg").write_bytes(b"dwg")
    record = c3d_window.run(config, "cutsheet", launcher=lambda a, e: pytest.fail("started"), memory=lambda: 1.5)
    assert record["skipped"] and record["pids"] == []


def test_r160_stops_only_its_own_window_when_memory_falls(tmp_path, monkeypatch):
    config = book_config(tmp_path)
    drawing = tmp_path / "book" / "work" / "Base Plans" / "cut.dwg"
    drawing.write_bytes(b"dwg")
    stopped = []
    monkeypatch.setattr(c3d_run, "stop_own", lambda process: stopped.append(process.pid) or True)
    started = iter([Fake(100, polls=[None, None]), Fake(200, polls=[0])])
    memory = iter([4.0, 3.0, 1.5])
    record = c3d_window.run(config, "cutsheet", launcher=lambda a, e: next(started), memory=lambda: next(memory), sleep=lambda s: None)
    assert record["pids"] == [100, 200] and stopped == [100]
    assert "under 2 GB" in record["stopped"] and record["unchanged"] and record["done"] is False
    assert drawing.read_bytes() == b"dwg"


def test_r160_probe_done_needs_the_probe_reader_written_in_this_run(tmp_path):
    summary = tmp_path / "data.json"
    summary.write_text(json.dumps({"reader": "CadReader", "done": True}))
    assert not c3d_window.probe_done(summary, 0)
    summary.write_text(json.dumps({"reader": "C3DProbe", "done": True}))
    assert c3d_window.probe_done(summary, 0)
    assert not c3d_window.probe_done(summary, time.time() + 60)


# R163 border

def test_r163_rings_close_where_ends_meet():
    pieces = [{"kind": "polyline", "layer": "P1", "handle": "A", "pts": [[0, 0], [10, 0], [10, 10], [0, 0]]},
              {"kind": "line", "layer": "P1", "pts": [[0, 0], [1, 1]]}]
    rings = data_book_border.rings(pieces)
    assert len(rings) == 1 and rings[0]["points"] == [(0, 0), (10, 0), (10, 10)] and rings[0]["endGap"] == 0


def test_r163_pick_by_layer_and_label_else_the_one_that_holds_all():
    one = {"layer": "Phase 1 Bound", "handle": "A", "points": [(0, 0), (10, 0), (10, 10), (0, 10)], "closedFlag": False, "endGap": 0}
    two = {"layer": "Phase 2 Bound", "handle": "B", "points": [(20, 0), (30, 0), (30, 10), (20, 10)], "closedFlag": True, "endGap": 0}
    texts = [{"text": "PHASE - 1", "x": 5, "y": 5}]
    ring, how, _ = data_book_border.pick([one, two], texts, "Phase 1 Bound", "PHASE - 1", [(1, 1), (2, 2)])
    assert ring["handle"] == "A" and how == "layer and label"
    ring, how, _ = data_book_border.pick([one, two], [], "Other", "PHASE - 9", [(1, 1)])
    assert ring["handle"] == "A" and "holds every" in how
    with pytest.raises(data_book_border.BorderError):
        data_book_border.pick([one, two], [], "Other", "X", [(1, 1), (25, 5)])


def test_r163_write_border_holds_area_and_perimeter(tmp_path):
    ring = {"layer": "L", "handle": "H", "points": [(0, 0), (10, 0), (10, 10), (0, 10)]}
    polygon = data_book_border.write_border(tmp_path / "b.csv", ring, "d.dwg")
    rows = list(csv.DictReader(open(tmp_path / "b.csv", encoding="utf-8")))
    assert polygon.area == 100 and len(rows) == 4 and rows[0]["perimeter"] == "40.000" and rows[2]["vertex"] == "2"


# R164 stand-in tables

def test_r164_tables_keep_inside_cut_lines_and_count_outside(tmp_path):
    folder = tmp_path / "cad"
    folder.mkdir()
    (folder / "pieces.json").write_text(json.dumps({"pieces": [
        {"kind": "line", "layer": "RD", "handle": "1", "pts": [[50, 50, 170], [150, 50, 171]]},
        {"kind": "polyline", "layer": "RD", "handle": "2", "pts": [[10, 10], [20, 10], [20, 20]], "closed": True},
        {"kind": "line", "layer": "RD", "handle": "3", "pts": [[200, 200], [300, 200]]},
        {"kind": "circle", "layer": "C", "handle": "4", "center": [50, 80], "radius": 1},
        {"kind": "hatch", "layer": "H", "handle": "5", "loops": [[[1, 1], [3, 1], [3, 3], [1, 3]]]},
        {"kind": "insert", "layer": "B", "handle": "6", "attributes": {"NO": "7"}}]}), encoding="utf-8")
    (folder / "blocks.json").write_text(json.dumps({"blocks": [
        {"name": "x$0$Villa_VL4", "layer": "B", "handle": "6", "insertPoint": [5, 5, 0], "rotationDeg": 90, "mirrored": True, "outline": [[5, 5], [17, 5], [17, 30], [5, 30]]},
        {"name": "Far", "layer": "B", "handle": "7", "insertPoint": [500, 5, 0]}]}), encoding="utf-8")
    (folder / "texts.json").write_text(json.dumps({"texts": [{"text": "2311", "x": 1, "y": 1, "layer": "T"}, {"text": "out", "x": -5, "y": 1, "layer": "T"}]}), encoding="utf-8")
    (folder / "layers.json").write_text(json.dumps({"units": "Meters", "layers": [{"name": "RD"}, {"name": "EMPTY"}], "xrefs": []}), encoding="utf-8")
    out = tmp_path / "out"
    counts, units, _ = data_book_cad.tables(folder, SQUARE, out)
    assert counts["entities"] == {"keptCut": 1, "kept": 3, "outside": 1}
    assert counts["blocks"] == {"kept": 1, "outside": 1} and counts["texts"] == {"kept": 1, "outside": 1} and units == "Meters"
    entities = list(csv.DictReader(open(out / "entities.csv", encoding="utf-8")))
    cut = next(r for r in entities if r["handle"] == "1")
    assert cut["cut"] == "yes" and float(cut["x2"]) == pytest.approx(100) and float(cut["length"]) == pytest.approx(50)
    assert next(r for r in entities if r["handle"] == "2")["area"] == "50.0"
    blocks = list(csv.DictReader(open(out / "blocks.csv", encoding="utf-8")))
    assert blocks[0]["name"] == "Villa_VL4" and blocks[0]["attr_NO"] == "7" and float(blocks[0]["cx"]) == pytest.approx(11)
    layers = {r["layer"]: r for r in csv.DictReader(open(out / "layers.csv", encoding="utf-8"))}
    assert layers["RD"]["countsByType"] == "LINE=1 POLYLINE=1" and layers["EMPTY"]["entitiesInside"] == "0"


def test_r164_arc_points_and_split_by_layer(tmp_path):
    points = data_book_cad.arc_points([0, 0], 1, [1, 0], [0, 1])
    assert points[0] == pytest.approx((1, 0, 0)) and points[-1] == pytest.approx((0, 1, 0), abs=1e-9)
    folder = tmp_path / "drawing"
    folder.mkdir()
    with open(folder / "entities.csv", "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["handle", "layer"])
        for k in range(200):
            writer.writerow([k, "A" if k % 2 else "B/C"])
    assert data_book.split_large(folder, 0.001) == ["entities.csv"]
    assert not (folder / "entities.csv").exists()
    parts = sorted(p.name for p in (folder / "entities").glob("*.csv"))
    assert parts and all(n.startswith(("A", "B_C")) for n in parts)
    assert sum(len(list(csv.DictReader(open(p, encoding="utf-8")))) for p in (folder / "entities").glob("*.csv")) == 200


# R165 workbook

def test_r165_every_sheet_as_csv_with_a_sheet_list(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Phase1"
    sheet.append([None, "Title"])
    sheet.append([None, "Land-use", "Typology", "Plot Count"])
    sheet.append([None, None, "Villa (C10)", 280])
    other = book.create_sheet("Plots")
    other.append(["Plot No.", "Easting", "Northing"])
    book.save(tmp_path / "w.xlsx")
    listing = data_book_excel.write_sheets(tmp_path / "w.xlsx", tmp_path / "out")
    assert [e["sheet"] for e in listing] == ["Phase1", "Plots"]
    assert listing[0]["headerRow"] == 2 and listing[0]["plotOrLocationColumns"] == ""
    assert "Plot No." in listing[1]["plotOrLocationColumns"]
    rows = list(csv.reader(open(tmp_path / "out" / "Phase1.csv", encoding="utf-8")))
    assert rows[0] == ["A", "B", "C", "D"] and rows[3][3] == "280"
    assert (tmp_path / "out" / "sheets.csv").is_file()


# R166 join and README

def test_r166_bearing_depth_cell_and_points_in_face():
    assert data_book.bearing((0, 1)) == 0 and data_book.bearing((1, 0)) == 90
    face = Polygon([(0, 0), (10, 0), (10, 20), (0, 20)])
    assert data_book.depth_of(face, {"n": (0, 1), "left": (0, 0)}) == 20
    assert data_book.cell(True) == "yes" and data_book.cell(None) == "" and data_book.cell(1.234567) == "1.2346"
    from shapely.strtree import STRtree
    items = [{"x": 5, "y": 5}, {"x": 50, "y": 5}]
    tree = STRtree([Point(i["x"], i["y"]) for i in items])
    assert data_book.in_face(items, face, tree) == [items[0]]


def test_r166_readme_lists_files_rows_and_columns(tmp_path):
    (tmp_path / "plots.csv").write_text("officialNumber,level_Top Surface-02_left,attr_NO\n1,170,7\n", encoding="utf-8")
    (tmp_path / "work").mkdir()
    (tmp_path / "work" / "ignored.csv").write_text("a\n", encoding="utf-8")
    data_book_readme.write(tmp_path)
    text = (tmp_path / "README.md").read_text(encoding="utf-8")
    assert "- plots.csv (1 rows):" in text and "ignored" not in text
    assert "Level of surface Top Surface-02 at the left front corner" in text and "Block attribute NO." in text


# R167 analysis

def row(**values):
    base = {k: "" for k in data_book.BASE_COLUMNS}
    base.update(values)
    return base


def test_r167_conflicts_widths_setback_and_types():
    analysis = {"unitPlotTypes": ["DP2", "C10"], "widthsByClientType": {"DP2": [10, 12]}, "widthToleranceM": 0.1, "setbackM": {"default": 5.52, "C10": 5.5}, "setbackToleranceM": 0.05}
    rows = [row(baseplanLabel="1", clientType="DP2", widthM="12.0", setbackFromFrontM="5.0", footprintBlock="Duplex_DP2", v9PlotNumber="1"),
            row(baseplanLabel="2", clientType="C10", widthM="16.0", setbackFromFrontM="5.5", footprintBlock="Villa_VL4", officialNumber="12", v9PlotNumber="12"),
            row(baseplanLabel="3", clientType="OS")]
    result = data_book_analysis.conflicts(rows, {"DP2", "C10"}, analysis)
    assert result["typeBaseplanAgainst_footprintBlock"] == {"sameType": 1, "otherType VL4": 1}
    assert result["widthAgainstRev6"]["DP2"] == {"matches": 1} and result["widthAgainstRev6"]["C10"] == {"notChecked": 1}
    assert result["setbackAgainstRev6"] == {"differs": 1, "within": 1} and result["unitPlotsInPhase1"] == 2


def test_r167_v9_is_reference_and_vl4_width_comes_from_cad():
    analysis = {"unitPlotTypes": ["VL4", "VL3"], "widthsByClientType": {"VL3": [12, 16]}, "widthFromCad": ["VL4"], "widthToleranceM": 0.1,
                "setbackM": {"default": 5.52}, "setbackToleranceM": 0.05}
    rows = [row(baseplanLabel="1", clientType="VL4", widthM="15.0", v9PlotNumber="9"),
            row(baseplanLabel="2", clientType="VL3", widthM="20.0", officialNumber="12", v9PlotNumber="12"),
            row(baseplanLabel="3", clientType="OS")]
    result = data_book_analysis.conflicts(rows, {"VL4", "VL3"}, analysis)
    assert "plotNumberV9" not in result and "VL4" not in result["widthAgainstRev6"]
    assert result["widthFromCad"] == {"VL4": 1} and result["widthAgainstRev6"]["VL3"] == {"differs 20.0": 1}
    assert data_book_analysis.v9_reference(rows) == {"differs": 1, "equalsOfficialNumber": 1, "noV9Number": 1}


def test_r167_unit_source_entry_bands_and_coverage():
    analysis = {"unitPlotTypes": ["DP2"], "bandsM": [0.005, 0.03]}
    rows = [row(clientType="DP2", footprintBlock="Duplex_VD2", neighbourhood="IMC", pastEntry="170.0", roadLevelLeft="170.002", roadLevelRight="170.0"),
            row(clientType="DP2", neighbourhood="", pastEntry="", roadLevelLeft="")]
    rows[0]["unitRev6"] = "VD2"
    found = data_book_analysis.unit_source(rows, {"DP2", "VD2"}, analysis)
    assert found["unitPlotsWithoutUnit"] == 1 and found["codesOnlyInRev6"] == ["VD2"]
    rows[1]["footprintBlock"] = "Duplex_VD2"
    assert data_book_analysis.unit_source(rows, {"DP2", "VD2"}, analysis)["plotsWhereColumnNamesARev6OnlyCode"] == {"footprintBlock": 1}
    entry = data_book_analysis.entry_source(rows, analysis)
    assert entry["roadLevel"]["within5mm"] == 1
    cover = data_book_analysis.coverage(rows, ["footprintBlock", "pastEntry"])
    assert cover["share"]["pastEntry"] == {"IMC": 1.0, "rest": 0.0}


# R168 second method

def test_r168_ray_cast_area_centroid_and_edge_distance():
    ring = [(0, 0), (10, 0), (10, 10), (0, 10)]
    assert data_book_check.inside(ring, 5, 5) and not data_book_check.inside(ring, 15, 5)
    assert data_book_check.area(ring) == 100 and data_book_check.centroid(ring) == pytest.approx((5, 5))
    assert data_book_check.edge_distance(ring, 10.3, 5) == pytest.approx(0.3)


def test_r168_counts_read_split_tables(tmp_path):
    folder = tmp_path / "d"
    (folder / "entities").mkdir(parents=True)
    (folder / "entities" / "A.csv").write_text("handle,layer\n1,A\n2,A\n", encoding="utf-8")
    (folder / "texts.csv").write_text("handle,x,y\n1,1,1\n", encoding="utf-8")
    (folder / "blocks.csv").write_text("handle,x,y\n", encoding="utf-8")
    (folder / "data.json").write_text(json.dumps({"counts": {"entities": {"kept": 1, "keptCut": 1, "outside": 9}, "texts": {"kept": 1}}}), encoding="utf-8")
    assert data_book_check.check_counts(folder) == {"entities": "match", "blocks": "match", "texts": "match"}


def test_r168_plot_rows_worked_out_again(tmp_path):
    faces = {"1": [[(0, 0), (10, 0), (10, 10), (0, 10)]]}
    rows = [row(baseplanLabel="1", faceAreaM2="100.0", neighbourhood="IMC", footprintBlock="Villa_VL4", v9PlotNumber="77")]
    blocks = [{"name": "Villa_VL4", "x": 5, "y": 5}]
    v9 = [("77", 2, 2), ("78", 50, 50)]
    result = data_book_check.check_plots(rows, faces, [(-1, -1), (20, -1), (20, 20), (-1, 20)], {"IMC": (10, {"11"})}, blocks, v9, 14, 5)
    assert result == {"insideMatch": 1, "areaMatch": 1, "neighbourhoodMatch": 1, "blockMatch": 1, "v9Match": 1}


def test_r385_a_float_cell_of_plots_csv_is_written_by_the_rule_at_4_decimals():
    # plots.csv is read by the build and by the face step of the parity
    assert data_book.cell(12.34565) == "12.3457" and data_book.cell(7.00625) == "7.0063" and data_book.cell(0.50005) == "0.5001"
    assert repr(round(12.34565, 4)) == "12.3456"
    assert data_book.cell(-0.00004) == "0.0" and data_book.cell(3) == "3" and data_book.cell(1.234567) == "1.2346"
