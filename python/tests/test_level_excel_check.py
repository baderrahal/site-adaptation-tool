"""R364. The second route of the level Excel, python/level_excel_check.py, on small workbooks written here with openpyxl in the layout the
worker writes, from a made up plan, override list, step clusters, step levels and plots.csv, the same as LevelExcelTests of C# (R10).
Each planted fault is caught by its own check."""
import hashlib
import json

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill

import level_excel_check as check


def levels(road, entry, ffl, ssl, height, steps):
    return {key: {"value": value} for key, value in (("ffl", ffl), ("ssl", ssl), ("entry", entry), ("height", height), ("steps", steps), ("roadAtFrontM", road))}


def plot(number, unit, lv, rule_ffl, surfaces, *applied):
    item = {"plot": number, "unit": {"value": unit}, "levels": lv, "levelRule": {"ffl": rule_ffl, "steps": 1}, "roadSurfaces": surfaces}
    if applied:
        item["overrides"] = {"applied": list(applied), "notApplied": {}}
    return item


def entry(number, field, value, reason, source="parsons"):
    return {"plot": number, "field": field, "value": value, "source": source, "model": "model_%s.rvt" % source, "key": "k1", "reason": reason}


A = ["Road A", "Road A"]
PLAN = {"neighbourhood": "all", "summary": {"roadSurfacesOfFrontCorners": {"Road A": 5, "none": 1}, "noRoadLevel": [{"plot": "104", "corners": ["left"]}],
                                            "rules": {"levelProfile": "parsons2024Rev6", "overrideSource": ["parsons"], "overrideModel": []}},
        "plots": [plot("101", "VD4", levels(9.9012, 9.9234, 10.123, 10.041, 0.2101, 1), 10.122, ["Road A", None], "heightM", "parameter CASE"),
                  plot("102", "DP1", levels(10.3801, 10.4, 10.5, 10.43, 0.11, 1), 10.501, ["Road A", "Road B"], "entryM", "rampFrontM"),
                  plot("103", "DP1", levels(10.2857, 10.3115, 10.5, 10.42, 0.19, 2), 10.499, A, "steps"),
                  plot("201", "VD4", levels(11.1, 11.2, 11.4, 11.3, 0.2, 1), 11.4, ["Road C", "Road C"]),
                  plot("999", "VD4", levels(12.1, 12.2, 12.4, 12.3, 0.2, 1), 12.4, A)],
        "missing": [{"plot": "104", "written": False, "missing": [{"key": "levels.ffl", "why": "no road level"}]},
                    {"plot": "106", "written": False, "missing": [{"key": "levels.ffl", "why": "no road level"}]}]}
OVERRIDES = {"overrides": [entry("101", "heightM", {"gapToRuleM": 0.0006}, "breaks oneFflPerPieceAtTheHighestEntry"),
                           entry("101", "heightM", {"gapToRuleM": 0.5}, "native", "native"), entry("101", "parameter CASE", "D", "no level field"),
                           entry("102", "entryM", {"gapToRuleM": 0.0123}, "breaks roadLevel"), entry("103", "steps", "2", "breaks stepsFromHeight"),
                           entry("102", "rampFrontM", {"gapToRuleM": -0.002}, "breaks roadLevel")]}
CLUSTERS = {"plots": {key: {"written": {"plots": key.split("|")}} for key in ("102|103|108", "104|105")}}
STEP_LEVELS = {"plots": {number: {"written": {"piece": piece}} for number, piece in (("101", ["101"]), ("102", ["102"]), ("103", ["108", "103"]), ("201", ["201"]), ("999", ["999"]))}}
HANDS = {"plots": {number: {"written": {"cluster": ["108", "103", "102"]}} for number in ("102", "103")}}
# The split of the units the build takes, the typologyRev9 of plots.csv in the row order of the hands, VL1 DP1 DP1 cut after the second unit
RULES = {"units_mm": True, "levels": {"ground_ssl": -75}, "levelProfiles": {"parsons2024Rev6": {"frontRoad": {"surface": "roadLevel"},
         "keyRules": {"ffl": {"rule": "oneFflPerPieceAtTheHighestEntry"}, "levelSplit": {"rule": "units", "table": {"VL1 DP1 DP1": "walls 2"}}}}}}
CONFIG = {"siteBuild": {"roadSurfaces": ["Road A", "Road B"]}, "placerCheck": {"level": {"passMm": 1.0}}}
PLOTS = ("locationPlanNumber,zone,nbhId,typologyRev9\n101,Zone 2,012,VD4\n102,Zone 2,012,DP1\n103,Zone 2,012,DP1\n104|106,Zone 2,012,DT1\n105,Zone 3,012,DT1\n"
         "108,Zone 2,,VL1\n201,Zone 3,013,VD4\n")
FILL = PatternFill(fill_type="solid", fgColor="FFDDEBF7")


def inputs_of(paths):
    return check.load_inputs(paths["plan"], paths["overrides"], paths["clusters"], paths["levels"], paths["hands"], paths["plots"], paths["rules"], paths["config"])


@pytest.fixture
def made(tmp_path):
    paths = {}
    for name, content in (("plan", PLAN), ("overrides", OVERRIDES), ("clusters", CLUSTERS), ("levels", STEP_LEVELS), ("hands", HANDS), ("rules", RULES), ("config", CONFIG)):
        paths[name] = tmp_path / (name + ".json")
        paths[name].write_text(json.dumps(content), encoding="utf-8")
    paths["plots"] = tmp_path / "plots.csv"
    paths["plots"].write_text(PLOTS, encoding="utf-8")
    mock = check.mock_labels()
    return paths, mock, check.expected(inputs_of(paths), "parsons2024Rev6", mock)


def write_book(path, want, mock, change=None):
    """A workbook as the worker lays it out, change applied to it before it is saved."""
    book = openpyxl.Workbook()
    levels = book.active
    levels.title = mock["sheets"][0]
    levels.append([want["topLine"] + want["formulas"]])
    levels["A1"].font = Font(bold=True)
    levels.merge_cells("A1:J1")
    levels.append(mock["columns"])
    for column, name in enumerate(mock["columns"], 1):
        if name in check.EDITABLE:
            levels.cell(row=2, column=column).fill = FILL
    for row in want["levels"]:
        levels.append(row)
        for column, name in enumerate(mock["columns"], 1):
            if name in check.PLAN_KEYS:
                levels.cell(row=levels.max_row, column=column).number_format = check.number_format(mock["decimals"].get(name, 0))
    for name, header, rows in ((mock["sheets"][1], check.CLUSTER_COLUMNS, want["clusters"]), (mock["sheets"][2], check.OVERRIDE_COLUMNS, want["used"])):
        sheet = book.create_sheet(name)
        sheet.append(list(header))
        for row in rows:
            sheet.append(row)
    settings = book.create_sheet(mock["sheets"][3])
    settings.append(["Setting", "Value"])
    values = dict(want["settings"], **{"Run id": "r1", "Run time": "2026-10-28T10:00:00Z"})
    for name in check.SETTINGS:
        settings.append([name, values[name]])
    if change:
        change(book)
    book.save(path)


def write_folder(folder, want, mock, change=None):
    folder.mkdir(exist_ok=True)
    for nbh, book in want["books"].items():
        write_book(folder / book["file"], book, mock, change if nbh == "012" else None)


def test_r364_the_mock_labels_read_from_both_mock_files_agree_on_the_ten_columns_the_four_sheets_and_the_decimals():
    mock = check.mock_labels()
    assert mock["columns"] == ["Plot No", "Unit", "Cluster", "Road at front", "El @ Entry", "FFL", "SSL", "Height", "Steps", "Source"]
    assert mock["sheets"] == ["Levels", "Clusters", "Overrides used", "Settings"]
    assert mock["file"] == "<NBH>_Levels_<profile>.xlsx"
    assert mock["decimals"] == {"Road at front": 4, "El @ Entry": 4, "Height": 4, "FFL": 3, "SSL": 3}
    assert mock["formulas"].startswith("El @ Entry = the road at the two front corners") and "SSL = FFL - 0.075 m" in mock["formulas"]


def test_r364_one_workbook_per_nbh_and_the_cells_worked_out_again_from_the_plan_the_list_the_clusters_and_the_units(made):
    paths, _, want = made
    assert sorted(want["books"]) == ["012", "013"] and want["gaps"]["plots with no nbhId"] == 1
    book = want["books"]["012"]
    # SSL and Height as the plan writes them, not as the FFL would give them
    assert book["levels"][0] == ["101", "VD4", "none", 9.9012, 9.9234, 10.123, 10.041, 0.2101, 1, "Override heightM (breaks oneFflPerPieceAtTheHighestEntry)"]
    assert book["levels"][1][-1] == "Override entryM (breaks roadLevel), rampFrontM (breaks roadLevel)"
    # Review 2 of PR 91: 104 and 106 take the NBH of the one joined row 104|106 of plots.csv by area_key.key_of
    assert book["levels"][3] == ["104", None, "104-105", None, None, None, None, None, None, "Not written: levels.ffl (no road level)"]
    assert book["levels"][4] == ["106", None, "none", None, None, None, None, None, None, "Not written: levels.ffl (no road level)"]
    # Review 2 of PR 91: the cluster named in its row order on both sheets
    assert book["levels"][1][2] == "108-103-102" and book["levels"][2][2] == "108-103-102"
    # The pieces cut here by the typologyRev9 of plots.csv and keyRules.levelSplit, the same as the step levels give them
    assert book["clusters"] == [["108-103-102", "108-103", "outside the scope DP1 DP1", "walls 2", 10.499, "10.500"],
                                ["108-103-102", "102", "outside the scope DP1 DP1", "walls 2", 10.501, "10.500"]]
    assert want["gaps"]["cluster plots whose piece is the step piece"] == 2 and "cluster plots whose piece is not the step piece" not in want["gaps"]
    assert [row[:3] for row in book["used"]] == [["101", "heightM", 0.0006], ["102", "entryM", 0.0123], ["102", "rampFrontM", -0.002], ["103", "steps", 1]]
    # Review 2 of PR 91: the corners of the written plots of the NBH only, one of 101 with no surface, never the corner of 104 not written
    assert book["topLine"] == "NBH 012, Zone 2, road surface Road A on 4, Road B on 1 of 6 front corners, front corners with no road surface 1, formulas: "
    assert want["books"]["013"]["topLine"] == "NBH 013, Zone 3, road surface Road C on 2 of 2 front corners, formulas: "
    assert book["settings"]["Plan SHA-256"] == hashlib.sha256(paths["plan"].read_bytes()).hexdigest()
    assert book["settings"]["Plan build id"] == check.NO_BUILD_ID


def test_r364_the_build_id_is_held_against_the_run_stamp_of_the_plan(made):
    paths, mock, _ = made
    paths["plan"].write_text(json.dumps(dict(PLAN, run={"runId": "build_7"})), encoding="utf-8")
    assert check.expected(inputs_of(paths), "parsons2024Rev6", mock)["books"]["012"]["settings"]["Plan build id"] == "build_7"


def test_r364_workbooks_as_the_worker_lays_them_out_hold_every_check(made, tmp_path):
    _, mock, want = made
    write_folder(tmp_path / "out", want, mock)
    result, bad = check.check_folder(tmp_path / "out", want, mock)
    assert bad == [] and sum(count for key, count in result.items() if key.startswith("match ")) > 150


def setter(where, value):
    return lambda book: setattr(where(book), "value", value)


# Each fault and the check that must catch it
FAULTS = {
    "label": (setter(lambda book: book["Levels"]["D2"], "Road level"), "mismatch columns"),
    "cell": (setter(lambda book: book["Levels"]["F3"], 10.124), "mismatch Levels FFL"),
    "ssl": (setter(lambda book: book["Levels"]["G3"], 10.048), "mismatch Levels SSL"),
    "source": (setter(lambda book: book["Levels"]["J5"], "Override heightM"), "mismatch Levels Source"),
    "key": (setter(lambda book: book["Levels"]["A4"], "101"), "mismatch Plot No key unique"),
    "fill": (lambda book: setattr(book["Levels"]["H2"], "fill", FILL), "mismatch editable fill"),
    "unfill": (lambda book: setattr(book["Levels"]["I2"], "fill", PatternFill()), "mismatch editable fill"),
    "merge": (lambda book: book["Levels"].merge_cells("A3:B3"), "mismatch merged only the top line"),
    "blank": (lambda book: book["Levels"].insert_rows(4), "mismatch Levels rows"),
    "format": (lambda book: setattr(book["Levels"]["F3"], "number_format", "0.0000"), "mismatch format FFL"),
    "sha": (setter(lambda book: book["Settings"]["B11"], "0" * 64), "mismatch Settings Plan SHA-256"),
    "build": (setter(lambda book: book["Settings"]["B10"], "r1"), "mismatch Settings Plan build id"),
    "editable": (setter(lambda book: book["Settings"]["B9"], "FFL, Steps"), "mismatch Settings Editable columns"),
    "order": (lambda book: book.move_sheet("Settings", offset=-1), "mismatch sheets"),
    "piece": (setter(lambda book: book["Clusters"]["B2"], "103-108"), "mismatch Clusters Plots"),
    "split": (setter(lambda book: book["Clusters"]["D2"], "none"), "mismatch Clusters Split"),
    "pieceffl": (setter(lambda book: book["Clusters"]["E2"], 10.5), "mismatch Clusters FFL by rule"),
    "aswritten": (setter(lambda book: book["Clusters"]["F2"], "10.499"), "mismatch Clusters FFL as written"),
    "clusterheader": (setter(lambda book: book["Clusters"]["E1"], "FFL"), "mismatch Clusters header"),
    "sequence": (setter(lambda book: book["Clusters"]["C2"], "DP1 DP1 outside the scope"), "mismatch Clusters Unit sequence"),
    "gap": (setter(lambda book: book["Overrides used"]["C4"], "2"), "mismatch Overrides used gap a number"),
    "none": (setter(lambda book: book["Levels"]["A1"], "NBH 012, Zone 2, road surface Road A on 5, none on 1 of 6 front corners, formulas: "), "mismatch top line prefix"),
}


@pytest.mark.parametrize("fault", sorted(FAULTS))
def test_r364_each_planted_fault_is_a_mismatch_of_its_own_check(made, tmp_path, fault):
    _, mock, want = made
    change, key = FAULTS[fault]
    write_folder(tmp_path / "out", want, mock, change)
    result, bad = check.check_folder(tmp_path / "out", want, mock)
    assert result[key] >= 1, (fault, sorted(name for name in result if name.startswith("mismatch")))


def test_r364_a_missing_or_extra_workbook_is_a_mismatch(made, tmp_path):
    _, mock, want = made
    write_folder(tmp_path / "out", want, mock)
    (tmp_path / "out" / "013_Levels_parsons2024Rev6.xlsx").rename(tmp_path / "out" / "all_Levels_parsons2024Rev6.xlsx")
    result, _ = check.check_folder(tmp_path / "out", want, mock)
    assert result["mismatch workbooks"] == 1


def test_r364_the_plan_side_counts_a_level_override_the_list_lacks_and_a_piece_off_the_step(made):
    paths, mock, _ = made
    paths["overrides"].write_text(json.dumps(dict(OVERRIDES, overrides=OVERRIDES["overrides"][:5])), encoding="utf-8")
    paths["levels"].write_text(json.dumps({"plots": dict(STEP_LEVELS["plots"], **{"102": {"written": {"piece": ["103", "102"]}}})}), encoding="utf-8")
    want = check.expected(inputs_of(paths), "parsons2024Rev6", mock)
    assert want["gaps"]["level overrides with no entries"] == 1 and want["gaps"]["cluster plots whose piece is not the step piece"] == 1
    assert want["books"]["012"]["levels"][1][-1] == "Override entryM (breaks roadLevel), rampFrontM (no entry in the override list)"


def test_r364_ffl_as_written_lists_the_written_ffls_an_override_moved_off_the_rule_in_row_order_and_is_empty_when_none_moved(made):
    # Q122: 108 written at 10.52 and 103 at 10.5 against their rule 10.499, two values in the row order of the piece, 102 at its rule, none
    paths, mock, _ = made
    plots = [dict(item, levels=dict(item["levels"], ffl={"value": 10.501})) if item["plot"] == "102" else item for item in PLAN["plots"]]
    plots.append(plot("108", "VL1", levels(10.3, 10.31, 10.52, 10.445, 0.21, 2), 10.499, A))
    paths["plan"].write_text(json.dumps(dict(PLAN, plots=plots)), encoding="utf-8")
    want = check.expected(inputs_of(paths), "parsons2024Rev6", mock)
    assert [row[4:] for row in want["books"]["012"]["clusters"]] == [[10.499, "10.520, 10.500"], [10.501, None]]
    assert check.CLUSTER_COLUMNS[4:] == ("FFL by rule", "FFL as written")
