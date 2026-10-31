"""R365. The second route of the report Excel the worker command reportExcel writes in C# (CadReader ReportWorkbook): each xlsx read with
openpyxl, every label and its order held against the words of docs/addin-mock.md item 6 and docs/design/07-mock.md 7.6 with items 9 to 12 of
7.7, read from those files and never from the C#, and every cell held against the plan, its override list, the step checkRows of the build and
plots.csv, worked out again here by its own walk with json and csv only.

- workbooks: one per NBH of the plan by nbhId of plots.csv, each named <NBH>_Report.xlsx, no other xlsx in the folder
- labels: the sheets of 7.6 with Assumed after Overrides (7.7 item 9) and Settings last (item 12), the Summary table and its steps 1 to 8 by
  the order of docs/design/13-workflow-changes.md with the column Manual rule of 7.6 (R381), the two boxes with their columns and items, the
  columns of Placed with Mirror after type (item 11) and of Skipped
- R381: the manual rule IDs of each Summary and Issues line from the block manualRules of the rules profile, the faces of plots.csv in the
  scope for 2 Read data (Q123), one Issues row per weak wall decision (Q125) and the Revit column from the counts the add-in read (Q124)
- Summary: the top line as the filled example of 7.6 words it, the table from row 3 with its counts walked again from the plan, the box
  Counts agree across outputs with the plan held against the rows of Placed, the box Rules applied with the plots that took each rule
  counted here from the status details, the boxes bordered and the share in its format
- Placed: plot in column A, one row per written plot and kind, the villa type and mirror flag of a link, no blank row, no merged range,
  the headers of type and Mirror alone in the light fill (D4, items 10 and 11)
- Skipped, Overrides, Assumed, Issues, Rules: each row worked out again from the plan, the list, the status and the check rows
- Settings: the editable columns against the fills, the build id against the run stamp of the plan, the SHA-256 of the plan file bytes by
  hashlib, the SHA-256 of the rules and the list

With no --xlsx-folder the plan side runs alone. Counts, plot numbers and row numbers only, no coordinate or level, to --out outside the repo.
"""
import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path

import area_key
import check_rows
import project_profile
import value_status
from level_excel_check import NO_BUILD_ID, listed, section, short

ROOT = Path(__file__).resolve().parent.parent
LINK = "link"
FLOOR_KINDS = ("generic", "yard", "ramp")
# The words the mock does not give, CHOSEN in Round ADDIN 4g3 PR 5, held here word for word against what the worker writes
NOT_RUN = "not run yet"
NOT_PLACED = "not placed yet"
NO_LEVEL_RUN = "no run of button 6 yet"
NOT_MEASURED = "not measured"
RULE_COLUMNS = ("Decision", "Rule", "Share", "Strength", "Plots")
OVERRIDE_COLUMNS = ("Plot", "Field", "Outcome", "Why", "Reason", "Source", "Model")
ISSUE_COLUMNS = ("Step", "Plot", "Check", "Field", "Source", "Use", "Why")
RULE_SHEET_COLUMNS = ("Rule", "Status", "Basis", "Plots", "Values")
BASIS = "Basis"
# R381, CHOSEN in 5e: the words of a line with no block manualRules or no rule of its step, the row of a weak wall decision (Q125)
NO_MANUAL_BLOCK = "no manual rule block"
NO_MANUAL_RULE = "no manual rule"
WEAK_CHECK = "Weak rule"
BAD_MANUAL = ("the block manualRules of the rules profile holds no list rules each with a text id and a list checkedIn of the steps that check it, a report "
              "line names the manual rules of its step by them (R381)")
REVIT_FORMAT = "siteAdaptationRevitCounts"
NOT_COUNTS = ("the input revitCounts is no file of the counts the add-in read from the model, format " + REVIT_FORMAT
              + " with plots by number, each with links, floors by kind and level (R381)")
NO_REPORT_BLOCK = ("the settings hold no block addin.report with stepOfField, defaultStep and wallDecisions, the report counts the overrides of each "
                   "button and the weak wall decisions by it (R365)")
SETTINGS = ("Run id", "Profile", "Rules SHA-256", "Override list SHA-256", "Run time", "Editable columns", "Plan build id", "Plan SHA-256")
SHARE_FORMAT = "0.000"


def title(word):
    return word[:1].upper() + word[1:]


def mock_labels(root=ROOT):
    """The labels of the report Excel by the words of the two mock files, each read on its own and held against the other, with the
    additions of items 9 to 12 of 7.7."""
    lines = (root / "docs" / "addin-mock.md").read_text(encoding="utf-8").splitlines()
    item = next(line for line in lines if line.startswith("6. Report Excel"))
    design = (root / "docs" / "design" / "07-mock.md").read_text(encoding="utf-8")
    part = section(design, "### 7.6", "### 7.7")
    extra = section(design, "### 7.7", "### 7.8")
    # R381, J12: the steps of Bader's changes of 28 Oct 2026, which prevail over the ribbon of item 1, one row per step before 9 Report
    order = (root / "docs" / "design" / "13-workflow-changes.md").read_text(encoding="utf-8")
    nine = listed(re.search(r"^- Nine numbered steps: (.*?)\.$", order, re.M).group(1))
    steps = nine[:-1]
    boxes = re.findall(r'A box "(.*?)"', item)
    found = {
        "file": re.search(r"file (\S+\.xlsx)", item).group(1),
        "topLine": re.search(r"A top line with (.*?)\.", item).group(1),
        "stepColumns": listed(re.search(r"Summary table with columns (.*?), one row per button", item).group(1)),
        "countsBox": boxes[0], "rulesBox": boxes[1],
        "countColumns": listed(re.search(r"with columns (Item.*?) for ", item).group(1)),
        "countItems": listed(re.search(r" for (Links[^.]*)\.", item).group(1)),
        "sheets": listed(re.search(r"Sheets: (.*?)\.\s*$", item).group(1)),
        "steps": steps,
    }
    table = [line for line in part.splitlines() if line.startswith("| ")]
    rows = lambda head: [[cell.strip() for cell in line.strip("|").split("|")] for line in table if line.startswith("| " + head)]
    mine = {
        "file": re.search(r"file (\S+\.xlsx)", part).group(1),
        "topLine": re.search(r"^A top line with (.*?)\.$", part, re.M).group(1),
        "stepColumns": rows("Step |")[0][:5],
        "countsBox": re.search(r'^A box "(.*?)" with columns', part, re.M).group(1),
        "rulesBox": re.findall(r'^A box "(.*?)"', part, re.M)[1],
        "countColumns": rows("Item |")[0],
        "countItems": [row[0] for line_head in found["countItems"] for row in rows(line_head + " |")],
        "sheets": listed(re.search(r"^Sheets: (.*?)\.$", part, re.M).group(1)),
        "steps": [row[0] for row in rows("") if re.match(r"[1-9] ", row[0])],
    }
    disagree = [key for key in found if found[key] != mine[key]]
    if disagree:
        raise ValueError("docs/addin-mock.md item 6 and 7.6 of docs/design/07-mock.md differ on " + ", ".join(disagree))
    if any(row[1] != NOT_PLACED for name in found["countItems"] for row in rows(name + " |")):
        raise ValueError("7.6 no longer writes the Revit count as " + NOT_PLACED)
    filled = re.search(r'Filled: "(\[file name\]\.rvt, NBH .*?)"', part).group(1)
    items = {int(match.group(1)): match.group(2) for match in re.finditer(r"^(\d+)\. (.*)$", extra, re.M)}
    if "the sheet after Overrides" not in items[9] or "after its sheets of 7.6 and the sheet Assumed" not in items[12]:
        raise ValueError("items 9 and 12 of 7.7 no longer place the sheets Assumed and Settings")
    sheets = list(found["sheets"])
    sheets.insert(sheets.index("Overrides") + 1, "Assumed")
    sheets.append("Settings")
    placed = [title(word) for word in listed(re.search(r"Placed: one row per placed element, (.*?)\.", part).group(1).replace("the element id", "element id"))]
    mirror = re.search(r"The column (\w+) of the sheet Placed, after (\w+)", items[11])
    placed.insert(placed.index(title(mirror.group(2))) + 1, mirror.group(1))
    editable = [title(word) for word in re.search(r"Steps of Levels and of (\w+) and (\w+) of Placed", items[10]).groups()]
    manual = re.search(r'the column "(.*?)" after Issues', part).group(1)
    found.update(
        stepColumns=found["stepColumns"] + [manual], issueColumns=list(ISSUE_COLUMNS) + [manual], reportStep=nine[-1],
        sheets=sheets, placed=placed, editable=editable,
        skipped=[title(word) for word in listed(re.search(r"Skipped: (.*?)\.", part).group(1))],
        assumed=[title(word) for word in re.split(r", | and ", re.search(r"its columns (.*?) and the Q ID", items[9]).group(1))] + [BASIS],
        plotCount=re.search(r'the count written "(.*?)"', items[9]).group(1),
        filled=re.sub(r"NBH \S+, [\d,]+ plots, \S+, ", "NBH %(nbh)s, %(n)d plots, %(profile)s, ", filled))
    return found


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def basis(entry):
    """The Q ID, the share or the source of one entry: its question, else its rule with its share, else the rule check_rows.rule_name names."""
    question = entry.get("question")
    if isinstance(question, str) and question:
        return question
    if is_number(entry.get("share")):
        return "%s share %s" % (check_rows.rule_name(entry), short(entry["share"]))
    return check_rows.rule_name(entry)


def cell(value):
    """A value of the plan as the one cell it fills: a list or an object as its JSON, read back and compared as JSON."""
    return value if value is None or isinstance(value, (str, bool, int, float)) else ("json", value)


def load_inputs(plan_path, overrides_path, check_rows_path, summary_path, plots_csv, rules_path, config_path, revit_path=None):
    read = lambda path: json.loads(Path(path).read_text(encoding="utf-8"))
    with open(plots_csv, encoding="utf-8-sig", newline="") as handle:
        table = [row for row in csv.DictReader(handle)]
    return {"plan": read(plan_path), "planBytes": Path(plan_path).read_bytes(), "overrides": read(overrides_path), "checkRows": read(check_rows_path), "summary": read(summary_path),
            "table": table, "rows": {row["locationPlanNumber"]: row for row in table if row.get("locationPlanNumber")}, "rules": read(rules_path),
            "config": read(config_path), "revit": read(revit_path) if revit_path else None,
            "sha": {name: hashlib.sha256(Path(path).read_bytes()).hexdigest() for name, path in (("rules", rules_path), ("overrides", overrides_path))}}


def report_block(config):
    """Review 1 of PR 96: addin.report of the settings, refused when it lacks the button of each field, the default button or the wall decisions."""
    block = (config.get("addin") or {}).get("report")
    if not isinstance(block, dict) or not isinstance(block.get("stepOfField"), dict) or not isinstance(block.get("defaultStep"), int) \
            or not isinstance(block.get("wallDecisions"), list):
        raise ValueError(NO_REPORT_BLOCK)
    return block


def step_of(field, block):
    """The step of the Summary table an override field counts under, by addin.report as rows 5 to 7 of 7.6 list the fields (R381)."""
    return block["stepOfField"].get(field, block["defaultStep"])


def manual_block(rules):
    """R381, W7: the list rules of the block manualRules of the rules profile (R403) as (id, checkedIn) pairs in its order, None when the
    profile holds no block."""
    block = rules.get("manualRules")
    if block is None:
        return None
    entries = block.get("rules") if isinstance(block, dict) else None
    if not isinstance(entries, list) or any(not isinstance(item, dict) or not isinstance(item.get("id"), str) or not isinstance(item.get("checkedIn"), list)
                                            or any(not isinstance(step, str) for step in item["checkedIn"]) for item in entries):
        raise ValueError(BAD_MANUAL)
    return [(item["id"], set(item["checkedIn"])) for item in entries]


def manual_of(block, step):
    if block is None:
        return NO_MANUAL_BLOCK
    return ", ".join(name for name, steps in block if step in steps) or NO_MANUAL_RULE


def revit_counts(counts):
    """R381, Q124: the counts the add-in read from the model, refused when they are not of the form RevitCounts of Common writes."""
    plots = counts.get("plots") if isinstance(counts, dict) and counts.get("format") == REVIT_FORMAT else None
    good = lambda item: isinstance(item, dict) and type(item.get("links")) is int and isinstance(item.get("level"), bool) \
        and isinstance(item.get("floors"), dict) and all(type(value) is int for value in item["floors"].values())
    if not isinstance(plots, dict) or not all(good(item) for item in plots.values()):
        raise ValueError(NOT_COUNTS)
    return plots


def counts_box(placed, written, mock, revit=None, manual=NO_MANUAL_BLOCK):
    """The box Counts agree across outputs: Revit the totals the add-in read or not placed yet, Excel the rows of the Placed sheet, and one Issues
    row per item whose plan, Revit and Excel counts differ on any two, a text count compared with none."""
    links = sum(1 for row in placed if row[1] == LINK)
    floors = sum(1 for row in placed if row[1] in FLOOR_KINDS)
    shown = list(revit) if revit is not None else [NOT_PLACED] * 3
    items = [(mock["countItems"][0], len(written), shown[0], links), (mock["countItems"][1], sum(len(plot.get("floors") or {}) for plot in written), shown[1], floors),
             (mock["countItems"][2], len(written), shown[2], NO_LEVEL_RUN)]
    rows = [[item, revit_count, excel] for item, held, revit_count, excel in items]
    issues = []
    for item, held, revit_count, excel in items:
        if len({value for value in (held, revit_count, excel) if is_number(value)}) > 1:
            issues.append([mock["reportStep"], None, "Counts", item, None, None, "the plan holds %s, Revit %s, Excel %s (%s)" % (held, revit_count, excel, "R365" if revit is None else "R381"),
                           manual])
    return rows, issues


def expected(inputs, profile_name, mock):
    """The workbooks of the plan by NBH, each with every cell it must hold worked out here, and the gaps found on the way."""
    plan, rows, config, gaps = inputs["plan"], inputs["rows"], inputs["config"], Counter()
    rules = (plan.get("summary") or {}).get("rules") or {}
    sources, models = set(rules.get("overrideSource") or []), set(rules.get("overrideModel") or [])
    entries = [item for item in inputs["overrides"].get("overrides") or [] if item.get("source") in sources and (not models or item.get("model") in models)]
    first = {}
    for item in entries:
        first.setdefault((str(item.get("plot")), item.get("field")), item)
    details = plan.get("statusFrom") or {}
    status = config.get("valueStatus") or {}
    assumed, ranked = status.get("assumed") or [], status.get("worstFirst") or []
    written = plan.get("plots") or []
    not_written = [item for item in plan.get("missing") or [] if item.get("written") is False]
    numbers = [plot["plot"] for plot in written] + [item["plot"] for item in not_written]
    for number in numbers:
        if number not in rows:
            joined = [key for key in rows if "|" in key and number in key.split("|")]
            found, _ = area_key.key_of(joined, {key: (rows[key].get("zone") or "", rows[key].get("nbhId") or "") for key in joined})
            if found is not None:
                rows[number] = {"zone": found[0], "nbhId": found[1]}
    nbh_of = {number: (rows.get(number) or {}).get("nbhId") or None for number in numbers}
    gaps["plots with no nbhId"] = sum(1 for nbh in nbh_of.values() if nbh is None)
    key_rules = inputs["rules"]["levelProfiles"][profile_name].get("keyRules") or {}
    prefix = "%s levelProfiles.%s.keyRules." % (config.get("profile") or "", profile_name)
    skipped_list = set((plan.get("summary") or {}).get("skipped") or [])
    step_plots = inputs["checkRows"].get("plots") or {}
    block = report_block(config)
    manual = manual_block(inputs["rules"])
    revit = revit_counts(inputs["revit"]) if inputs.get("revit") is not None else None
    scope = {number: item.get("written") or item.get("raw") or {} for number, item in (inputs["summary"].get("plots") or {}).items()}
    scope_kind = inputs["summary"].get("scope")
    books = {}
    for nbh in sorted({nbh for nbh in nbh_of.values() if nbh}):
        mine_written = [plot for plot in written if nbh_of[plot["plot"]] == nbh]
        mine_missing = [item for item in plan.get("missing") or [] if nbh_of.get(item["plot"]) == nbh]
        mine_not = [item for item in not_written if nbh_of[item["plot"]] == nbh]
        names = {plot["plot"] for plot in mine_written}
        order = [plot["plot"] for plot in mine_written] + [item["plot"] for item in mine_not]
        walls = [wall for wall in plan.get("walls") or [] if wall.get("plot") in names]
        placed = []
        for plot in mine_written:
            placed.append([plot["plot"], LINK, (plot.get("unit") or {}).get("value"), (plot.get("mirrored") or {}).get("value"), None])
            for kind, floor in (plot.get("floors") or {}).items():
                placed.append([plot["plot"], kind, ((floor or {}).get("type") or {}).get("value"), None, None])
        skipped = []
        for item in mine_missing:
            lacks = item.get("missing") or []
            if item.get("written") is False:
                why = ", ".join("%s (%s)" % (found.get("key"), found.get("why")) for found in lacks)
                skipped.extend([item["plot"], kind, why] for kind in (LINK,) + FLOOR_KINDS)
            else:
                skipped.extend([item["plot"], found["key"][len("floors."):] if found["key"].startswith("floors.") else found["key"], found.get("why")] for found in lacks)
        overrides = []
        for number, record in [(plot["plot"], plot.get("overrides")) for plot in mine_written] + [(item["plot"], item.get("overrides")) for item in mine_not]:
            for field in (record or {}).get("applied") or []:
                entry = first.get((number, field)) or {}
                overrides.append([number, field, "Applied", None, entry.get("reason"), entry.get("source"), entry.get("model")])
            for field, why in ((record or {}).get("notApplied") or {}).items():
                entry = first.get((number, field)) or {}
                overrides.append([number, field, "Left", cell(why), entry.get("reason"), entry.get("source"), entry.get("model")])
        assumed_rows, values = [], []
        for plot in mine_written:
            for path, value in value_status.fields_of(plot).items():
                detail = details.get(value.get("statusFrom")) or {}
                values.append((plot["plot"], detail))
                name = value_status.status_label(value.get("status"), detail)
                if name in assumed:
                    shown = list(dict.fromkeys(basis(entry) for entry in value_status.worst_entries(detail.get("entries") or [], ranked)))
                    assumed_rows.append([plot["plot"], path, cell(value.get("value")), name, " | ".join(shown)])
        for wall in walls:
            values.extend((wall["plot"], details.get(value.get("statusFrom")) or {}) for value in wall.values() if isinstance(value, dict) and "value" in value and "source" in value)
        issues = []
        for number in order:
            step = step_plots.get(number) or {}
            for row in (step.get("written") or step.get("raw") or {}).get("rows") or []:
                issues.append([mock["steps"][3], number, row.get("check"), row.get("field"), row.get("from"), row.get("use"), cell(row.get("why")),
                               manual_of(manual, mock["steps"][3])])
        grouped, took = {}, {}
        for number, detail in values:
            for entry in detail.get("entries") or []:
                key = (check_rows.rule_name(entry) or "", value_status.entry_label(entry))
                group = grouped.setdefault(key, [basis(entry), set(), 0])
                group[1].add(number)
                group[2] += 1
            # Review 1 of PR 96: a value whose detail names an override took the value of the override, never that of a keyRules entry
            if detail.get("overrides"):
                continue
            for path in {entry.get("entry") for entry in detail.get("entries") or [] if entry.get("entry")}:
                took.setdefault(path, set()).add(number)
        rule_rows = [[rule, name, group[0], len(group[1]), group[2]] for (rule, name), group in sorted(grouped.items())]
        box = []
        for decision, item in key_rules.items():
            if not isinstance(item, dict) or not isinstance(item.get("rule"), str):
                continue
            share, strong = item.get("share"), item.get("strong")
            box.append([decision, item["rule"], share if is_number(share) else NOT_MEASURED,
                        NOT_MEASURED if not isinstance(strong, bool) else "strong" if strong else "weak",
                        mock["plotCount"].replace("[n]", str(len(took.get(prefix + decision + ".rule", ())))) ])
        totals = None
        if revit is not None:
            # The plots the add-in read whose nbhId of plots.csv is this NBH, a tagged plot outside the plan counted too
            here = [item for number, item in revit.items() if (rows.get(number) or {}).get("nbhId") == nbh]
            totals = [sum(item["links"] for item in here), sum(sum(item["floors"].values()) for item in here), sum(1 for item in here if item["level"])]
        placed_links = sum(1 for row in placed if row[1] == LINK)
        placed_floors = sum(1 for row in placed if row[1] in FLOOR_KINDS)
        counts, unequal = counts_box(placed, mine_written, mock, totals, manual_of(manual, mock["reportStep"]))
        gaps["count boxes with a mismatch"] += 1 if unequal else 0
        outcome = lambda step, word: sum(1 for row in overrides if row[2] == word and step_of(row[1], block) == step)
        moved = sum(1 for wall in walls if any((details.get(value.get("statusFrom")) or {}).get("overrides") for value in wall.values()
                                               if isinstance(value, dict) and "statusFrom" in value))
        # R381, Q123: 2 Read data over the faces of plots.csv of the NBH in the scope, every face for a build of the whole project or of the NBH,
        # else the faces of the plots the step summary lists, a face with a unit of typologyRev9 Done and the rest Skipped
        whole = scope_kind in ("project", "nbh=" + nbh)
        faces = [row for row in inputs["table"] if row.get("nbhId") == nbh and (whole or row.get("locationPlanNumber") in scope)]
        with_unit = sum(1 for row in faces if row.get("typologyRev9"))
        gaps["plots of the plan not in the step summary"] += sum(1 for number in order if number not in scope)
        weak = [name for name in block["wallDecisions"] if isinstance(key_rules.get(name), dict) and key_rules[name].get("strong") is False]
        checked = len(issues)
        for name in weak:
            share = key_rules[name].get("share")
            issues.append([mock["steps"][7], None, WEAK_CHECK, name, prefix + name + ".rule", None,
                           "%s share %s, weak" % (key_rules[name].get("rule") or "", short(share) if is_number(share) else NOT_MEASURED), manual_of(manual, mock["steps"][7])])
        not_run = [NOT_RUN] * 4
        steps = [not_run, [with_unit, len(faces) - with_unit, 0, 0], not_run,
                 [len(order), sum(1 for number in order if number in skipped_list), sum(1 for item in entries if str(item.get("plot")) in set(order)), checked],
                 [placed_floors, sum(1 for row in skipped if row[1] in FLOOR_KINDS), outcome(5, "Applied"), outcome(5, "Left")],
                 [len(mine_written), len(mine_not), outcome(6, "Applied"), outcome(6, "Left")],
                 [placed_links, sum(1 for row in skipped if row[1] == LINK), outcome(7, "Applied"), outcome(7, "Left")],
                 [len(walls), 0, moved, len(weak)]]
        steps = [[name] + values + [manual_of(manual, name)] for name, values in zip(mock["steps"], steps)]
        issues.extend(unequal)
        stamp = plan.get("run")
        books[nbh] = {"file": mock["file"].replace("<NBH>", nbh), "topLine": mock["filled"] % {"nbh": nbh, "n": len(order), "profile": profile_name},
                      "steps": steps, "counts": counts, "box": box,
                      "sheets": {"Placed": placed, "Skipped": skipped, "Overrides": overrides, "Assumed": assumed_rows, "Issues": issues, "Rules": rule_rows},
                      "settings": {"Profile": profile_name, "Rules SHA-256": inputs["sha"]["rules"], "Override list SHA-256": inputs["sha"]["overrides"],
                                   "Editable columns": ", ".join(name for name in mock["placed"] if name in mock["editable"]),
                                   "Plan build id": stamp["runId"] if isinstance(stamp, dict) and isinstance(stamp.get("runId"), str) else NO_BUILD_ID,
                                   "Plan SHA-256": hashlib.sha256(inputs["planBytes"]).hexdigest()}}
    return {"books": books, "gaps": gaps, "written": len(written), "notWritten": len(not_written)}


def same(found, wanted):
    """A cell against its value: a flag as a flag, a number as the same double, a list or an object by its JSON read back."""
    if isinstance(wanted, tuple) and wanted[:1] == ("json",):
        try:
            return isinstance(found, str) and json.loads(found) == wanted[1]
        except ValueError:
            return False
    if isinstance(wanted, bool) or isinstance(found, bool):
        return found is wanted
    if is_number(wanted) and is_number(found):
        return float(found) == float(wanted)
    return found == wanted


def row_values(sheet, row, width):
    return [sheet.cell(row=row, column=column).value for column in range(1, width + 1)]


def check_xlsx(path, want, mock, run_id=None):
    """Every label, layout, fill, border and cell of one written workbook against the mock and the plan: matches and mismatches by check."""
    import openpyxl
    book = openpyxl.load_workbook(path)
    result, bad = Counter(), []

    def hold(what, ok, where):
        result[("match " if ok else "mismatch ") + what] += 1
        if not ok:
            bad.append(where)

    def rows(sheet, start, wanted, header, what):
        for index, values in enumerate(wanted):
            found = row_values(sheet, start + index, len(header))
            for column, (got, value) in enumerate(zip(found, values)):
                hold("%s %s" % (what, header[column]), same(got, value), "%s row %d %s" % (what, start + index, header[column]))

    hold("sheets", book.sheetnames == mock["sheets"], "sheets %s" % book.sheetnames)
    if book.sheetnames != mock["sheets"]:
        return result, bad
    summary = book["Summary"]
    columns = mock["stepColumns"]
    hold("top line", summary.cell(row=1, column=1).value == want["topLine"], "Summary row 1")
    hold("top line bold", bool(summary.cell(row=1, column=1).font.bold), "Summary A1 not bold")
    merged = sorted(str(item) for item in summary.merged_cells.ranges)
    hold("merged only the top line", merged == ["A1:%s1" % openpyxl.utils.get_column_letter(len(columns))], "Summary merged %s" % merged)
    hold("row 2 blank", all(value is None for value in row_values(summary, 2, 8)), "Summary row 2")
    hold("step columns from row 3", row_values(summary, 3, len(columns)) == columns, "Summary row 3")
    rows(summary, 4, want["steps"], columns, "Summary")
    # R381: the rows after the step table move with its count of steps, one blank row before each box
    blank = 4 + len(want["steps"])
    counts_at, rules_at = blank + 1, blank + 7
    hold("row %d blank" % blank, all(value is None for value in row_values(summary, blank, 8)), "Summary row %d" % blank)
    hold("counts box title", summary.cell(row=counts_at, column=1).value == mock["countsBox"] and bool(summary.cell(row=counts_at, column=1).font.bold), "Summary row %d" % counts_at)
    hold("counts box columns", row_values(summary, counts_at + 1, 3) == mock["countColumns"], "Summary row %d" % (counts_at + 1))
    rows(summary, counts_at + 2, want["counts"], mock["countColumns"], "Counts")
    hold("row %d blank" % (rules_at - 1), all(value is None for value in row_values(summary, rules_at - 1, 8)), "Summary row %d" % (rules_at - 1))
    hold("rules box title", summary.cell(row=rules_at, column=1).value == mock["rulesBox"] and bool(summary.cell(row=rules_at, column=1).font.bold), "Summary row %d" % rules_at)
    hold("rules box columns", row_values(summary, rules_at + 1, len(RULE_COLUMNS)) == list(RULE_COLUMNS), "Summary row %d" % (rules_at + 1))
    rows(summary, rules_at + 2, want["box"], RULE_COLUMNS, "Rules applied")
    hold("Summary rows", summary.max_row == rules_at + 1 + len(want["box"]), "Summary max_row %d against %d" % (summary.max_row, rules_at + 1 + len(want["box"])))
    boxed = [(row, column) for row in range(counts_at + 1, counts_at + 5) for column in range(1, 4)] \
        + [(row, column) for row in range(rules_at + 1, rules_at + 2 + len(want["box"])) for column in range(1, 6)]
    thin = lambda side: side is not None and side.style == "thin"
    hold("boxes bordered", all(thin(summary.cell(row=row, column=column).border.left) and thin(summary.cell(row=row, column=column).border.bottom)
                               for row, column in boxed), "Summary box borders")
    hold("step table not bordered", not thin(summary.cell(row=4, column=1).border.left), "Summary row 4 bordered")
    for index, row in enumerate(want["box"]):
        if is_number(row[2]):
            hold("share format", summary.cell(row=rules_at + 2 + index, column=3).number_format == SHARE_FORMAT, "Rules applied row %d format" % (rules_at + 2 + index))
    headers = {"Placed": mock["placed"], "Skipped": mock["skipped"], "Overrides": list(OVERRIDE_COLUMNS), "Assumed": mock["assumed"],
               "Issues": mock["issueColumns"], "Rules": list(RULE_SHEET_COLUMNS)}
    for name, header in headers.items():
        sheet, wanted = book[name], want["sheets"][name]
        hold(name + " header", row_values(sheet, 1, len(header)) == header, name + " row 1")
        hold(name + " rows", sheet.max_row == 1 + len(wanted), "%s max_row %d against %d" % (name, sheet.max_row, 1 + len(wanted)))
        hold(name + " no merged range", not sheet.merged_cells.ranges, name + " merged")
        rows(sheet, 2, wanted, header, name)
    placed = book["Placed"]
    keys = [(placed.cell(row=row, column=1).value, placed.cell(row=row, column=2).value) for row in range(2, placed.max_row + 1)]
    hold("Placed plot key in column A", all(isinstance(key[0], str) and key[0] for key in keys), "Placed column A blank")
    hold("Placed one row per plot and kind", len(keys) == len(set(keys)), "Placed plot and kind repeated")
    hold("Placed mirror only on links", all((placed.cell(row=row, column=4).value is None) != (placed.cell(row=row, column=2).value == LINK)
                                           for row in range(2, placed.max_row + 1)), "Placed Mirror off a link or missing on one")
    fills = [placed.cell(row=1, column=column).value for column in range(1, len(mock["placed"]) + 1) if placed.cell(row=1, column=column).fill.fill_type == "solid"]
    hold("editable fill", fills == [name for name in mock["placed"] if name in mock["editable"]], "fills on %s" % fills)
    colours = {placed.cell(row=1, column=column).fill.fgColor.rgb for column in range(1, len(mock["placed"]) + 1) if placed.cell(row=1, column=column).fill.fill_type == "solid"}
    hold("one fill colour", len(colours) == 1, "fill colours %s" % sorted(colours))
    others = [(name, item.coordinate) for name in book.sheetnames for line in book[name].iter_rows() for item in line
              if item.fill.fill_type == "solid" and not (name == "Placed" and item.row == 1)]
    hold("no other fill", not others, "other fills %s" % others[:5])
    settings = book["Settings"]
    hold("Settings header", row_values(settings, 1, 2) == ["Setting", "Value"], "Settings row 1")
    names = [settings.cell(row=row, column=1).value for row in range(2, settings.max_row + 1)]
    hold("Settings names", names == list(SETTINGS), "Settings names %s" % names)
    values = {settings.cell(row=row, column=1).value: settings.cell(row=row, column=2).value for row in range(2, settings.max_row + 1)}
    for name, value in want["settings"].items():
        hold("Settings " + name, values.get(name) == value, "Settings " + name)
    if run_id is not None:
        hold("Settings Run id", values.get("Run id") == run_id, "Settings Run id")
    hold("Settings sha256 lowercase hex", bool(re.fullmatch(r"[0-9a-f]{64}", values.get("Plan SHA-256") or "")), "Settings Plan SHA-256 form")
    hold("Settings run time", bool(re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", values.get("Run time") or "")), "Settings Run time form")
    hold("Settings editable against the fills", values.get("Editable columns") == ", ".join(fills), "Settings Editable columns against the fills")
    return result, bad


def check_folder(folder, want, mock, run_id=None):
    """Each workbook the plan must give in the folder, and no other xlsx there."""
    result, bad = Counter(), []
    names = sorted(path.name for path in Path(folder).glob("*.xlsx"))
    wanted = sorted(book["file"] for book in want["books"].values())
    result["match workbooks" if names == wanted else "mismatch workbooks"] += 1
    if names != wanted:
        bad.append("workbooks %s against %s" % (names, wanted))
    for nbh, book in sorted(want["books"].items()):
        path = Path(folder) / book["file"]
        if path.exists():
            found, missed = check_xlsx(path, book, mock, run_id)
            result.update(found)
            bad.extend(nbh + ": " + item for item in missed)
    return result, bad


def run(args):
    mock = mock_labels()
    inputs = load_inputs(args.plan, args.overrides, args.check_rows, args.summary, args.plots_csv, args.rules, args.config, args.revit_counts)
    want = expected(inputs, args.profile, mock)
    books = want["books"]
    report = {"plan": Path(args.plan).name, "workbooks": sorted(book["file"] for book in books.values()),
              "rows": {"written": want["written"], "notWritten": want["notWritten"],
                       **{name: sum(len(book["sheets"][name]) for book in books.values()) for name in ("Placed", "Skipped", "Overrides", "Assumed", "Issues", "Rules")}},
              "gaps": dict(sorted(want["gaps"].items()))}
    if args.xlsx_folder:
        result, bad = check_folder(args.xlsx_folder, want, mock, args.run_id)
        report["xlsx"] = {"matches": sum(count for key, count in result.items() if key.startswith("match ")),
                          "mismatches": sum(count for key, count in result.items() if key.startswith("mismatch ")),
                          "byCheck": dict(sorted(result.items())), "first": bad[:20]}
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--plan", required=True)
    parser.add_argument("--overrides", required=True)
    parser.add_argument("--check-rows", required=True, help="the step checkRows file of the build, step_checkRows_python.json or its C# twin")
    parser.add_argument("--summary", required=True, help="the step summary file of the build, step_summary_python.json, the unit of each plot of the scope")
    parser.add_argument("--plots-csv", required=True)
    parser.add_argument("--rules", help="the rules profile, by default the one config names")
    parser.add_argument("--config", default=str(project_profile.SETTINGS_PATH))
    parser.add_argument("--profile", default="parsons2024Rev6")
    parser.add_argument("--revit-counts", help="the counts the add-in read from the model, the input revitCounts of the request when it gives one (R381)")
    parser.add_argument("--run-id", help="the run id of the reportExcel request, held against Settings when given")
    parser.add_argument("--xlsx-folder", help="the folder of the report Excel files the worker wrote of this plan, left out to check the plan side alone")
    parser.add_argument("--out", required=True, help="the report, outside the repo")
    args = parser.parse_args(argv)
    if args.rules is None:
        args.rules = str(project_profile.profile_path(project_profile.load_settings(Path(args.config))))
    report = run(args)
    Path(args.out).write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps({key: report[key] for key in ("plan", "rows", "gaps")}))
    if "xlsx" in report:
        print("xlsx matches %d mismatches %d" % (report["xlsx"]["matches"], report["xlsx"]["mismatches"]))
    return 1 if report.get("xlsx", {}).get("mismatches") else 0


if __name__ == "__main__":
    sys.exit(main())
