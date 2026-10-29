"""R364. The second route of the level Excel the worker command levelExcel writes in C# (CadReader LevelWorkbook): each xlsx read with
openpyxl, every label and its order held against the words of docs/addin-mock.md item 5 and docs/design/07-mock.md 7.5, read from those
files and never from the C#, and every cell held against the plan, its override list, the step clusters, the step levels and plots.csv,
worked out again here with json and csv only.

- workbooks: one per NBH of the plan by nbhId of plots.csv (7.5), each named <NBH>_Levels_<profile>.xlsx, no other xlsx in the folder
- labels: the sheets, the ten columns, the top line with its formulas as the filled example of 7.5 words them, the number format of each
  column by the decimals 7.5 gives
- Levels: Plot No the key in column A, one row per written plot of the NBH in the order of the plan then one per plot not written, no blank
  row and no merged range under the top line, each value equal to the plan value, the cluster of the step clusters, the source Rule or
  Override with the reason of the one entry of the list of the sources and models of the plan
- Clusters: the pieces worked out here by their own code, each cluster in the row order of the step hands cut where keyRules.levelSplit
  names a split for the typologyRev9 of its plots in plots.csv, the units the build takes, and those pieces held against the piece the
  step levels give each plot as well, FFL by rule the levelRule.ffl of the piece and FFL as written each written FFL of the piece that rounds
  off the rule of its plot at 3 decimals, once each in the row order, as text with 3 decimals, empty when none (Q122)
- Overrides used: each row worked out again, the gap a number for every field, the step count less the steps of the rule for steps
- Settings: the build id held against the run stamp of the plan, the SHA-256 of the plan file bytes by hashlib, the SHA-256 of the rules and
  the list, and the editable columns, which are the only headers with a fill (Bader, 2026-10-28)

With no --xlsx-folder the plan side runs alone on real plans. Counts, plot numbers and row numbers only, no level, to --out outside the repo.
"""
import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path

import area_key
import project_profile
from common import fixed

ROOT = Path(__file__).resolve().parent.parent
LEVEL_FIELDS = ("entryM", "heightM", "steps", "rampFrontM")
# Bader, 2026-10-28: the columns the team may edit for the button Update from report of Phase 5f
EDITABLE = ("El @ Entry", "FFL", "Steps")
SETTINGS = ("Run id", "Profile", "Rules SHA-256", "Override list SHA-256", "Road surface order", "Pass limit", "Run time", "Editable columns",
            "Plan build id", "Plan SHA-256")
# Q122 (Bader, 2026-10-29): the FFL of the rule of the piece and the written FFLs an override moved off it, new items of 7.7
CLUSTER_COLUMNS = ("Cluster", "Plots", "Unit sequence", "Split", "FFL by rule", "FFL as written")
OVERRIDE_COLUMNS = ("Plot No", "Field", "Gap", "Reason", "Source", "Model")
PLAN_KEYS = {"Road at front": "roadAtFrontM", "El @ Entry": "entry", "FFL": "ffl", "SSL": "ssl", "Height": "height", "Steps": "steps"}
PIECE_RULE = "oneFflPerPieceAtTheHighestEntry"
# The words the mock does not give, CHOSEN in review 1 of PR 91, held here word for word against what the worker writes
NO_BUILD_ID = "none, the plan holds no run stamp"
NO_SURFACE = ", front corners with no road surface %d"
OUTSIDE = "outside the scope"
NOT_WRITTEN = "not written"


def section(text, start, end):
    """The lines of a Markdown file from the heading start up to the heading end."""
    lines = text.splitlines()
    first = next(index for index, line in enumerate(lines) if line.startswith(start))
    last = next((index for index in range(first + 1, len(lines)) if lines[index].startswith(end)), len(lines))
    return "\n".join(lines[first:last])


def listed(text):
    """The names of a list of the mock, a part in brackets dropped, split on its commas."""
    return [part.strip() for part in re.sub(r"\s*\([^)]*\)", "", text).split(",") if part.strip()]


def mock_labels(root=ROOT):
    """The labels of the level Excel by the words of the two mock files, each read on its own and held against the other."""
    item = next(line for line in (root / "docs" / "addin-mock.md").read_text(encoding="utf-8").splitlines() if line.startswith("5. Level Excel"))
    part = section((root / "docs" / "design" / "07-mock.md").read_text(encoding="utf-8"), "### 7.5", "### 7.6")
    found = {
        "file": re.search(r"file (\S+\.xlsx)", item).group(1),
        "columns": listed(re.search(r"Columns in this order: (.*?)\. Sheets:", item).group(1)),
        "sheets": listed(re.search(r"Sheets: (.*?)\.\s*$", item).group(1)),
        "topLine": re.search(r"A top line with (.*?)\.", item).group(1),
    }
    table = next(line for line in part.splitlines() if line.startswith("| Plot No"))
    design = {
        "file": re.search(r"file (\S+\.xlsx)", part).group(1),
        "columns": listed(re.search(r"^Columns in this order: (.*?)\.$", part, re.M).group(1)),
        "table": [cell.strip() for cell in table.strip("|").split("|")],
        "sheets": listed(re.search(r"^Sheets: (.*?)\.$", part, re.M).group(1)),
        "topLine": re.search(r"^A top line with (.*?)\.$", part, re.M).group(1),
    }
    disagree = [key for key in ("file", "columns", "sheets", "topLine") if found[key] != design[key]]
    if design["table"] != found["columns"]:
        disagree.append("table")
    if disagree:
        raise ValueError("docs/addin-mock.md item 5 and 7.5 of docs/design/07-mock.md differ on " + ", ".join(disagree))
    decimals = {}
    match = re.search(r"(Road at front[^.]*?) with (\d) decimals as R2 keeps them, (\w+ and \w+) with (\d)", part)
    for names, count in ((match.group(1), match.group(2)), (match.group(3), match.group(4))):
        for name in re.split(r", | and ", names):
            decimals[name.strip()] = int(count)
    filled = re.search(r'Filled: "(NBH .*?)"', part).group(1)
    found.update(decimals=decimals, formulas=filled.split("formulas: ", 1)[1])
    return found


def number_format(decimals):
    return "0" if decimals == 0 else "0." + "0" * decimals


def short(value):
    return ("%.6f" % value).rstrip("0").rstrip(".")


def plan_build_id(plan):
    """The run id of the run stamp a C# build writes into its plan, else the words for none, never the run id of the Excel."""
    stamp = plan.get("run")
    return stamp["runId"] if isinstance(stamp, dict) and isinstance(stamp.get("runId"), str) else NO_BUILD_ID


def as_written(plots):
    """Q122. The written FFLs of a piece an override moved off the rule of their plot, by R385 at 3 decimals, once each in the row order."""
    moved = []
    for plot in plots:
        written, ruled = ((plot.get("levels") or {}).get("ffl") or {}).get("value"), (plot.get("levelRule") or {}).get("ffl")
        if isinstance(written, (int, float)) and isinstance(ruled, (int, float)) and fixed(written) != fixed(ruled):
            text = "%.3f" % fixed(written)
            if text not in moved:
                moved.append(text)
    return ", ".join(moved) or None


def whole(value):
    try:
        found = float(value)
    except (TypeError, ValueError):
        return None
    return int(found) if found.is_integer() else None


def load_inputs(plan_path, overrides_path, clusters_path, levels_path, hands_path, plots_csv, rules_path, config_path):
    read = lambda path: json.loads(Path(path).read_text(encoding="utf-8"))
    with open(plots_csv, encoding="utf-8-sig", newline="") as handle:
        rows = {row["locationPlanNumber"]: row for row in csv.DictReader(handle) if row.get("locationPlanNumber")}
    return {"plan": read(plan_path), "planBytes": Path(plan_path).read_bytes(), "overrides": read(overrides_path), "clusters": read(clusters_path),
            "levels": read(levels_path), "hands": read(hands_path), "rows": rows, "rules": read(rules_path), "config": read(config_path),
            "sha": {name: hashlib.sha256(Path(path).read_bytes()).hexdigest() for name, path in (("rules", rules_path), ("overrides", overrides_path))}}


def cut(cluster, pattern):
    """The walls a split pattern names in a cluster, read here on its own: none, middle of an even count, or walls n,m after unit n."""
    count = len(cluster)
    if pattern in (None, "none"):
        return []
    if pattern == "middle":
        return [count // 2] if count % 2 == 0 else []
    return sorted({int(part) for part in pattern.split(" ", 1)[1].split(",") if 0 < int(part) < count})


def split_label(walls, count):
    if not walls:
        return "none"
    return "middle" if count % 2 == 0 and walls == [count // 2] else "walls " + ",".join(str(wall) for wall in walls)


def expected(inputs, profile_name, mock):
    """The workbooks of the plan by NBH, each with the cells it must hold worked out here from the inputs, and the gaps found on the way."""
    plan, rows, gaps = inputs["plan"], inputs["rows"], Counter()
    rules = (plan.get("summary") or {}).get("rules") or {}
    sources, models = set(rules.get("overrideSource") or []), set(rules.get("overrideModel") or [])
    entries = {}
    for item in inputs["overrides"].get("overrides") or []:
        if item.get("field") in LEVEL_FIELDS and item.get("source") in sources and (not models or item.get("model") in models):
            entries.setdefault((str(item["plot"]), item["field"]), []).append(item)
    cluster_of = {}
    for item in (inputs["clusters"].get("plots") or {}).values():
        plots = (item.get("written") or item.get("raw") or {}).get("plots") or []
        if len(plots) >= 2:
            for plot in plots:
                cluster_of[plot] = plots
    row_of = {number: (item.get("written") or item.get("raw") or {}).get("cluster") for number, item in (inputs["hands"].get("plots") or {}).items()}
    step_piece = {number: (item.get("written") or item.get("raw") or {}).get("piece") for number, item in (inputs["levels"].get("plots") or {}).items()}
    written = plan.get("plots") or []
    not_written = [item for item in plan.get("missing") or [] if item.get("written") is False]
    by_number = {plot["plot"]: plot for plot in written}
    missing = {item["plot"] for item in not_written}
    # Review 2 of PR 91: a number with no row of its own takes the zone and nbhId of the joined rows whose parts hold it, by area_key.key_of
    for number in [plot["plot"] for plot in written] + list(missing):
        if number not in rows:
            joined = [key for key in inputs["rows"] if "|" in key and number in key.split("|")]
            found, _ = area_key.key_of(joined, {key: (inputs["rows"][key].get("zone") or "", inputs["rows"][key].get("nbhId") or "") for key in joined})
            if found is not None:
                rows[number] = {"zone": found[0], "nbhId": found[1]}
    nbh_of = {number: (rows.get(number) or {}).get("nbhId") or None for number in [plot["plot"] for plot in written] + list(missing)}
    gaps["plots with no nbhId"] = sum(1 for nbh in nbh_of.values() if nbh is None)
    profile = inputs["rules"]["levelProfiles"][profile_name]
    key = profile.get("keyRules") or {}
    rule = (key.get("ffl") or {}).get("rule") or profile.get("height")
    split = key.get("levelSplit") or {}
    table = (split.get("table") or {}) if rule == PIECE_RULE and split.get("rule") == "units" else {}
    ground = inputs["rules"]["levels"]["ground_ssl"] / 1000.0
    formulas = mock["formulas"].replace("FFL - 0.075 m", "FFL %s %s m" % ("-" if ground <= 0 else "+", short(abs(ground))))
    front = profile.get("frontRoad") or {}
    order = inputs["config"]["siteBuild"]["roadSurfaces"] if front.get("surface") == "roadLevel" else [front.get("surface")]
    passed = ((inputs["config"].get("placerCheck") or {}).get("level") or {}).get("passMm")
    # Review 2 of PR 91: each cluster in the row order the step hands give it, named so on both sheets, the sorted order when none does
    ordered = {}
    for cluster in {tuple(plots) for plots in cluster_of.values()}:
        row = next((row_of[member] for member in cluster if row_of.get(member) and sorted(row_of[member]) == sorted(cluster)), None)
        if row is None:
            gaps["clusters with no row order"] += 1
        for member in cluster:
            ordered[member] = list(row or cluster)
    books = {}
    for nbh in sorted({nbh for nbh in nbh_of.values() if nbh}):
        levels, used = [], []
        mine_written = [plot for plot in written if nbh_of[plot["plot"]] == nbh]
        for plot in mine_written:
            number = plot["plot"]
            applied = [field for field in (plot.get("overrides") or {}).get("applied") or [] if field in LEVEL_FIELDS]
            parts = []
            for field in applied:
                found = entries.get((number, field)) or []
                gaps["level overrides applied"] += 1
                gaps["level overrides with %s entries" % ("no" if not found else "one" if len(found) == 1 else "several")] += 1
                if not found:
                    parts.append("%s (no entry in the override list)" % field)
                    continue
                entry = found[0]
                parts.append("%s (%s)" % (field, entry.get("reason") or ""))
                value = entry.get("value")
                if field == "steps":
                    count, ruled = whole(value), whole((plot.get("levelRule") or {}).get("steps"))
                    gap = None if count is None or ruled is None else count - ruled
                else:
                    gap = value.get("gapToRuleM") if isinstance(value, dict) else None
                    gap = gap if isinstance(gap, (int, float)) and not isinstance(gap, bool) else None
                if gap is None:
                    gaps["overrides used with no gap"] += 1
                used.append([number, field, gap, entry.get("reason"), entry.get("source"), entry.get("model")])
            row = [number, (plot.get("unit") or {}).get("value"), "-".join(ordered[number]) if number in ordered else "none"]
            for column in ("Road at front", "El @ Entry", "FFL", "SSL", "Height", "Steps"):
                value = ((plot.get("levels") or {}).get(PLAN_KEYS[column]) or {}).get("value")
                if value is None:
                    gaps["written plots with no " + PLAN_KEYS[column]] += 1
                elif column != "Steps" and round(value, mock["decimals"][column]) != value:
                    gaps["values with more decimals than 7.5 gives, " + column] += 1
                row.append(value)
            row.append("Override " + ", ".join(parts) if parts else "Rule")
            levels.append(row)
        for item in not_written:
            if nbh_of[item["plot"]] != nbh:
                continue
            reasons = ", ".join("%s (%s)" % (found["key"], found["why"]) for found in item.get("missing") or [])
            levels.append([item["plot"], None, "-".join(ordered[item["plot"]]) if item["plot"] in ordered else "none", None, None, None, None, None, None, "Not written: " + reasons])
        pieces, seen = [], set()
        for plot in mine_written:
            cluster = cluster_of.get(plot["plot"])
            if not cluster or tuple(cluster) in seen:
                continue
            seen.add(tuple(cluster))
            row = ordered[plot["plot"]]
            sequence = " ".join(((by_number[member].get("unit") or {}).get("value") or "") if member in by_number else NOT_WRITTEN if member in missing else OUTSIDE for member in row)
            walls = cut(row, table.get(" ".join((rows.get(member) or {}).get("typologyRev9") or "" for member in row)))
            label = split_label(walls, len(row))
            for start, end in zip([0] + walls, walls + [len(row)]):
                piece = row[start:end]
                for member in piece:
                    # A plot not written holds no piece in the step levels, so it is not counted
                    if step_piece.get(member) is not None:
                        gaps["cluster plots whose piece %s the step piece" % ("is" if step_piece[member] == piece else "is not")] += 1
                rules_ffl = [by_number[member]["levelRule"]["ffl"] for member in piece if member in by_number]
                if rules_ffl:
                    pieces.append(["-".join(row), "-".join(piece), sequence, label, rules_ffl[0], as_written([by_number[member] for member in piece if member in by_number])])
        # Review 2 of PR 91: every front corner of the written plots of the NBH, one basis, a corner with no surface counted apart
        corners = [name for plot in mine_written for name in plot.get("roadSurfaces") or []]
        surfaces = Counter(name for name in corners if name not in (None, "none"))
        zones = sorted({(rows.get(number) or {}).get("zone") for number, found in nbh_of.items() if found == nbh and (rows.get(number) or {}).get("zone")})
        none = sum(1 for name in corners if name in (None, "none"))
        top_line = "NBH %s, %s, road surface %s of %d front corners%s, formulas: " % (
            nbh, " and ".join(zones) or "no zone", ", ".join("%s on %d" % item for item in surfaces.most_common()), len(corners), NO_SURFACE % none if none else "")
        books[nbh] = {"levels": levels, "clusters": pieces, "used": used, "topLine": top_line, "formulas": formulas if rule == PIECE_RULE else None,
                      "file": mock["file"].replace("<NBH>", nbh).replace("<profile>", profile_name),
                      "settings": {"Profile": profile_name, "Rules SHA-256": inputs["sha"]["rules"], "Override list SHA-256": inputs["sha"]["overrides"],
                                   "Road surface order": ", ".join(order), "Pass limit": "none" if passed is None else short(passed) + " mm",
                                   "Editable columns": ", ".join(name for name in mock["columns"] if name in EDITABLE), "Plan build id": plan_build_id(plan),
                                   "Plan SHA-256": hashlib.sha256(inputs["planBytes"]).hexdigest()}}
    return {"books": books, "gaps": gaps, "rule": rule, "written": len(written), "notWritten": len(not_written)}


def row_values(sheet, row, width):
    return [sheet.cell(row=row, column=column).value for column in range(1, width + 1)]


def same(found, wanted):
    """A cell against its value: a number as the same double, a whole step count read back as an int or a float."""
    if isinstance(wanted, (int, float)) and not isinstance(wanted, bool) and isinstance(found, (int, float)) and not isinstance(found, bool):
        return float(found) == float(wanted)
    return found == wanted


def check_xlsx(path, want, mock):
    """Every label, format, fill and cell of one written workbook against the mock and the plan: matches and mismatches by check."""
    import openpyxl
    book = openpyxl.load_workbook(path)
    result, bad = Counter(), []

    def hold(what, ok, where):
        result[("match " if ok else "mismatch ") + what] += 1
        if not ok:
            bad.append(where)

    hold("sheets", book.sheetnames == mock["sheets"], "sheets %s" % book.sheetnames)
    if book.sheetnames != mock["sheets"]:
        return result, bad
    levels = book[mock["sheets"][0]]
    width = len(mock["columns"])
    top = levels.cell(row=1, column=1).value or ""
    hold("top line prefix", top.startswith(want["topLine"]), "Levels row 1 prefix")
    if want["formulas"] is not None:
        hold("top line formulas", top == want["topLine"] + want["formulas"], "Levels row 1 formulas")
    for word in re.split(r", | and ", mock["topLine"]):
        word = re.sub(r"^the | used$", "", word)
        hold("top line names " + word, word.lower() in top.lower(), "Levels row 1 lacks " + word)
    hold("top line bold", bool(levels.cell(row=1, column=1).font.bold), "Levels A1 not bold")
    merged = [str(item) for item in levels.merged_cells.ranges]
    hold("merged only the top line", merged == ["A1:%s1" % openpyxl.utils.get_column_letter(width)], "Levels merged %s" % merged)
    hold("columns", row_values(levels, 2, width) == mock["columns"], "Levels row 2")
    fills = [levels.cell(row=2, column=column).value for column in range(1, width + 1) if levels.cell(row=2, column=column).fill.fill_type == "solid"]
    hold("editable fill", fills == [name for name in mock["columns"] if name in EDITABLE], "fills on %s" % fills)
    colours = {levels.cell(row=2, column=column).fill.fgColor.rgb for column in range(1, width + 1) if levels.cell(row=2, column=column).fill.fill_type == "solid"}
    hold("one fill colour", len(colours) == 1, "fill colours %s" % sorted(colours))
    others = [(name, cell.coordinate) for name in book.sheetnames for row in book[name].iter_rows() for cell in row
              if cell.fill.fill_type == "solid" and not (name == mock["sheets"][0] and cell.row == 2)]
    hold("no other fill", not others, "other fills %s" % others[:5])
    hold("Levels rows", levels.max_row == 2 + len(want["levels"]), "Levels max_row %d against %d" % (levels.max_row, 2 + len(want["levels"])))
    keys = [levels.cell(row=row, column=1).value for row in range(3, levels.max_row + 1)]
    hold("Plot No key unique", len(keys) == len(set(keys)) and all(isinstance(key, str) and key for key in keys), "Plot No keys not unique or blank")
    formats = {name: number_format(mock["decimals"][name]) if name in mock["decimals"] else "0" for name in PLAN_KEYS}
    for index, wanted in enumerate(want["levels"]):
        row = index + 3
        found = row_values(levels, row, width)
        for column, (cell, value) in enumerate(zip(found, wanted)):
            name = mock["columns"][column]
            hold("Levels " + name, same(cell, value), "Levels row %d %s" % (row, name))
            if name in formats and value is not None:
                hold("format " + name, levels.cell(row=row, column=column + 1).number_format == formats[name], "Levels row %d format %s" % (row, name))
    for name, header, rows in ((mock["sheets"][1], CLUSTER_COLUMNS, want["clusters"]), (mock["sheets"][2], OVERRIDE_COLUMNS, want["used"])):
        sheet = book[name]
        hold(name + " header", row_values(sheet, 1, len(header)) == list(header), name + " row 1")
        hold(name + " rows", sheet.max_row == 1 + len(rows), "%s max_row %d against %d" % (name, sheet.max_row, 1 + len(rows)))
        for index, wanted in enumerate(rows):
            found = row_values(sheet, index + 2, len(header))
            for column, (cell, value) in enumerate(zip(found, wanted)):
                hold("%s %s" % (name, header[column]), same(cell, value) and (value is None or isinstance(cell, type(value)) or isinstance(value, (int, float))), "%s row %d %s" % (name, index + 2, header[column]))
            if name == mock["sheets"][2]:
                hold("Overrides used gap a number", isinstance(found[2], (int, float)) and not isinstance(found[2], bool), "%s row %d gap" % (name, index + 2))
    settings = book[mock["sheets"][3]]
    hold("Settings header", row_values(settings, 1, 2) == ["Setting", "Value"], "Settings row 1")
    names = [settings.cell(row=row, column=1).value for row in range(2, settings.max_row + 1)]
    hold("Settings names", names == list(SETTINGS), "Settings names %s" % names)
    values = {settings.cell(row=row, column=1).value: settings.cell(row=row, column=2).value for row in range(2, settings.max_row + 1)}
    for name, value in want["settings"].items():
        hold("Settings " + name, values.get(name) == value, "Settings " + name)
    hold("Settings sha256 lowercase hex", bool(re.fullmatch(r"[0-9a-f]{64}", values.get("Plan SHA-256") or "")), "Settings Plan SHA-256 form")
    hold("Settings run time", bool(re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", values.get("Run time") or "")), "Settings Run time form")
    hold("Settings editable against the fills", values.get("Editable columns") == ", ".join(fills), "Settings Editable columns against the fills")
    return result, bad


def check_folder(folder, want, mock):
    """Each workbook the plan must give in the folder, and no other xlsx there."""
    result, bad = Counter(), []
    names = sorted(path.name for path in Path(folder).glob("*.xlsx"))
    wanted = sorted(book["file"] for book in want["books"].values())
    result["match workbooks" if names == wanted else "mismatch workbooks"] += 1
    if names != wanted:
        bad.append("workbooks %s against %s" % (names, wanted))
    for nbh, book in sorted(want["books"].items()):
        path = Path(folder) / book["file"]
        if path.exists():
            found, missed = check_xlsx(path, book, mock)
            result.update(found)
            bad.extend(nbh + ": " + item for item in missed)
    return result, bad


def run(args):
    mock = mock_labels()
    inputs = load_inputs(args.plan, args.overrides, args.clusters, args.levels, args.hands, args.plots_csv, args.rules, args.config)
    want = expected(inputs, args.profile, mock)
    books = want["books"]
    report = {"plan": Path(args.plan).name, "workbooks": sorted(book["file"] for book in books.values()), "heightRule": want["rule"],
              "rows": {"written": want["written"], "notWritten": want["notWritten"], "levelRows": sum(len(book["levels"]) for book in books.values()),
                       "clusterPieces": sum(len(book["clusters"]) for book in books.values()), "overridesUsed": sum(len(book["used"]) for book in books.values())},
              "gaps": dict(sorted(want["gaps"].items())),
              "sources": dict(Counter(row[-1].split(" ", 1)[0] for book in books.values() for row in book["levels"]))}
    if args.xlsx_folder:
        result, bad = check_folder(args.xlsx_folder, want, mock)
        report["xlsx"] = {"matches": sum(count for key, count in result.items() if key.startswith("match ")),
                          "mismatches": sum(count for key, count in result.items() if key.startswith("mismatch ")),
                          "byCheck": dict(sorted(result.items())), "first": bad[:20]}
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--plan", required=True)
    parser.add_argument("--overrides", required=True)
    parser.add_argument("--clusters", required=True, help="the step clusters file of the build, step_clusters_python.json or its C# twin")
    parser.add_argument("--levels", required=True, help="the step levels file of the build, the piece of each plot")
    parser.add_argument("--hands", required=True, help="the step hands file of the build, the row order of each cluster")
    parser.add_argument("--plots-csv", required=True)
    parser.add_argument("--rules", help="the rules profile, by default the one config names")
    parser.add_argument("--config", default=str(project_profile.SETTINGS_PATH))
    parser.add_argument("--profile", default="parsons2024Rev6")
    parser.add_argument("--xlsx-folder", help="the folder of the level Excel files the worker wrote of this plan, left out to check the plan side alone")
    parser.add_argument("--out", required=True, help="the report, outside the repo")
    args = parser.parse_args(argv)
    if args.rules is None:
        args.rules = str(project_profile.profile_path(project_profile.load_settings(Path(args.config))))
    report = run(args)
    Path(args.out).write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps({key: report[key] for key in ("plan", "rows", "gaps")}))
    if "xlsx" in report:
        print("xlsx matches %d mismatches %d" % (report["xlsx"]["matches"], report["xlsx"]["mismatches"]))
    return 1 if report.get("xlsx", {}).get("mismatches") else 0


if __name__ == "__main__":
    sys.exit(main())
