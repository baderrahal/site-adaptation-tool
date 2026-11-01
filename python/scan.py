"""R420 to R427, round 4j-1. The Python twin of the scanner of the worker command scan. It reads the scan_cad files the worker wrote
from each CAD file the user gives, every sheet of each Excel file, every page of each PDF file and the file names of the prototype folder,
proposes the plot outlines, the plot numbers, the villa blocks and the villa type source, and writes scan_python.json in the run folder in
the shape and order of docs/design/17-scanner.md. Every value is in the scan block of the rules profile, no file, layer or sheet name of a
project is in this file.
"""

import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from shapely.geometry import Point, box
from shapely.strtree import STRtree

import common
import input_guard
import location_plan
import plot_relations
import project_profile

OUTPUT = "scan_python.json"
RECORD = "scan_python_inputs.json"
KINDS = ("cadLayer", "plotOutline", "plotNumber", "villaBlock", "villaType")
LAYER_JOIN = " + "
TIE_JOIN = " or "
# The conflict fields a row counts, R424 on the plot outlines and R426 on the villa type sources, the other kinds count none
# The kinds of villa type source in the order the lead is taken from (R426, Bader 31 Oct 2026)
SOURCE_KINDS = ("block", "text", "excel", "pdf")
FIELDS_OF = {"plotOutline": ("number", "location", "label"), "villaType": ("villa type",)}
USE = {"plotOutline": "plot outlines", "plotNumber": "plot numbers", "villaBlock": "villa blocks", "villaType": "villa type"}


def contains(ring, x, y):
    """R421. The even odd ray test of PlotLabels.Contains of the CAD reader, the same sums in the same order, so a point on an edge
    counts as the C# side counts it."""
    inside = False
    j = len(ring) - 1
    for i in range(len(ring)):
        a, b = ring[i], ring[j]
        if (a[1] > y) != (b[1] > y) and x < (b[0] - a[0]) * (y - a[1]) / (b[1] - a[1]) + a[0]:
            inside = not inside
        j = i
    return inside


def label_number(text, pattern):
    """R421. The number of a label as PlotLabels.Parse takes it: the last line with text, matched whole by labelPattern, its group 2."""
    lines = [line.strip() for line in (text or "").split("\n") if line.strip()]
    if not lines:
        return None
    match = pattern.match(lines[-1])
    return match.group(2) if match else None


def labels_of(texts, pattern):
    """R421. Every text of a file that matches labelPattern, with its number, text, layer and point."""
    found = []
    for text in texts:
        number = label_number(text.get("text"), pattern)
        if number is not None:
            found.append({"number": number, "text": text["text"], "layer": text["layer"], "x": text["x"], "y": text["y"]})
    return found


def face_of(face):
    ring = [(point[0], point[1]) for point in face["ring"]]
    xs, ys = [point[0] for point in ring], [point[1] for point in ring]
    return {"ring": ring, "centroid": (face["centroid"][0], face["centroid"][1]), "box": (min(xs), min(ys), max(xs), max(ys))}


def inside(face, x, y):
    low_x, low_y, high_x, high_y = face["box"]
    return low_x <= x <= high_x and low_y <= y <= high_y and contains(face["ring"], x, y)


class Faces:
    """The faces held in a tree of their boxes, so a point is tested only against the faces whose box holds it, in the order given."""

    def __init__(self, faces):
        self.faces = list(faces)
        self.tree = STRtree([box(*face["box"]) for face in self.faces]) if self.faces else None

    def holding(self, x, y):
        if self.tree is None:
            return []
        found = sorted(int(index) for index in self.tree.query(Point(x, y)))
        return [self.faces[index] for index in found if inside(self.faces[index], x, y)]


def covered(faces, labels):
    """R421. The faces that cover a plot, each holding exactly one label, with its number, label text and label layer, then the label
    numbers of each face that holds two or more, and the labels inside no face."""
    faces = [face_of(face) for face in faces]
    held = defaultdict(list)
    index = Faces(faces)
    outside = []
    for label in labels:
        found = index.holding(label["x"], label["y"])
        for face in found:
            held[id(face)].append(label)
        if not found:
            outside.append(label)
    plots, many = [], []
    for face in faces:
        if len(held[id(face)]) == 1:
            label = held[id(face)][0]
            plots.append(dict(face, number=label["number"], label=label["text"], layer=label["layer"]))
        elif len(held[id(face)]) > 1:
            many.append(sorted(label["number"] for label in held[id(face)]))
    return plots, many, outside


def pick(scores):
    """The source of the highest score above 0, None on a tie or when none scores, and the tied sources in ordinal order."""
    best = max(scores.values(), default=0)
    if best <= 0:
        return None, []
    top = sorted(source for source, score in scores.items() if score == best)
    return (top[0], []) if len(top) == 1 else (None, top)


def check(what, source, plot, value):
    return {"what": what, "source": source, "plot": plot, "value": value}


def tie_check(kind, tied):
    return check("tie", kind, "", TIE_JOIN.join(tied))


def row(kind, source, found, score, proposed):
    return {"kind": kind, "source": source, "found": found, "proposedUse": USE[kind] if proposed else ("" if kind == "cadLayer" else "check"),
            "score": score, "conflicts": 0}


def point_text(point):
    return "%.3f, %.3f" % (common.fixed(float(point[0]), 3), common.fixed(float(point[1]), 3))


def prototype_code(name, prototype_names):
    """R425. The code of a prototype file name by prototypeNames of config, the part after prefix, discipline and a dash up to the suffix,
    as the build reads it, else None."""
    head = prototype_names["prefix"] + prototype_names["discipline"] + "-"
    if not name.startswith(head) or prototype_names["suffix"] not in name[len(head):]:
        return None
    return name[len(head):].split(prototype_names["suffix"])[0] or None


def known_codes(profile, prototype_names, names):
    """R425. The codes a value is matched with: the keys of typologies and their keo_file, the old and new codes of typology_codes and the
    code of each prototype file name."""
    codes = set()
    for code, typology in profile.get("typologies").items():
        codes.add(code)
        if (typology or {}).get("keo_file"):
            codes.add(typology["keo_file"])
    for item in profile.get("typology_codes.rows"):
        codes.update(code for code in (item.get("old"), item.get("new")) if code)
    codes.update(code for code in (prototype_code(name, prototype_names) for name in names) if code)
    return codes


def old_new_pairs(profile):
    """R426, Q245. The old and the new code of each row of typology_codes that differ, as a set of two."""
    return {frozenset((item["old"], item["new"])) for item in profile.get("typology_codes.rows")
            if item.get("old") and item.get("new") and item["old"] != item["new"]}


def value_codes(value, codes, split):
    """R425. The codes a value holds, matched with case, in ordinal order: the whole value when it is a code, else each code found as a
    run bounded by the start, the end or a character of tokenSplit, the longest code first and a matched span never used again, so a code
    that holds a blank or a dash is found whole and not as its shorter head. Two or more make the value a check."""
    text = str(value if value is not None else "").strip()
    if not text:
        return []
    if text in codes:
        return [text]
    bound = re.compile(split)
    edge = lambda at: at < 0 or at >= len(text) or bound.fullmatch(text[at]) is not None
    used = [False] * len(text)
    found = set()
    for code in sorted((code for code in codes if code), key=lambda code: (-len(code), code)):
        start = text.find(code)
        while start >= 0:
            end = start + len(code)
            if edge(start - 1) and edge(end) and not any(used[start:end]):
                used[start:end] = [True] * len(code)
                found.add(code)
            start = text.find(code, start + 1)
    return sorted(found)


def plot_of(text, shift, plots):
    """R426. The CAD label a number text gives under a shift, the number minus the shift, when a plot holds it."""
    text = str(text).strip()
    if not (text.isascii() and text.isdigit()):
        return None
    number = str(int(text) - shift)
    return number if number in plots else None


def cell_text(value):
    """R426. The text of a cell: a number with no fraction as integer text, else the repr of the float."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value)


def sheet_of(title, rows):
    """R426. A sheet as the scanner reads it: the header of each column from the first row with any text, the column letter when that
    cell is empty, and the rows under it."""
    rows = [[cell_text(value) for value in found] for found in rows]
    width = max((len(found) for found in rows), default=0)
    head = next((index for index, found in enumerate(rows) if any(text.strip() for text in found)), None)
    header_row = rows[head] if head is not None else []
    headers = [header_row[column].strip() if column < len(header_row) and header_row[column].strip() else get_column_letter(column + 1) for column in range(width)]
    return {"sheet": title, "headers": headers, "rows": rows[head + 1:] if head is not None else []}


def read_excel(path):
    """R426. Every sheet of a workbook, read only and with the values the file holds, as the readers of the repo open a workbook."""
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    try:
        sheets = [sheet_of(sheet.title, sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets]
    finally:
        workbook.close()
    return {"name": Path(path).name, "sheets": sheets}


def pdf_lines(chars, gap, tolerance):
    """R426, R179. The text lines of a page as the location plan makes them, the upright characters of any size: the same top within
    tolerance and each character starting within gap of the end of the one before, each line with its text, top and centre in page
    points, as LocationPlan.TextLines of the C# side."""
    return location_plan.text_lines([c for c in chars if c.get("upright", True)], (float("-inf"), float("inf")), gap, tolerance)


def read_pdf(path, gap, tolerance):
    """R426. The text lines of every page of a PDF file, by pdfplumber."""
    with pdfplumber.open(str(path)) as pdf:
        pages = [pdf_lines(page.chars, gap, tolerance) for page in pdf.pages]
    return {"name": Path(path).name, "pages": pages}


def number_token(text, split):
    """R426. The one number token of a line, None when it holds none or more than one."""
    found = [token for token in re.split(split, text) if token.isascii() and token.isdigit()]
    return found[0] if len(found) == 1 else None


def page_pairs(lines, codes, split, dy_band, dx_max):
    """R426, R180. Each line that holds a code paired with a line of no code and one number token whose top lies within dy_band under
    its top and whose centre lies within dx_max, nearest first and one to one, by the pairs of location_plan. Returns the code text and
    the number token of each pair."""
    code_lines = [line for line in lines if value_codes(line["text"], codes, split)]
    number_lines = [line for line in lines if not value_codes(line["text"], codes, split) and number_token(line["text"], split) is not None]
    pairs, _, _ = location_plan.pair_labels(code_lines, number_lines, dy_band, dx_max)
    return [(code_lines[i]["text"], number_token(number_lines[k]["text"], split)) for i, k in pairs]


def cad_rows(cad):
    """R420. One cadLayer row per layer of a file."""
    return [row("cadLayer", cad["source"] + "|" + layer["layer"], "%d closed outlines, %d texts, %d blocks" % (layer["closedOutlines"], layer["texts"], layer["blocks"]), 0, False)
            for layer in cad["layers"]]


def outline_proposal(cad, pattern, rows, checks):
    """R421. The plotOutline rows of a file, its proposed source and score and the plots of that source. Each face of the proposed
    outline with two or more labels is a manyLabels check and each label inside no face of it a labelOutside check."""
    labels = labels_of(cad["texts"], pattern)
    plots_of, scores = {}, {}
    for candidate in cad["candidates"]:
        source = cad["source"] + "|" + LAYER_JOIN.join(sorted(candidate["layers"]))
        plots_of[source] = covered(candidate["faces"], labels)
        scores[source] = len(plots_of[source][0])
        rows.append(row("plotOutline", source, "%d faces" % len(candidate["faces"]), scores[source], False))
    proposed, tied = pick(scores)
    if tied:
        checks.append(tie_check("plotOutline", tied))
    # R429. A file the time cap left pairs of proposes no outline, the best pair may be among those left and the notRead row says so
    if proposed is None or cad.get("notTraced"):
        return None
    for found in rows:
        if found["kind"] == "plotOutline" and found["source"] == proposed:
            found["proposedUse"] = USE["plotOutline"]
    plots, many, outside = plots_of[proposed]
    checks.extend(check("manyLabels", proposed, "", ", ".join(numbers)) for numbers in many)
    checks.extend(check("labelOutside", cad["source"] + "|" + label["layer"], label["number"], label["text"]) for label in outside)
    return {"source": proposed, "score": scores[proposed], "plots": plots, "cad": cad}


def number_rows(proposal, rows, checks):
    """R422. A plotNumber row per text layer of the file of a proposed outline, scored by the plots whose label lies on it."""
    cad = proposal["cad"]
    scores = {cad["source"] + "|" + layer: 0 for layer in sorted({text["layer"] for text in cad["texts"]})}
    for plot in proposal["plots"]:
        scores[cad["source"] + "|" + plot["layer"]] += 1
    proposed, tied = pick(scores)
    if tied:
        checks.append(tie_check("plotNumber", tied))
    for source, score in scores.items():
        rows.append(row("plotNumber", source, "%d labels in plots" % score, score, source == proposed))


def block_rows(cads, index, rows, checks):
    """R423. A villaBlock row per insert layer of each file, scored by the plots of the plot drawing that hold an insert point of it."""
    held, counts = defaultdict(set), defaultdict(int)
    for cad in cads:
        for insert in cad["inserts"]:
            source = cad["source"] + "|" + insert["layer"]
            counts[source] += 1
            held[source].update(face["number"] for face in index.holding(insert["x"], insert["y"]))
    scores = {source: len(held[source]) for source in counts}
    proposed, tied = pick(scores)
    if tied:
        checks.append(tie_check("villaBlock", tied))
    for source in sorted(counts):
        rows.append(row("villaBlock", source, "%d inserts" % counts[source], scores[source], source == proposed))


def conflict(plot, field, value_a, source_a, value_b, source_b):
    return {"plot": plot, "field": field, "valueA": value_a, "sourceA": source_a, "valueB": value_b, "sourceB": source_b}


def link_conflicts(drawing, plots, index, other):
    """R424. The plots of another file held against the plots of the plot drawing. Two of one number link when the label text is the
    same and the centroid of each lies inside the other, else a conflict row of each field that fails. A plot whose centroid lies inside
    a plot of the drawing with another number is a conflict of the number. Returns the links and the conflicts."""
    links, found = [], []
    for theirs in other["plots"]:
        ours = plots.get(theirs["number"])
        if ours is not None:
            located = inside(theirs, *ours["centroid"]) and inside(ours, *theirs["centroid"])
            if located and ours["label"] == theirs["label"]:
                links.append((ours["number"], other["source"]))
            if not located:
                found.append(conflict(ours["number"], "location", point_text(ours["centroid"]), drawing, point_text(theirs["centroid"]), other["source"]))
            if ours["label"] != theirs["label"]:
                found.append(conflict(ours["number"], "label", ours["label"], drawing, theirs["label"], other["source"]))
        for face in index.holding(*theirs["centroid"]):
            if face["number"] != theirs["number"]:
                found.append(conflict(face["number"], "number", face["number"], drawing, theirs["number"], other["source"]))
    return links, found


def add_value(sources, kind, source, plot, value):
    if str(value).strip():
        sources[source][plot].append(value)
        sources.kinds[source] = kind


def cad_sources(cads, index, sources):
    """R426. The block names and the texts inside each plot of the plot drawing, by file and by insert or text layer."""
    for cad in cads:
        for insert in cad["inserts"]:
            for face in index.holding(insert["x"], insert["y"]):
                add_value(sources, "block", cad["source"] + "|block|" + insert["layer"], face["number"], insert["block"])
        for text in cad["texts"]:
            for face in index.holding(text["x"], text["y"]):
                add_value(sources, "text", cad["source"] + "|text|" + text["layer"], face["number"], text["text"])


def excel_sources(excel, plots, shifts, sources, checks):
    """R426. Each sheet joined to the plots by the column and the shift that match the most plot numbers, a tie joins none, then every
    other column of the sheet one source."""
    for sheet in excel["sheets"]:
        prefix = excel["name"] + "|" + sheet["sheet"] + "|"
        scores, by_text = {}, {}
        for column, header in enumerate(sheet["headers"]):
            for shift in shifts:
                hits = {plot_of(found[column], shift, plots) for found in sheet["rows"] if column < len(found)} - {None}
                text = prefix + header + "|shift " + str(shift)
                scores[text], by_text[text] = len(hits), (column, shift)
        joined, tied = pick(scores)
        if tied:
            checks.append(tie_check("excelJoin", tied))
        if joined is None:
            continue
        column, shift = by_text[joined]
        for found in sheet["rows"]:
            plot = plot_of(found[column], shift, plots) if column < len(found) else None
            if plot is None:
                continue
            for other, header in enumerate(sheet["headers"]):
                if other != column and other < len(found):
                    add_value(sources, "excel", prefix + header + "|shift " + str(shift), plot, found[other])


def pdf_sources(pdf, plots, block, codes, sources, checks):
    """R426. The code line and number line pairs of R180 of each page of a PDF file, joined to the plots by the number under the shift
    that matches the most plots of the file, by page."""
    split = block["tokenSplit"]
    pages = [page_pairs(page, codes, split, block["pdfPairDyPt"], block["pdfPairDxPt"]) for page in pdf["pages"]]
    scores = {}
    for shift in block["plotNumberShifts"]:
        scores[pdf["name"] + "|shift " + str(shift)] = len({plot_of(token, shift, plots) for pairs in pages for _, token in pairs} - {None})
    chosen, tied = pick(scores)
    if tied:
        checks.append(tie_check("pdfShift", tied))
    if chosen is None:
        return
    shift = int(chosen.rsplit(" ", 1)[1])
    for number, pairs in enumerate(pages, start=1):
        for text, token in pairs:
            plot = plot_of(token, shift, plots)
            if plot is not None:
                add_value(sources, "pdf", "%s|page %d|shift %d" % (pdf["name"], number, shift), plot, text)


def source_codes(source, values_by_plot, codes, split, checks):
    """R425. The one code each plot of a source gives. A value or a plot with two codes is a check row, and a plot that holds such a
    value or two codes gets none from the source."""
    given = {}
    for plot, values in values_by_plot.items():
        found, many = set(), False
        for value in values:
            matched = value_codes(value, codes, split)
            if len(matched) > 1:
                checks.append(check("twoCodes", source, plot, str(value).strip()))
                many = True
            elif matched:
                found.add(matched[0])
        if len(found) > 1:
            checks.append(check("twoCodes", source, plot, ", ".join(sorted(found))))
        elif found and not many:
            given[plot] = found.pop()
    return given


def share_text(share):
    return "%.3f" % common.fixed(share, 3)


def type_rows(given, kinds, prototypes, strong, pairs, rows, conflicts, checks):
    """R426. A villaType row per source whose score is above 0, the score the plots it gives a prototype code. The lead is the source of
    the highest score in the first kind of SOURCE_KINDS with a score above 0, a tie in that kind no lead. It is proposed when its share of
    the plots any source gives a prototype code is at least strong, else every source is a check and noStrongSource names the lead and
    the best of each later kind with their shares. A plot where any other source gives another code than the lead, a profile code too,
    is a conflict, or a check when the two are the old and the new code of one row of typology_codes."""
    scores = {source: sum(1 for code in found.values() if code in prototypes) for source, found in given.items()}
    scores = {source: score for source, score in scores.items() if score > 0}
    by_kind = defaultdict(dict)
    for source, score in scores.items():
        by_kind[kinds[source]][source] = score
    order = [kind for kind in SOURCE_KINDS if kind in by_kind]
    lead, proposed = None, None
    if order:
        lead, tied = pick(by_kind[order[0]])
        if tied:
            checks.append(tie_check("villaType", tied))
    if lead is not None:
        total = len({plot for found in given.values() for plot, code in found.items() if code in prototypes})
        share = scores[lead] / total
        if share >= strong:
            proposed = lead
        else:
            parts = [lead + " " + share_text(share)]
            for kind in order[1:]:
                best = min(by_kind[kind], key=lambda source: (-by_kind[kind][source], source))
                parts.append(best + " " + share_text(by_kind[kind][best] / total))
            checks.append(check("noStrongSource", "villaType", "", " | ".join(parts)))
    for source in sorted(scores):
        rows.append(row("villaType", source, ", ".join(sorted(set(given[source].values()))), scores[source], source == proposed))
    if lead is None:
        return
    for source in sorted(given):
        if source == lead:
            continue
        for plot, code in sorted(given[source].items()):
            ours = given[lead].get(plot)
            if ours is None or ours == code:
                continue
            if frozenset((ours, code)) in pairs:
                checks.append(check("oldNewPair", source, plot, ours + ", " + code))
            else:
                conflicts.append(conflict(plot, "villa type", ours, lead, code, source))


class Sources(defaultdict):
    """The values of each villa type source by plot, with the kind of each source."""

    def __init__(self):
        super().__init__(lambda: defaultdict(list))
        self.kinds = {}


def scan_result(block, codes, pairs, prototypes, strong, cads, excels, pdfs, inputs):
    """R421 to R427. The scan result of the files of a run, every list in the order of docs/design/17-scanner.md. prototypes are the codes
    of the prototype file names, strong the share majority.strongShare."""
    pattern = re.compile(block["labelPattern"])
    split, shifts = block["tokenSplit"], block["plotNumberShifts"]
    rows, conflicts, checks = [], [], []
    proposals = []
    for cad in cads:
        rows.extend(cad_rows(cad))
        checks.extend(check("notRead", item["source"], "", item["value"]) for item in cad.get("notRead") or [])
        proposal = outline_proposal(cad, pattern, rows, checks)
        if proposal is not None:
            proposals.append(proposal)
            number_rows(proposal, rows, checks)
    chosen, tied = pick({proposal["cad"]["source"]: proposal["score"] for proposal in proposals})
    if tied:
        checks.append(tie_check("plotDrawing", tied))
    drawing = next((proposal for proposal in proposals if proposal["cad"]["source"] == chosen), None)
    plots = {}
    if drawing is None:
        checks.append(check("noPlotDrawing", "", "", ""))
    else:
        by_number = defaultdict(list)
        for plot in drawing["plots"]:
            by_number[plot["number"]].append(plot)
        for number, found in by_number.items():
            if len(found) > 1:
                checks.append(check("numberTwice", drawing["source"], number, str(len(found))))
            else:
                plots[number] = found[0]
    index = Faces(plots[number] for number in sorted(plots))
    block_rows(cads, index, rows, checks)
    if drawing is not None:
        for proposal in proposals:
            if proposal is not drawing:
                conflicts.extend(link_conflicts(drawing["source"], plots, index, proposal)[1])
    sources = Sources()
    cad_sources(cads, index, sources)
    for excel in excels:
        excel_sources(excel, plots, shifts, sources, checks)
    for pdf in pdfs:
        pdf_sources(pdf, plots, block, codes, sources, checks)
    given = {source: source_codes(source, sources[source], codes, split, checks) for source in sorted(sources)}
    type_rows(given, sources.kinds, prototypes, strong, pairs, rows, conflicts, checks)
    for found in rows:
        fields = FIELDS_OF.get(found["kind"], ())
        found["conflicts"] = sum(1 for item in conflicts if item["field"] in fields and found["source"] in (item["sourceA"], item["sourceB"]))
    rows.sort(key=lambda item: (KINDS.index(item["kind"]), item["source"]))
    conflicts.sort(key=lambda item: (item["plot"], item["field"], item["sourceA"], item["sourceB"], item["valueA"], item["valueB"]))
    unique = {tuple(item.values()): item for item in checks}
    checks = [unique[found] for found in sorted(unique)]
    plot_rows = [{"plot": number, "label": plots[number]["label"], "centroid": point_text(plots[number]["centroid"]),
                  "villaType": {source: given[source][number] for source in sorted(given) if number in given[source]}} for number in sorted(plots)]
    return {"inputs": inputs, "rows": rows, "conflicts": conflicts, "checks": checks, "plots": plot_rows}


def strong_share(settings):
    """R426. majority.strongShare of config, a run with no share between 0 and 1 there is refused."""
    share = settings.get("majority.strongShare")
    if isinstance(share, bool) or not isinstance(share, (int, float)) or not 0 < share <= 1:
        raise project_profile.ProfileError(settings.path.name + " gives no majority.strongShare between 0 and 1: " + repr(share))
    return share


def load_json(path):
    with open(path, encoding="utf-8-sig") as handle:
        return json.load(handle)


def input_row(role, path):
    found = input_guard.mark(path)
    return {"role": role, "name": Path(path).name, "size": found["size"], "sha256": found["sha256"]}


def run(settings_path, cad_paths, excel_paths, pdf_paths, prototypes, out):
    """R427. One scan under the input guard of D19, every input recorded before and after, the result written once the reads are done."""
    out = Path(out)
    plot_relations.refuse_inside_repo(out.resolve())
    names = sorted(path.name for path in Path(prototypes).iterdir() if path.is_file()) if prototypes else []
    files = [Path(settings_path)] + [Path(path) for path in cad_paths + excel_paths + pdf_paths] + [Path(prototypes) / name for name in names]
    with input_guard.Guard(input_guard.roots(settings_path) + files, record=out / RECORD):
        settings = project_profile.load_settings(settings_path)
        profile = project_profile.load_profile(project_profile.profile_path(settings))
        block = profile.get("scan")
        inputs = [input_row("config", settings_path)]
        cads = [load_json(path) for path in cad_paths]
        inputs += [{"role": "scanCad:%d" % n, "name": cad["source"], "size": cad["size"], "sha256": cad["sha256"]} for n, cad in enumerate(cads, start=1)]
        inputs += [input_row("scanExcel:%d" % n, path) for n, path in enumerate(excel_paths, start=1)]
        inputs += [input_row("scanPdf:%d" % n, path) for n, path in enumerate(pdf_paths, start=1)]
        inputs += [input_row("prototype:" + name, Path(prototypes) / name) for name in names]
        excels = [read_excel(path) for path in excel_paths]
        pdfs = [read_pdf(path, block["pdfJoinGapPt"], block["pdfLineTolerancePt"]) for path in pdf_paths]
        codes = known_codes(profile, settings.get("prototypeNames"), names)
        pairs = old_new_pairs(profile)
        prototypes = {code for code in (prototype_code(name, settings.get("prototypeNames")) for name in names) if code}
        strong = strong_share(settings)
    result = scan_result(block, codes, pairs, prototypes, strong, cads, excels, pdfs, inputs)
    common.write_json(out / OUTPUT, result)
    return result


def main(argv=None):
    parser = argparse.ArgumentParser(description="R420 to R427. Scan the CAD reads, Excel and PDF files a user gives for the plots and the villa type, "
                                                 "and write scan_python.json in the run folder.")
    parser.add_argument("--settings", default=str(project_profile.SETTINGS_PATH))
    parser.add_argument("--cad-read", action="append", default=[], help="a scan_cad_<n>.json of the worker, in the order of n")
    parser.add_argument("--excel", action="append", default=[])
    parser.add_argument("--pdf", action="append", default=[])
    parser.add_argument("--prototypes", help="the folder of the prototype files, read by file name")
    parser.add_argument("--out", required=True, help="the run folder")
    args = parser.parse_args(argv)
    result = run(args.settings, args.cad_read, args.excel, args.pdf, args.prototypes, args.out)
    for kind in KINDS:
        found = [item for item in result["rows"] if item["kind"] == kind]
        proposed = [item["source"] for item in found if item["proposedUse"] not in ("", "check")]
        print("%s: %d rows, proposed %s" % (kind, len(found), ", ".join(proposed) or "none"))
    print("conflicts %d, checks %d, plots %d" % (len(result["conflicts"]), len(result["checks"]), len(result["plots"])))
    return 0


if __name__ == "__main__":
    sys.exit(main())
